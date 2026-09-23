import importlib.util
import io
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook


MODULE_PATH = Path(__file__).parents[1] / "hrms" / "hr" / "training_importer.py"
SPEC = importlib.util.spec_from_file_location("training_importer", MODULE_PATH)
IMPORTER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(IMPORTER)


def workbook_bytes(workbook):
	output = io.BytesIO()
	workbook.save(output)
	return output.getvalue()


class TrainingWorkbookImportTest(unittest.TestCase):
	def test_plan_total_sheet_preserves_source_fields_and_audience_matrix(self):
		workbook = Workbook()
		sheet = workbook.active
		sheet.title = IMPORTER.PLAN_SHEET
		sheet["B2"] = "2026年度教育训练工作表（内培+外培）"
		headers = ["部门", "分类", "培训类型", "培训内容", "内/外", "课时", "召集\n人员", "召集\n部门", "地点", "主要培训岗位/人员", "预计上课时间\n（月份）"]
		for column, label in enumerate(headers, start=2):
			sheet.cell(3, column).value = label
		sheet["P3"] = "连续课"
		sheet.merge_cells("P3:R3")
		sheet["P4"] = "生产组"
		sheet["Q4"] = "保养组"
		sheet["R4"] = "维护组"
		sheet["AH3"] = "备注"
		values = ["工程", "计划", "岗位应知应会", "点镀作业规范", "内", 1, "陆体廷", "工程", "会议室", "直接人员", "1月"]
		for column, value in enumerate(values, start=2):
			sheet.cell(5, column).value = value
		sheet["M5"] = datetime(2026, 1, 10)
		sheet["P5"] = "●"
		sheet["AH5"] = "原表备注"

		result = IMPORTER.parse_plan_workbook(workbook_bytes(workbook))
		self.assertEqual(len(result["rows"]), 1)
		row = result["rows"][0]
		self.assertEqual(row["actual_dates"], ["2026-01-10"])
		self.assertEqual(row["audience_matrix"], [{"unit": "连续课 / 生产组", "requirement": "●"}])
		self.assertEqual(row["remarks"], "原表备注")
		self.assertTrue(row["source_key"].startswith("TRAIN-PLAN-"))

	def test_record_summary_groups_people_into_one_event_and_reads_compact_dates(self):
		workbook = Workbook()
		sheet = workbook.active
		sheet.title = IMPORTER.RECORD_SHEET
		headers = ["序号", "月份", "实际上课时间", "部门", "姓名", "课程类型", "培训内容", "课程归\n属部门", "内/\n外", "课件\n方式", "课时", "学时", "授课人", "地点", "培训对象", "成绩", "备注（评价标准）"]
		for column, label in enumerate(headers, start=3):
			sheet.cell(3, column).value = label
		for offset, (name, code) in enumerate((("张三", 85), ("李四", None)), start=4):
			values = [offset - 3, 8, "260825\n/260831", "连续课", name, "岗位应知应会", "导线连接方式讲解", "连续课", "内", "实操", 0.5, 0.75, "李旭", "会议室", "连续线长", code, ""]
			for column, value in enumerate(values, start=3):
				sheet.cell(offset, column).value = value

		result = IMPORTER.parse_record_workbook(workbook_bytes(workbook))
		self.assertEqual(len(result["rows"]), 2)
		self.assertEqual(len(result["events"]), 1)
		self.assertEqual(result["events"][0]["actual_dates"], ["2026-08-25", "2026-08-31"])
		self.assertEqual(len(result["events"][0]["participants"]), 2)
		self.assertFalse(result["events"][0]["errors"])

	def test_record_sheet_is_detected_by_headers_when_title_changes(self):
		workbook = Workbook()
		sheet = workbook.active
		sheet.title = "26年教育训练登记表"
		headers = ["序号", "月份", "实际上课时间", "部门", "姓名", "课程类型", "培训内容", "课程归\n属部门", "内/\n外", "课件\n方式", "课时", "学时", "授课人", "地点", "培训对象", "成绩", "备注（评价标准）"]
		for column, label in enumerate(headers, start=1):
			sheet.cell(1, column).value = label
		values = [1, 1, datetime(2026, 1, 5), "量试组", "陆卫国", "会议类", "周会", "量试组", "内", "excel", 0.5, 0.75, "时雷", "会议室", "量试组全体", None, None]
		for column, value in enumerate(values, start=1):
			sheet.cell(2, column).value = value

		result = IMPORTER.parse_record_workbook(workbook_bytes(workbook))
		self.assertEqual(result["sheet_name"], "26年教育训练登记表")
		self.assertEqual(result["header_row"], 1)
		self.assertEqual(len(result["rows"]), 1)
		self.assertEqual(result["events"][0]["source_sheet"], "26年教育训练登记表")

	def test_preview_token_changes_when_source_identity_changes(self):
		plan_rows = [{"source_key": "PLAN-1"}]
		rows = [{"source_row": 4, "source_serial": "1", "identity_key": "张三|工程课"}]
		first = IMPORTER.preview_token("永新", "a", "b", plan_rows, rows)
		rows[0]["identity_key"] = "李四|工程课"
		second = IMPORTER.preview_token("永新", "a", "b", plan_rows, rows)
		self.assertNotEqual(first, second)

	def test_identical_plan_rows_keep_distinct_source_keys(self):
		workbook = Workbook()
		sheet = workbook.active
		sheet.title = IMPORTER.PLAN_SHEET
		sheet["B2"] = "2026年度教育训练工作表"
		headers = ["部门", "分类", "培训类型", "培训内容", "内/外", "课时", "召集人员", "召集部门", "地点", "主要培训岗位/人员", "预计上课时间（月份）"]
		for column, label in enumerate(headers, start=2):
			sheet.cell(3, column).value = label
		for row_number in (5, 6):
			for column, value in enumerate(["工程", "计划", "部门内月会", "月会", "内", 1, "张三", "工程", "会议室", "工程全体", "1月"], start=2):
				sheet.cell(row_number, column).value = value

		rows = IMPORTER.parse_plan_workbook(workbook_bytes(workbook))["rows"]
		self.assertEqual(len(rows), 2)
		self.assertNotEqual(rows[0]["source_key"], rows[1]["source_key"])


if __name__ == "__main__":
	unittest.main()
