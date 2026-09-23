import ast
import importlib.util
import io
import unittest
from pathlib import Path

from openpyxl import Workbook


MODULE_PATH = Path(__file__).parents[1] / "hrms" / "hr" / "employee_relationship_importer.py"
API_PATH = Path(__file__).parents[1] / "hrms" / "hr" / "page" / "employee_relationship" / "employee_relationship.py"
SPEC = importlib.util.spec_from_file_location("employee_relationship_importer", MODULE_PATH)
IMPORTER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(IMPORTER)


def workbook_bytes(workbook):
	output = io.BytesIO()
	workbook.save(output)
	return output.getvalue()


def source_workbook(relationship="同村"):
	workbook = Workbook()
	sheet = workbook.active
	sheet.title = "人员关系表——在职"
	headers = ["", "籍贯", "部门", "姓名", "聘用日期", "性别", "年龄", "部门", "姓名", "聘用日期", "性别", "年龄", "关系", "更新人员/时间"]
	for column, value in enumerate(headers, start=1):
		sheet.cell(3, column).value = value
	values = ["", "河南", "连续课", "张三", "2024-01-01", "男", 30, "品管课", "李四", "2024-02-01", "女", 28, relationship, "余琳"]
	for column, value in enumerate(values, start=1):
		sheet.cell(4, column).value = value
	return workbook


class EmployeeRelationshipWorkbookImportTest(unittest.TestCase):
	def test_apply_endpoint_annotates_every_request_parameter(self):
		module = ast.parse(API_PATH.read_text(encoding="utf-8"))
		function = next(
			node
			for node in module.body
			if isinstance(node, ast.FunctionDef) and node.name == "apply_employee_relationship_import"
		)
		missing = [argument.arg for argument in function.args.args if argument.annotation is None]
		self.assertEqual(missing, [])

	def test_parses_existing_two_person_layout_without_guessing_codes(self):
		result = IMPORTER.parse_employee_relationship_workbook(workbook_bytes(source_workbook()))
		self.assertEqual(result["sheet_names"], ["人员关系表——在职"])
		self.assertEqual(len(result["rows"]), 1)
		row = result["rows"][0]
		self.assertEqual(row["employee_a_name"], "张三")
		self.assertEqual(row["employee_b_name"], "李四")
		self.assertEqual(row["relationship"], "同村")
		self.assertEqual(row["employee_a_code"], "")
		self.assertEqual(row["employee_a_identity"], "姓名部门:张三|连续课")
		self.assertFalse(row["errors"])

	def test_rejects_relationship_outside_business_categories(self):
		row = IMPORTER.parse_employee_relationship_workbook(workbook_bytes(source_workbook("直属上级")))["rows"][0]
		self.assertTrue(any("不在允许范围" in error for error in row["errors"]))

	def test_prefers_explicit_company_codes_when_template_contains_them(self):
		workbook = Workbook()
		sheet = workbook.active
		sheet.title = "员工关系导入"
		headers = ["员工一工号", "员工一姓名", "部门", "员工二工号", "员工二姓名", "部门", "员工关系大类"]
		for column, value in enumerate(headers, start=1):
			sheet.cell(1, column).value = value
		for column, value in enumerate(["A001", "张三", "连续课", "B002", "李四", "品管课", "朋友"], start=1):
			sheet.cell(2, column).value = value
		row = IMPORTER.parse_employee_relationship_workbook(workbook_bytes(workbook))["rows"][0]
		self.assertEqual(row["employee_a_identity"], "工号:A001")
		self.assertEqual(row["employee_b_identity"], "工号:B002")

	def test_preview_token_changes_with_relationship_category(self):
		rows = IMPORTER.parse_employee_relationship_workbook(workbook_bytes(source_workbook()))["rows"]
		first = IMPORTER.preview_token("永新", "digest", rows)
		rows[0]["relationship"] = "朋友"
		second = IMPORTER.preview_token("永新", "digest", rows)
		self.assertNotEqual(first, second)

	def test_duplicate_pair_with_conflicting_categories_requires_resolution(self):
		workbook = source_workbook("朋友")
		sheet = workbook.active
		for column in range(1, 15):
			sheet.cell(5, column).value = sheet.cell(4, column).value
		sheet.cell(5, 13).value = "旁系亲属"
		rows = IMPORTER.parse_employee_relationship_workbook(workbook_bytes(workbook))["rows"]
		conflicts, duplicates = IMPORTER.summarize_source_conflicts(rows)
		self.assertEqual(duplicates, 0)
		self.assertEqual(len(conflicts), 1)
		self.assertEqual(conflicts[0]["categories"], ["旁系亲属", "朋友"])


if __name__ == "__main__":
	unittest.main()
