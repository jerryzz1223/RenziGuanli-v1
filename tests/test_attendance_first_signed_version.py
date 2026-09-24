from __future__ import annotations

import importlib.util
import io
import sys
import types
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
API_PATH = ROOT / "hrms" / "api" / "attendance_processing_center.py"


def load_processing_api():
	frappe = types.ModuleType("frappe")
	frappe._ = lambda value, *args: value
	frappe.whitelist = lambda function=None: function or (lambda decorated: decorated)
	frappe.throw = lambda message, *args, **kwargs: (_ for _ in ()).throw(RuntimeError(message))
	frappe.db = types.SimpleNamespace()
	frappe.get_all = lambda *args, **kwargs: []
	frappe.get_doc = lambda *args, **kwargs: None
	frappe.get_roles = lambda *args, **kwargs: []
	frappe.session = types.SimpleNamespace(user="test@example.com")
	frappe.DoesNotExistError = type("DoesNotExistError", (Exception,), {})
	frappe_utils = types.ModuleType("frappe.utils")
	frappe_utils.cint = lambda value: int(value or 0)
	frappe_utils.flt = lambda value: float(value or 0)
	frappe_utils.getdate = lambda value: __import__("datetime").date.fromisoformat(str(value)[:10])
	frappe_utils.now_datetime = lambda: __import__("datetime").datetime.now()
	file_manager = types.ModuleType("frappe.utils.file_manager")
	export_watermark = types.ModuleType("hrms.utils.export_watermark")
	file_manager.save_file = lambda *args, **kwargs: types.SimpleNamespace(file_url="/private/files/first-signed.xlsx", file_name=args[0])
	export_watermark.save_workbook_with_logo_watermark = lambda book, output: book.save(output)
	frappe_utils.file_manager = file_manager
	hrms = types.ModuleType("hrms")
	hrms.__path__ = [str(ROOT / "hrms")]
	hrms_api = types.ModuleType("hrms.api")
	hrms_api.__path__ = [str(ROOT / "hrms" / "api")]
	hrms_processors = types.ModuleType("hrms.api.attendance_processors")
	hrms_processors.__path__ = [str(ROOT / "hrms" / "api" / "attendance_processors")]
	hrms_utils = types.ModuleType("hrms.utils")
	hrms_utils.__path__ = [str(ROOT / "hrms" / "utils")]
	sys.modules.update({
		"frappe": frappe,
		"frappe.utils": frappe_utils,
		"frappe.utils.file_manager": file_manager,
		"hrms": hrms,
		"hrms.api": hrms_api,
		"hrms.api.attendance_processors": hrms_processors,
		"hrms.utils": hrms_utils,
		"hrms.utils.export_watermark": export_watermark,
	})
	spec = importlib.util.spec_from_file_location("attendance_processing_center_first_signed_test", API_PATH)
	module = importlib.util.module_from_spec(spec)
	sys.modules[spec.name] = module
	spec.loader.exec_module(module)
	return module


class TestAttendanceFirstSignedVersion(unittest.TestCase):
	@classmethod
	def setUpClass(cls):
		cls.api = load_processing_api()

	def test_rows_use_attendance_population_and_keep_missing_card_amount(self):
		api = self.api
		attendance = types.SimpleNamespace(source_type="attendance_draft", company="测试公司")
		missing = types.SimpleNamespace(source_type="missing_card", company="测试公司")
		special = types.SimpleNamespace(source_type="special_hours", company="测试公司", attendance_month="2026-07")
		api._result_rows = lambda batch, *args, **kwargs: (
			[{
				"employee_code": "E-001", "employee_name": "张三", "department": "工程课",
				"eligible_for_downstream": True, "processed_value": {
					"standard_hours": 176, "actual_attendance_hours": 168, "workday_overtime_hours": 2,
				},
			}] if batch.source_type == "attendance_draft" else [{
				"employee_code": "E-001", "employee_name": "张三", "department": "工程课",
				"eligible_for_downstream": True, "processed_value": {"red_apples": 1, "amount": 5},
			}] if batch.source_type == "missing_card" else [{
				"employee_code": "E-001", "employee_name": "张三", "department": "工程课",
				"eligible_for_downstream": True, "processed_value": {"special_hours": 2, "special_hours_days": []},
			}]
		)
		api._special_hours_breakdown = lambda *args, **kwargs: {"special_workday_hours": 2, "special_restday_hours": 0, "special_holiday_hours": 0}
		api._employee_directory = lambda company: [{"employee_code": "E-001", "date_of_joining": "2008-06-16"}]

		rows = api._monthly_first_signed_rows({"attendance_draft": attendance, "missing_card": missing, "special_hours": special})

		self.assertEqual(len(rows), 1)
		self.assertEqual(rows[0]["sequence"], 1)
		self.assertEqual(rows[0]["date_of_joining"], "2008-06-16")
		self.assertEqual(rows[0]["red_apple_amount"], 5)
		self.assertEqual(rows[0]["special_workday_hours"], 2)

	def test_confirmed_special_hours_override_same_day_derived_from_attendance(self):
		api = self.api
		attendance = types.SimpleNamespace(
			source_type="attendance_draft", company="测试公司", attendance_month="2026-07",
		)
		special = types.SimpleNamespace(
			source_type="special_hours", company="测试公司", attendance_month="2026-07",
		)
		api._result_rows = lambda batch, *args, **kwargs: [{
			"employee_code": "E-001", "employee_name": "张三", "department": "工程课",
			"eligible_for_downstream": True,
			"processed_value": {
				"standard_hours": 176, "actual_attendance_hours": 176,
				"special_hours": 1, "special_hours_days": [{"day": 1, "hours": 1}],
			},
		}] if batch.source_type == "attendance_draft" else [{
			"employee_code": "E-001", "employee_name": "张三", "department": "工程课",
			"eligible_for_downstream": True,
			"processed_value": {
				"special_hours": 2.5,
				"special_hours_days": [{"day": 1, "hours": 0.5}, {"day": 2, "hours": 2}],
			},
		}]
		api._special_hours_breakdown = lambda entries, *_args: {
			"special_workday_hours": sum(entry["hours"] for entry in entries),
			"special_restday_hours": 0,
			"special_holiday_hours": 0,
		}

		rows = api._monthly_final_rows({"attendance_draft": attendance, "special_hours": special})

		self.assertEqual(rows[0]["special_hours_days"], [{"day": 1, "hours": 0.5}, {"day": 2, "hours": 2.0}])
		self.assertEqual(rows[0]["special_workday_hours"], 2.5)

	def test_workbook_matches_first_signature_layout(self):
		api = self.api
		captured = {}

		def capture_save(name, content, *args, **kwargs):
			captured["name"] = name
			captured["content"] = content
			return types.SimpleNamespace(file_url="/private/files/first-signed.xlsx", file_name=name)

		api_save_file = sys.modules["frappe.utils.file_manager"].save_file
		sys.modules["frappe.utils.file_manager"].save_file = capture_save
		try:
			daily_row = ["张三", "E-001", "26-07-01 星期三", "工程课", "工作日", "长白班", "08:01", "17:00"] + [None] * 32 + ["08:00", "17:00", 0, 0, "无申请", "2026-07-01迟到1分钟（半小时以内）"]
			result = api._save_monthly_first_signed_confirmation_file("2026-07", [{
				"sequence": 1, "department": "工程课", "employee_name": "张三", "employee_code": "E-001",
				"date_of_joining": "2008-06-16", "standard_hours": 176, "actual_attendance_hours": 168,
				"special_workday_hours": 2, "workday_overtime_hours": 2, "large_night_shifts": 1, "absence_hours": 0,
			}], [daily_row])
		finally:
			sys.modules["frappe.utils.file_manager"].save_file = api_save_file

		self.assertEqual(result["file_name"], "2026-07_一次签字版.xlsx")
		book = load_workbook(io.BytesIO(captured["content"]), data_only=False)
		self.assertEqual(book.sheetnames, ["每日统计", "工时汇总"])
		daily = book["每日统计"]
		self.assertEqual(daily.freeze_panes, "B3")
		self.assertEqual(daily["A3"].value, "张三")
		self.assertIn("T1:AD1", {str(item) for item in daily.merged_cells.ranges})
		self.assertEqual(daily.max_column, 46)
		self.assertEqual(daily["AO1"].value, "计划上班")
		self.assertEqual(daily["AT3"].value, "2026-07-01迟到1分钟（半小时以内）")
		sheet = book["工时汇总"]
		self.assertIn("B2:Z2", {str(item) for item in sheet.merged_cells.ranges})
		self.assertIn("I3:K3", {str(item) for item in sheet.merged_cells.ranges})
		self.assertIn("R3:X3", {str(item) for item in sheet.merged_cells.ranges})
		self.assertEqual(sheet.freeze_panes, "C5")
		self.assertEqual(sheet["B2"].value, "7月工时汇总")
		self.assertEqual(sheet["E5"].value, "E-001")
		self.assertEqual(sheet["G5"].value, 176)
		self.assertEqual(sheet["I5"].value, 2)
		self.assertIsNone(sheet["J5"].value)
		self.assertIsNone(sheet["K5"].value)
		self.assertEqual(sheet.max_column, 29)

	def test_daily_projection_exports_calculated_and_confirmed_overtime_separately(self):
		api = self.api
		batch = types.SimpleNamespace(source_type="attendance_draft")
		api._result_rows = lambda *_args, **_kwargs: [{
			"employee_code": "E-001", "employee_name": "张三", "department": "工程课", "eligible_for_downstream": True,
			"source_file": "sample.xlsx", "source_sheet": "每日统计",
			"original_value": {"rows": [{
				"姓名": "张三", "工号": "E-001", "日期": "26-07-01", "实际部门": "工程课", "工作类型": "工作日",
				"班次": "白班 08:00-17:00", "上班时间": "08:30", "下班时间": "18:00", "source_file": "sample.xlsx",
				"source_sheet": "每日统计", "source_row": 3,
			}]},
			"processed_value": {"attendance_details": [{
				"attendance_date": "2026-07-01", "source_file": "sample.xlsx", "source_sheet": "每日统计", "source_row": 3,
				"personal_leave_hours": 0.5, "late_count": 1, "scheduled_start": "08:00", "scheduled_end": "17:00",
				"raw_outside_shift_hours": 1.75, "calculated_workday_overtime_hours": 1.5,
				"confirmed_overtime_hours": 0.75, "overtime_approval_status": "人工确认",
				"attendance_note": "2026-07-01迟到30分钟（半小时以内）",
			}]},
		}]

		rows = api._monthly_first_signed_daily_rows({"attendance_draft": batch})

		self.assertEqual(rows[0][14], 1.5)
		self.assertEqual(rows[0][19], 0.5)
		self.assertEqual(rows[0][4], "工作日")
		self.assertEqual(rows[0][38], 1)
		self.assertEqual(rows[0][40:46], ["08:00", "17:00", 1.75, 0.75, "人工确认", "2026-07-01迟到30分钟（半小时以内）"])

	def test_second_signed_workbook_matches_supplied_header_contract(self):
		api = self.api
		captured = {}
		file_manager = sys.modules["frappe.utils.file_manager"]
		old_save_file = file_manager.save_file
		file_manager.save_file = lambda name, content, *args, **kwargs: (captured.update(name=name, content=content) or types.SimpleNamespace(file_url="/private/files/second-signed.xlsx", file_name=name))
		try:
			api._save_monthly_signed_confirmation_file("2026-07", [{
				"employee_code": "E-001", "employee_name": "张三", "department": "工程课", "standard_hours": 176,
				"actual_attendance_hours": 168, "reunion_leave_hours": 40, "special_workday_hours": 2,
				"review_note": "2026-07-01迟到30分钟（半小时以内）",
			}])
		finally:
			file_manager.save_file = old_save_file

		self.assertEqual(captured["name"], "2026-07_第二次员工签字版.xlsx")
		book = load_workbook(io.BytesIO(captured["content"]), data_only=False)
		sheet = book["第二次员工签字版"]
		self.assertIn("D1:BJ1", {str(item) for item in sheet.merged_cells.ranges})
		self.assertIn("J2:K2", {str(item) for item in sheet.merged_cells.ranges})
		self.assertIn("AE2:AG2", {str(item) for item in sheet.merged_cells.ranges})
		self.assertIn("B6:H6", {str(item) for item in sheet.merged_cells.ranges})
		self.assertEqual(sheet["D1"].value, "7月工时奖惩确认表")
		self.assertEqual(sheet["AG3"].value, "团圆假\n工时")
		self.assertEqual(sheet["BJ2"].value, "备注")
		self.assertEqual(sheet["BJ5"].value, "2026-07-01迟到30分钟（半小时以内）")
		self.assertEqual(sheet["X5"].value, 40)
		self.assertEqual(sheet["AJ5"].value, "=J5+K5")
		self.assertEqual(sheet["AR4"].value, 7)
		self.assertEqual(sheet["BH4"].value, 20)
		self.assertEqual(sheet.freeze_panes, "J4")
		self.assertEqual(sheet["R10"].value, "审核：")
		self.assertEqual(sheet["AS10"].value, "审核：")
		self.assertEqual(sheet["BA10"].value, "复核：")
		self.assertEqual(sheet["BH10"].value, "制表：李微微2026.6.11")

	def test_all_disabled_company_shift_rules_do_not_use_builtin_fallback(self):
		api = self.api
		row = {
			"name": "RULE-1", "enabled": 0, "rule_code": "SHIFT-002", "rule_name": "生产夜班",
			"sequence": 2, "match_tokens": "生产|夜班", "basic_time": "20:00-4:30",
			"weekday_overtime_time": "4:30-8:00", "weekday_overtime_hours": 3.5,
			"weekday_overtime_mode": "不提交加班单", "weekend_overtime_mode": "不提交加班单",
			"punch_in_range": "18:00-20:29", "punch_out_range": "20:30-10:00",
		}
		old_get_all = api.frappe.get_all
		old_exists = getattr(api.frappe.db, "exists", None)
		try:
			api.frappe.get_all = lambda *args, **kwargs: [row]
			api.frappe.db.exists = lambda *args, **kwargs: True
			bundle = api._attendance_shift_rule_bundle("永新")
		finally:
			api.frappe.get_all = old_get_all
			if old_exists is None:
				delattr(api.frappe.db, "exists")
			else:
				api.frappe.db.exists = old_exists

		self.assertEqual(bundle["rules"], [])
		self.assertFalse(bundle["version"].startswith("builtin-"))

	def test_complete_rule_center_returns_shift_policy_and_system_rules(self):
		api = self.api
		fake_import = types.ModuleType("hrms.api.attendance_import")
		fake_import.list_attendance_custom_rules = lambda page_length=0: [{
			"rule_code": "ATT-LATE-30", "rule_name": "迟到提示", "enabled": 1,
		}]
		old_import = sys.modules.get("hrms.api.attendance_import")
		old_bundle = api._attendance_shift_rule_bundle
		old_scheduling_policies = api.list_attendance_scheduling_policies
		old_permission = api._require_processing_manager
		old_company = api._require_company
		try:
			sys.modules["hrms.api.attendance_import"] = fake_import
			api._attendance_shift_rule_bundle = lambda company: {
				"items": [{"rule_code": "SHIFT-001", "rule_name": "生产白班"}],
				"rules": [{"rule_code": "SHIFT-001"}], "version": "rules-v1",
			}
			api.list_attendance_scheduling_policies = lambda company: {"items": []}
			api._require_processing_manager = lambda: None
			api._require_company = lambda company: company
			result = api.get_complete_attendance_rules("永新")
		finally:
			api._attendance_shift_rule_bundle = old_bundle
			api.list_attendance_scheduling_policies = old_scheduling_policies
			api._require_processing_manager = old_permission
			api._require_company = old_company
			if old_import is None:
				del sys.modules["hrms.api.attendance_import"]
			else:
				sys.modules["hrms.api.attendance_import"] = old_import

		self.assertEqual(result["shift_rule_version"], "rules-v1")
		self.assertEqual(result["shift_rules"][0]["rule_code"], "SHIFT-001")
		self.assertEqual(result["policy_rules"][0]["rule_code"], "ATT-LATE-30")
		self.assertTrue(any(row["rule_name"] == "中班" for row in result["builtin_shift_rules"]))
		self.assertEqual(len(result["system_boundaries"]), 6)
		self.assertTrue(any(item["name"] == "周末与调班边界" for item in result["system_boundaries"]))
		self.assertTrue(any(item["name"] == "周末未排班中班夜班" for item in result["system_boundaries"]))

	def test_shift_match_preview_uses_active_matcher_without_writing(self):
		api = self.api
		old_bundle = api._attendance_shift_rule_bundle
		old_permission = api._require_processing_manager
		old_company = api._require_company
		try:
			api._attendance_shift_rule_bundle = lambda company: {"rules": None, "version": "builtin-test"}
			api._require_processing_manager = lambda: None
			api._require_company = lambda company: company
			result = api.preview_attendance_shift_match("永新", "中班 13:00-22:00", "2026-07-05")
			api._attendance_shift_rule_bundle = lambda company: {"rules": [], "version": "company-disabled"}
			disabled = api.preview_attendance_shift_match("永新", "中班 13:00-22:00", "2026-07-05")
		finally:
			api._attendance_shift_rule_bundle = old_bundle
			api._require_processing_manager = old_permission
			api._require_company = old_company
		self.assertTrue(result["matched"])
		self.assertEqual(result["rule_name"], "中班")
		self.assertEqual(result["source"], "内置兼容规则")
		self.assertEqual(result["rule_version"], "builtin-test")
		self.assertFalse(disabled["matched"])


if __name__ == "__main__":
	unittest.main()
