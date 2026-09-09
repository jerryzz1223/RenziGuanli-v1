"""Rollback-only XLSX round trips, portable identities and validation failures."""
import copy
import io
import json

import frappe
from openpyxl import load_workbook

from hrms.api import organization_package as package
from hrms.api.organization_roster_sync import reconcile, enabled_companies


def run():
	chart = package.chart_module()
	before = frappe.get_all("Employee", fields=["name", "custom_employee_code", "employee_name", "department", "designation", "status"], order_by="name")
	node_count = frappe.db.count("Organization Node")
	try:
		from scripts.merge_organization_business_group import run as merge
		merge()
		assert not merge()["changed"]
		file = package.export_configuration("永新")
		data = package.read_package(file["file_url"])
		plan = package.prepare("永新", data)
		assert not plan["errors"], plan["errors"]
		old_cards = chart._build_manual_organization_tree("永新", chart._get_manual_organization_records("永新"))
		def people(root):
			result = {p["employee"] for p in root.get("people", []) if p.get("matched_employee")}
			for child in root.get("children", []): result |= people(child)
			return result
		before_people = people(old_cards)
		result = package.apply_configuration("永新", file["file_url"], plan["fingerprint"])
		assert result["created"] == 0 and result["updated"] == len(data["组织层级"])
		assert "永新" in enabled_companies()
		reconcile("永新")
		assert reconcile("永新")["changed"] == 0
		new_cards = chart._build_manual_organization_tree("永新", chart._get_manual_organization_records("永新"))
		assert people(new_cards) == before_people, (len(people(new_cards)), len(before_people))
		assert new_cards["current_headcount"] == old_cards["current_headcount"] == 199
		plan = package.preview_configuration("永新", file["file_url"])
		package.apply_configuration("永新", file["file_url"], plan["fingerprint"])
		assert len(chart._get_manual_organization_records("永新")["nodes"]) == len(data["组织层级"])
		assert frappe.get_all("Employee", fields=list(before[0]), order_by="name") == before
		print("PASS: full XLSX round trip, all 199 people, acting/proxy references, manual flags, repeat upsert, no Employee writes")

		dept = frappe.get_doc({"doctype": "Department", "department_name": "配置迁移测试课", "company": "永新", "parent_department": "All Departments"}).insert()
		job = frappe.get_doc({"doctype": "Designation", "designation_name": "配置迁移测试岗位"}).insert()
		emp = frappe.get_doc({"doctype": "Employee", "custom_employee_code": "9999090919", "first_name": "配置迁移测试", "gender": "Male", "date_of_birth": "1990-01-01", "date_of_joining": "2026-09-01", "company": "永新", "department": dept.name, "designation": job.name, "status": "Active"}).insert()
		# A server's Employee document ID is deliberately different from the business code.
		code = "000099990919"
		frappe.db.set_value("Employee", emp.name, "custom_employee_code", code)
		assert emp.name != code
		book = load_workbook(io.BytesIO(package.workbook_bytes("永新")))
		book["说明"]["B2"] = "来源服务器公司"
		for sheet in ("组织层级", "人员任职", "职级定义"):
			book[sheet].delete_rows(2, book[sheet].max_row)
		def node(key, parent, kind, label, **fields):
			row = {"portable_id": key, "parent": parent, "node_kind": kind, "display_name": label, "assignment_mode": "正式", **fields}
			book["组织层级"].append([row.get(k, 0 if k in package.BOOL_FIELDS else "") for k in package.NODE_COLUMNS.values()])
		node("SOURCE-UNIT", "", "课", "配置迁移测试课", department=dept.department_name, chart_grade_code="M1")
		node("SOURCE-JOB", "SOURCE-UNIT", "岗位", "测试职位（代）", department=dept.department_name, designation=job.name, chart_grade_code="M2", role_title="组长（代）")
		book["人员任职"].append(["SOURCE-JOB", "岗位成员", code, "来源服务器姓名", "", "", 0])
		book["职级定义"].append(["M1", "管理职级", 1, ""])
		book["职级定义"].append(["M2", "岗位职级", 2, "M1"])
		def save_book():
			stream = io.BytesIO(); book.save(stream)
			return package.save_private(stream.getvalue(), "配置测试", "永新")["file_url"]
		url = save_book()
		portable = package.read_package(url)
		plan = package.prepare("永新", portable)
		assert not plan["errors"], plan["errors"]
		assert plan["create_count"] == 2
		assert plan["_plans"]["SOURCE-JOB"]["config"]["assigned_employees"] == [emp.name]
		package.apply_configuration("永新", url, plan["fingerprint"])
		nodes = chart._get_manual_organization_records("永新")["nodes"]
		unit = next(n for n in nodes if n.manual_config.get("portable_id") == "SOURCE-UNIT")
		role = next(n for n in nodes if n.manual_config.get("portable_id") == "SOURCE-JOB")
		assert role.parent_node == unit.name and role.name != "SOURCE-JOB"
		assert role.manual_config["chart_grade"] == {"code": "M2", "label": "岗位职级", "rank": 2, "parent": "M1"}
		chart.save_manual_organization_node("岗位", node_name=role.name, company="永新", department=dept.name, designation=job.name,
			parent_node=unit.name, display_name="手动修改名称", assigned_employees=[emp.name], roster_auto_sync=0)
		config = chart._manual_node_config(frappe.get_doc("Organization Node", role.name).source_text)
		assert config["portable_id"] == "SOURCE-JOB" and config["chart_grade"]["code"] == "M2"
		reexport = package.read_package(package.export_configuration("永新")["file_url"])
		assert next(r for r in reexport["人员任职"] if r["node"] == "SOURCE-JOB")["code"] == code
		print("PASS: portable node IDs, parent IDs, leading-zero business code differing from Employee ID, two grade levels, manual edit preserves grade")

		def rejected(mutator, text):
			bad = copy.deepcopy(portable); mutator(bad)
			errors = package.prepare("永新", bad)["errors"]
			assert any(text in error for error in errors), errors
		rejected(lambda p: p["组织层级"][0].update(parent="SOURCE-JOB"), "循环")
		rejected(lambda p: p["组织层级"][0].update(parent="MISSING"), "上级不存在")
		rejected(lambda p: p["组织层级"].append(p["组织层级"][0].copy()), "重复")
		rejected(lambda p: p["组织层级"][0].update(department="不存在的课"), "未唯一匹配")
		rejected(lambda p: p["人员任职"][0].update(code="UNMATCHED"), "未唯一匹配")
		rejected(lambda p: p["组织层级"][0].update(chart_grade_code="MISSING"), "职级定义缺少")
		rejected(lambda p: p["职级定义"][0].update(parent="M2"), "循环")
		rejected(lambda p: p["职级定义"][1].update(rank="0"), "等级顺序")
		rejected(lambda p: p["组织层级"][0].update(roster_auto_sync="yes"), "必须为 0 或 1")
		# Stale preview must fail before any node writes.
		preview = package.preview_configuration("永新", url)
		frappe.db.set_value("Employee", emp.name, "employee_name", "修改后的姓名")
		try: package.apply_configuration("永新", url, preview["fingerprint"])
		except frappe.ValidationError as exc: assert "重新预览" in str(exc)
		else: raise AssertionError("stale preview accepted")
		# Duplicate business codes cannot be resolved by choosing the first row.
		frappe.db.set_value("Employee", before[0].name, "custom_employee_code", code)
		assert any("未唯一匹配" in e for e in package.prepare("永新", portable)["errors"])
		frappe.db.set_value("Employee", before[0].name, "custom_employee_code", before[0].custom_employee_code)
		# File reader rejects formulas and arbitrary URLs.
		book["组织层级"]["C2"] = "=1+1"
		try: package.read_package(save_book())
		except frappe.ValidationError as exc: assert "公式" in str(exc)
		else: raise AssertionError("formula accepted")
		try: package.read_package("https://example.com/config.xlsx")
		except frappe.ValidationError: pass
		else: raise AssertionError("external URL accepted")
		frappe.set_user("Guest")
		try: package.export_configuration("永新")
		except frappe.PermissionError: pass
		else: raise AssertionError("guest export accepted")
		finally: frappe.set_user("Administrator")
		print("PASS: missing/duplicate keys, cycles, invalid grades, unmatched/duplicate employee codes, stale preview, formula rejection, permissions")
		# Future employees in the roster Business department join the canonical subgroup.
		frappe.db.set_value("Employee", emp.name, {"department": "业务组", "employee_name": "配置迁移测试"})
		result = reconcile("永新")
		assert any(i.get("employee") == emp.name and "手动任职" in i["reason"] for i in result["issues"])
		chart.save_manual_organization_node("岗位", node_name=role.name, company="永新", department=dept.name, designation=job.name,
			parent_node=unit.name, assigned_employees=[], roster_auto_sync=0)
		reconcile("永新")
		nodes = chart._get_manual_organization_records("永新")["nodes"]
		business = next(n for n in nodes if n.manual_config.get("template_source_cell") == "CS18")
		bucket = next(n for n in nodes if n.manual_config.get("department") == "业务组" and n.manual_config.get("node_kind") == "岗位")
		assert bucket.parent_node == business.name and emp.name in bucket.manual_config["assigned_employees"]
		assert not any(n.display_name == "业务组" and not n.parent_node for n in nodes)
		print("PASS: business group merged under 总办室 and future roster allocation follows the explicit alias")
	finally:
		frappe.set_user("Administrator")
		frappe.db.rollback()
	assert frappe.db.count("Organization Node") == node_count
	assert frappe.get_all("Employee", fields=list(before[0]), order_by="name") == before
	print("PASS: all fixtures and imports rolled back")


if __name__ == "__main__":
	frappe.init(site="hrms.localhost", sites_path="/home/frappe/frappe-bench/sites")
	frappe.connect(); frappe.set_user("Administrator")
	try: run()
	finally: frappe.destroy()
