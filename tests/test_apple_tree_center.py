from __future__ import annotations

import importlib.util
import inspect
import sys
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path
from types import ModuleType
from types import SimpleNamespace


MODULE_PATH = Path(__file__).parents[1] / "hrms" / "hr" / "page" / "apple_tree_center" / "apple_tree_center.py"


def apple_tree_center_module():
	frappe = ModuleType("frappe")
	frappe.whitelist = lambda function=None, **_kwargs: (lambda decorated: decorated) if function is None else function
	frappe.throw = lambda message: (_ for _ in ()).throw(ValueError(message))
	frappe.get_doc = lambda *_args, **_kwargs: SimpleNamespace(check_permission=lambda *_permission: None, get=lambda field: "001" if field == "custom_employee_code" else None)
	frappe.get_all = lambda *_args, **_kwargs: []
	frappe_utils = ModuleType("frappe.utils")
	frappe_utils.getdate = lambda value: value if isinstance(value, date) else date.fromisoformat(str(value)[:10])
	frappe_utils.now_datetime = lambda: "2026-09-12 12:00:00"
	frappe_utils.nowdate = lambda: "2026-09-04"
	previous = {name: sys.modules.get(name) for name in ("frappe", "frappe.utils")}
	sys.modules.update({"frappe": frappe, "frappe.utils": frappe_utils})
	try:
		spec = importlib.util.spec_from_file_location("apple_tree_center_contract", MODULE_PATH)
		module = importlib.util.module_from_spec(spec)
		sys.modules[spec.name] = module
		spec.loader.exec_module(module)
		return module
	finally:
		for name, old_module in previous.items():
			if old_module is None:
				sys.modules.pop(name, None)
			else:
				sys.modules[name] = old_module


class AppleTreeCenterContractTest(unittest.TestCase):
	def test_download_template_matches_new_workbook_and_keeps_employee_code_as_text(self):
		center = apple_tree_center_module()
		workbook = center._build_history_import_template()
		sheet = workbook["苹果树合计"]
		self.assertEqual(sheet.max_column, 14)
		self.assertEqual(sheet.freeze_panes, "B4")
		self.assertEqual(sheet.auto_filter.ref, "B3:N203")
		self.assertEqual([sheet.cell(3, column).value for column in range(2, 15)], [
			"序号", "创建时间", "奖/惩日期", "受奖/惩人部门", "受奖/惩人工号", "受奖/惩人",
			"绿苹果", "红苹果", "奖/惩项目", "备注", "创建人", "签名", "备注",
		])
		self.assertEqual(sheet["F4"].number_format, "@")
		buffer = BytesIO()
		workbook.save(buffer)
		sheet_name, header_row, rows = center._history_sheet_rows(buffer.getvalue())
		self.assertEqual((sheet_name, header_row, rows), ("苹果树合计", 3, []))

	def test_export_workbook_uses_current_view_filters_and_sort_without_formula_cells(self):
		center = apple_tree_center_module()
		data = {"people": [
			{"department": "连续课", "employee_name": "=CMD()", "employee_code": "001", "green_apples": 3, "red_apples": 1, "net_apples": 2, "reward_amount": 10, "record_count": 2},
			{"department": "连续课", "employee_name": "李四", "employee_code": "002", "green_apples": 8, "red_apples": 1, "net_apples": 7, "reward_amount": 35, "record_count": 3},
			{"department": "品管课", "employee_name": "王五", "employee_code": "003", "green_apples": 99, "red_apples": 0, "net_apples": 99, "reward_amount": 495, "record_count": 1},
		]}
		workbook = center._build_export_workbook(
			"annual-summary", data, "2026年 苹果树个人汇总", '{"department":"连续课"}', "net_apples", "desc",
		)
		sheet = workbook["苹果树统计"]
		self.assertEqual(sheet["A1"].value, "2026年 苹果树个人汇总")
		self.assertEqual(sheet.freeze_panes, "A3")
		self.assertEqual(sheet.auto_filter.ref, "A2:H4")
		self.assertEqual([sheet.cell(2, column).value for column in range(1, 9)], ["部门", "姓名", "工号", "绿苹果", "红苹果", "净苹果", "苹果金额", "记录数"])
		self.assertEqual(sheet["C3"].value, "002")
		self.assertEqual(sheet["B4"].value, "'=CMD()")

	def test_monthly_export_derives_net_apples_and_final_status(self):
		center = apple_tree_center_module()
		columns, rows = center._export_rows("monthly-detail", {"records": [{
			"reward_date": "2026-08-12", "employee_code": "001", "green_apples": 4, "red_apples": 2,
			"approval_result": "已确认", "approval_status": "已锁定",
		}]})
		self.assertEqual(columns[0], ("attendance_month", "月份"))
		self.assertEqual(rows[0]["attendance_month"], "2026-08")
		self.assertEqual(rows[0]["net_apples"], 2)
		self.assertEqual(rows[0]["final_status"], "已确认 / 已锁定")

	def test_person_export_uses_apple_reward_headers_and_total_at_the_bottom(self):
		center = apple_tree_center_module()
		columns, rows = center._export_rows("person", {
			"year": "2026", "person": {"department": "连续课", "employee_name": "张三", "employee_code": "001"},
			"rows": [{"sequence": 1, "created_at": "2026-08-01 09:00:00", "reward_date": "2026-08-01", "employee_code": "001", "green_apples": 4, "red_apples": 1}],
			"totals": {"green_apples": 4, "red_apples": 1},
		})
		self.assertEqual(len(columns), 12)
		self.assertEqual([label for _field, label in columns], [
			"序号", "创建时间", "奖/惩日期", "受奖/惩人部门", "受奖/惩人",
			"绿苹果", "红苹果", "奖/惩项目", "备注", "创建人", "签名", "备注",
		])
		self.assertEqual(len(rows), 2)
		self.assertEqual(rows[0]["sequence"], 1)
		self.assertEqual(rows[-1]["sequence"], "合计")
		self.assertEqual(rows[-1]["green_apples"], 4)
		self.assertEqual(rows[-1]["red_apples"], 1)
		filtered = center._filtered_export_rows(columns, rows, sort_key="green_apples", sort_order="desc")
		self.assertEqual(filtered[-1]["sequence"], "合计")

	def test_person_detail_maps_monthly_workbook_fields_without_attendance_columns(self):
		center = apple_tree_center_module()
		detail = center._apple_detail_from_source({
			"sequence": 9, "created_at": "2026-06-02 20:00:23", "reward_date": "2026-06-02",
			"source_department": "连续课", "employee_code": "388", "employee_name": "孔红西",
			"green_apples": 5, "red_apples": 0, "reward_item": "连续课/绿苹果/保养",
			"note": "保养Y线", "created_by_name": "李鑫", "signature": "", "note_2": "",
		})
		self.assertEqual(detail["sequence"], 9)
		self.assertEqual(detail["created_at"], "2026-06-02 20:00:23")
		self.assertEqual(detail["department"], "连续课")
		self.assertEqual(detail["employee_code"], "388")
		self.assertEqual(detail["green_apples"], 5)
		self.assertNotIn("actual_attendance_hours", detail)

	def test_history_employee_prefers_code_and_rejects_name_mismatch(self):
		center = apple_tree_center_module()
		filters = []
		center.frappe.get_all = lambda _doctype, **kwargs: filters.append(kwargs["filters"]) or [{
			"name": "EMP-001", "employee_name": "张三", "custom_employee_code": "001", "department": "制造",
		}]
		employee, error = center._history_employee("永新", "001", "张三")
		self.assertEqual(employee["name"], "EMP-001")
		self.assertEqual(error, "")
		self.assertEqual(filters, [{"company": "永新", "custom_employee_code": "001"}])
		self.assertIn("工号与姓名不一致", center._history_employee("永新", "001", "李四")[1])

	def test_history_workbook_matches_supplied_third_row_header_format(self):
		from openpyxl import Workbook

		center = apple_tree_center_module()
		workbook = Workbook()
		sheet = workbook.active
		sheet.title = "苹果树合计"
		sheet.append([])
		sheet.append([None, "6月苹果树"])
		sheet.append([None, "序号", "创建时间", "奖/惩日期", "受奖/惩人部门", "受奖/惩人", "绿苹果", "红苹果", "奖/惩项目", "备注", "创建人", "签名", "备注"])
		sheet.append([None, 1, "2026-06-08 16:33:11", "2026-06-08", "工程课", "张三", "", 2, "红苹果项目", "处理说明", "管理员"])
		buffer = BytesIO()
		workbook.save(buffer)
		sheet_name, header_row, rows = center._history_sheet_rows(buffer.getvalue())
		self.assertEqual((sheet_name, header_row, len(rows)), ("苹果树合计", 3, 1))
		self.assertEqual(rows[0]["__source_row"], 4)
		self.assertEqual(rows[0]["受奖/惩人"], "张三")
		self.assertEqual(rows[0]["备注"], "处理说明")
		self.assertIn("备注_2", rows[0])

	def test_employee_summary_uses_employee_identity_and_the_same_active_final_scope(self):
		center = apple_tree_center_module()
		center.get_data = lambda **kwargs: {
			"filters": {"company": kwargs["company"], "year": 2026},
			"people": [
				{"employee": "EMP-001", "employee_code": "001", "employee_name": "同名", "green_apples": 3, "red_apples": 1, "net_apples": 2, "record_count": 2, "reward_amount": 10},
				{"employee": "EMP-002", "employee_code": "002", "employee_name": "同名", "green_apples": 99},
			],
		}
		center.frappe.get_doc = lambda _doctype, name: SimpleNamespace(
			check_permission=lambda *_permission: None,
			get=lambda field: "001" if field == "custom_employee_code" and name == "EMP-001" else "404" if field == "custom_employee_code" else None,
		)
		result = center.get_employee_summary("EMP-001", "2026", "永新")
		self.assertTrue(result["available"])
		self.assertEqual(result["person_key"], "001")
		self.assertEqual(result["person"]["green_apples"], 3)
		self.assertFalse(center.get_employee_summary("missing", "2026", "永新")["available"])

	def test_person_detail_accepts_employee_code_but_keeps_rows_on_that_code(self):
		center = apple_tree_center_module()
		records = [
			{"employee": "EMP-001", "employee_code": "001", "employee_name": "同名", "attendance_month": "2026-08", "green_apples": 3},
			{"employee": "EMP-002", "employee_code": "002", "employee_name": "同名", "attendance_month": "2026-08", "green_apples": 99},
		]
		center.get_data = lambda **kwargs: {"filters": {"company": kwargs["company"], "year": kwargs["year"]}, "people": records, "records": records}
		result = center.get_person_detail("001", "2026", "永新")
		self.assertTrue(result["available"])
		self.assertEqual([row["employee_code"] for row in result["rows"]], ["001"])

	def test_person_detail_forwards_outer_filters_and_applies_detail_filter(self):
		center = apple_tree_center_module()
		calls = []
		records = [{
			"employee": "EMP-001", "employee_code": "001", "employee_name": "张三", "department": "制造",
			"attendance_month": "2026-06", "green_apples": 3, "red_apples": 0,
			"reward_item": "月度考勤终稿",
		}]
		def fake_get_data(**kwargs):
			calls.append(kwargs)
			return {
				"filters": {"company": "永新", "year": "2026", "month": "2026-06", "search": "保养", "start_date": "2026-06-10", "end_date": "2026-06-18"},
				"available_years": [2026], "people": records, "records": records,
			}
		center.get_data = fake_get_data
		result = center.get_person_detail(
			"001", "2026", "永新", month="2026-06", search="保养",
			start_date="2026-06-10", end_date="2026-06-18", detail_search="终稿",
		)
		self.assertEqual(calls[0], {
			"year": "2026", "month": "2026-06", "search": "保养", "company": "永新",
			"start_date": "2026-06-10", "end_date": "2026-06-18",
		})
		self.assertEqual(len(result["rows"]), 1)
		self.assertEqual(result["filters"]["month"], "2026-06")
		self.assertEqual(result["filters"]["detail_search"], "终稿")

	def test_person_export_forwards_both_outer_and_detail_filters(self):
		center = apple_tree_center_module()
		calls = []
		center.frappe.local = SimpleNamespace(response=SimpleNamespace())
		center.get_person_detail = lambda **kwargs: calls.append(kwargs) or {
			"available": True, "year": "2026", "person": {"employee_name": "张三"}, "rows": [], "totals": {},
		}
		class ExportWorkbook:
			def save(self, output):
				output.write(b"xlsx")
		center._build_export_workbook = lambda *_args: ExportWorkbook()
		center.download_export(
			view="person", year="2026", month="2026-06", search="张三", company="永新",
			start_date="2026-06-01", end_date="2026-06-30", person="001",
			detail_start_date="2026-06-10", detail_end_date="2026-06-18", detail_search="保养Y",
		)
		self.assertEqual(calls[0], {
			"person": "001", "year": "2026", "company": "永新", "month": "2026-06", "search": "张三",
			"start_date": "2026-06-01", "end_date": "2026-06-30",
			"detail_start_date": "2026-06-10", "detail_end_date": "2026-06-18", "detail_search": "保养Y",
		})

	def test_quarters_select_only_the_three_months_in_the_chosen_year(self):
		center = apple_tree_center_module()
		center.frappe.get_list = lambda *_args, **_kwargs: [{"attendance_month": f"2026-{month:02d}"} for month in range(1, 13)] + [{"attendance_month": "2025-12"}]
		calls = []
		def active(company, attendance_month):
			calls.append(attendance_month)
			return [{"attendance_month": attendance_month, "employee": "001", "employee_code": "001", "green_apples": 2, "red_apples": 1}]
		center._list_active_month_records = active
		for quarter in range(1, 5):
			calls.clear()
			result = center.get_data(year="2026", month=f"2026-Q{quarter}", company="永新")
			self.assertEqual(calls, [f"2026-{month:02d}" for month in range(quarter * 3 - 2, quarter * 3 + 1)])
			self.assertEqual(result["summary"]["green_apples"], 6)
			self.assertEqual(result["summary"]["employee_count"], 1)
		for invalid in ("2026-Q0", "2026-Q5", "2025-Q4"):
			with self.assertRaises(ValueError):
				center.get_data(year="2026", month=invalid, company="永新")

	def test_person_detail_keeps_identity_version_and_unknown_values(self):
		center = apple_tree_center_module()
		records = [
			{"employee": "001", "employee_code": "001", "employee_name": "同名", "attendance_month": "2026-08", "standard_hours": 160, "actual_attendance_hours": 152, "green_apples": 3, "red_apples": 0, "approval_no": "处理终稿:old"},
			{"employee": "002", "employee_code": "002", "employee_name": "同名", "attendance_month": "2026-07", "green_apples": 999},
			{"employee": "001", "employee_code": "001", "employee_name": "同名", "attendance_month": "2026-07", "standard_hours": 184, "actual_attendance_hours": 180, "green_apples": 5, "red_apples": 2, "approval_no": "处理终稿:current"},
		]
		center.get_data = lambda **kwargs: {"filters": {"company": kwargs["company"], "year": kwargs["year"]}, "people": records[:2], "records": records}
		result = center.get_person_detail("001", "2026", "永新")
		self.assertEqual([row["green_apples"] for row in result["rows"]], [5, 3])
		self.assertEqual(result["totals"]["green_apples"], 8)
		self.assertEqual(result["totals"]["red_apples"], 2)
		self.assertEqual(len(result["columns"]), 12)
		self.assertNotIn("actual_attendance_hours", result["rows"][0])
		self.assertFalse(center.get_person_detail("missing", "2026", "永新")["available"])

	def test_rpc_arguments_have_string_annotations_for_frappe_type_validation(self):
		center = apple_tree_center_module()
		parameters = inspect.signature(center.get_data).parameters
		self.assertEqual({name: parameter.annotation for name, parameter in parameters.items()}, {
			"year": "str",
			"month": "str",
			"search": "str",
			"company": "str",
			"start_date": "str",
			"end_date": "str",
		})

	def test_custom_date_range_validation(self):
		center = apple_tree_center_module()
		self.assertEqual(center._parse_custom_date_range("2025-06-01", "2026-08-31"), (date(2025, 6, 1), date(2026, 8, 31)))
		for values in (("2026-01-01", ""), ("2026-02-01", "2026-01-31"), ("2026/01/01", "2026-01-31")):
			with self.assertRaises(ValueError):
				center._parse_custom_date_range(*values)

	def test_history_date_range_filters_source_rows_then_reaggregates_employee_month(self):
		center = apple_tree_center_module()
		center._active_history_batch = lambda *_args: {"name": "BATCH-1"}
		center.frappe.get_all = lambda *_args, **_kwargs: [{
			"name": "SUM-1", "employee": "EMP-001", "employee_code": "001", "employee_name": "张三",
			"department": "制造", "green_apples": 9, "red_apples": 4, "reward_amount": 25, "source_row_count": 3,
			"source_rows_json": '[{"reward_date":"2026-06-01","green_apples":2,"red_apples":0},{"reward_date":"2026-06-15","green_apples":3,"red_apples":1},{"reward_date":"2026-06-30","green_apples":4,"red_apples":3}]',
		}]
		rows = center._list_history_month_records("永新", "2026-06", date(2026, 6, 10), date(2026, 6, 20))
		self.assertEqual(len(rows), 1)
		self.assertEqual((rows[0]["green_apples"], rows[0]["red_apples"], rows[0]["reward_amount"], rows[0]["source_row_count"]), (3, 1, 10, 1))

	def test_cross_year_date_range_uses_exact_history_and_only_fully_covered_monthly_finals(self):
		center = apple_tree_center_module()
		center.frappe.get_list = lambda *_args, **_kwargs: [{"attendance_month": "2025-12"}, {"attendance_month": "2026-01"}]
		center._available_history_months = lambda _company: {"2025-12"}
		history_calls = []
		center._list_history_month_records = lambda company, month, start, end: history_calls.append((company, month, start, end)) or [{
			"attendance_month": month, "reward_date": f"{month}-01", "employee_code": "001", "employee_name": "张三", "green_apples": 2, "red_apples": 0,
		}]
		active_calls = []
		center._list_active_month_records = lambda company, month: active_calls.append((company, month)) or [{"attendance_month": month, "employee_code": "002", "employee_name": "李四", "green_apples": 3, "red_apples": 0}]
		partial = center.get_data(year="2026", company="永新", start_date="2025-12-15", end_date="2026-01-31")
		self.assertEqual(history_calls[0][1:], ("2025-12", date(2025, 12, 15), date(2026, 1, 31)))
		self.assertEqual(active_calls, [("永新", "2026-01")])
		self.assertEqual(partial["summary"]["green_apples"], 5)
		self.assertEqual(partial["filters"]["start_date"], "2025-12-15")
		self.assertIn("按奖/惩日期精确统计", partial["notice"])

	def test_employee_and_month_statistics_keep_green_red_and_net_separate(self):
		center = apple_tree_center_module()
		summary, people, months = center._summarize_records(
			[
				{"reward_date": "2026-06-01", "employee": "EMP-001", "employee_code": "001", "employee_name": "张三", "department": "制造", "green_apples": 3, "red_apples": 1, "reward_amount": 10},
				{"reward_date": "2026-06-12", "employee": "EMP-001", "employee_code": "001", "employee_name": "张三", "department": "制造", "green_apples": 2, "red_apples": 0, "reward_amount": 10},
				{"reward_date": "2026-07-03", "employee": "EMP-002", "employee_code": "002", "employee_name": "李四", "department": "品管", "green_apples": 1, "red_apples": 4, "reward_amount": -15},
			]
		)

		self.assertEqual(summary, {"record_count": 3, "green_apples": 6, "red_apples": 5, "reward_amount": 5, "net_apples": 1, "employee_count": 2})
		self.assertEqual(people[0]["employee_name"], "张三")
		self.assertEqual(people[0]["net_apples"], 4)
		self.assertEqual(months[0], {"month": "2026-07", "record_count": 1, "green_apples": 1, "red_apples": 4, "reward_amount": -15, "net_apples": -3})

	def test_month_must_belong_to_selected_year(self):
		center = apple_tree_center_module()
		self.assertEqual(center._parse_month("2026-09", 2026), "2026-09")
		self.assertEqual(center._date_range(2026, "2026-09"), (date(2026, 9, 1), date(2026, 10, 1)))
		with self.assertRaises(ValueError):
			center._parse_month("2025-12", 2026)

	def test_data_uses_the_current_company_active_monthly_final_for_the_selected_month(self):
		center = apple_tree_center_module()
		calls = []

		def get_list(doctype, **kwargs):
			calls.append((doctype, kwargs))
			self.assertEqual(doctype, center.MONTHLY_SUMMARY_DOCTYPE)
			self.assertEqual(kwargs["filters"], {"company": "永新"})
			self.assertEqual(kwargs["fields"], ["attendance_month"])
			return [{"attendance_month": "2026-09"}, {"attendance_month": "2025-12"}]

		center.frappe.get_list = get_list
		active_calls = []
		center._list_active_month_records = lambda company, attendance_month: active_calls.append((company, attendance_month)) or [{
			"name": "FINAL-001",
			"attendance_month": attendance_month,
			"employee": "EMP-001",
			"employee_code": "001",
			"employee_name": "张三",
			"department": "制造",
			"green_apples": 2,
			"red_apples": 0,
			"apple_reward_amount": 10,
			"attendance_lock_version": "处理终稿:abc",
			"lock_status": "已锁定",
			"status": "已确认",
		}]
		result = center.get_data(year="2026", month="2026-09", company="永新")
		self.assertEqual(len(calls), 1)
		self.assertEqual(active_calls, [("永新", "2026-09")])
		self.assertEqual(result["summary"]["green_apples"], 2)
		self.assertEqual(result["summary"]["employee_count"], 1)
		self.assertEqual(result["available_years"], [2026, 2025])
		self.assertEqual(result["records"][0]["reward_item"], "月度考勤终稿")
		self.assertEqual(result["records"][0]["approval_status"], "已锁定")

	def test_statistics_only_history_takes_precedence_for_its_month(self):
		center = apple_tree_center_module()
		center.frappe.get_list = lambda *_args, **_kwargs: [{"attendance_month": "2026-06"}]
		center._available_history_months = lambda _company: {"2026-06"}
		center._list_history_month_records = lambda _company, _month: [{
			"attendance_month": "2026-06", "reward_date": "2026-06-01",
			"employee": "EMP-001", "employee_code": "001", "employee_name": "张三",
			"green_apples": 8, "red_apples": 2, "reward_amount": 30,
		}]
		center._list_active_month_records = lambda *_args: self.fail("历史导入月份不应再读取考勤终稿")
		result = center.get_data(year="2026", month="2026-06", company="永新")
		self.assertEqual(result["summary"]["net_apples"], 6)
		self.assertIn("不修改考勤终稿或薪资", result["notice"])
