"""Run production controllers and selectors against memory-only documents and stores."""
import ast
from collections import defaultdict
from copy import deepcopy
from datetime import date, datetime
from io import BytesIO
import importlib.util
import inspect
import json
from pathlib import Path
import sys
from types import ModuleType, SimpleNamespace
import unittest
from unittest.mock import Mock, patch

ROOT = Path(__file__).resolve().parents[1]
class Row(dict):
    __getattr__ = dict.get
    __setattr__ = dict.__setitem__

class Document(Row):
    def get_doc_before_save(self): return self.get('_old')


def load_module(path, name):
    spec = importlib.util.spec_from_file_location(name, ROOT / path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

class StandingPayTests(unittest.TestCase):
    def setUp(self):
        self.frappe = ModuleType('frappe')
        self.frappe._ = lambda value: value
        self.frappe.whitelist = lambda: lambda fn: fn
        self.frappe.session = Row(user='editor')
        self.frappe.PermissionError = ValueError
        self.frappe.get_roles = lambda user: ['HR Manager']
        self.frappe.throw = lambda value, *args: (_ for _ in ()).throw(ValueError(value))
        self.frappe.get_doc = Mock(return_value=SimpleNamespace(check_permission=lambda *a: None))
        self.frappe.db = SimpleNamespace(get_value=Mock(return_value='ACME'), exists=Mock(return_value=False), sql=Mock())
        self.frappe.get_all = Mock(return_value=[])
        utils = ModuleType('frappe.utils')
        utils.flt = lambda value: float(value or 0)
        utils.getdate = lambda value: date.fromisoformat(str(value))
        utils.get_datetime = lambda value: datetime.fromisoformat(str(value))
        utils.nowdate = lambda: '2026-09-09'
        utils.now_datetime = lambda: datetime(2026, 9, 9, 12, 34, 56)
        doc_module = ModuleType('frappe.model.document'); doc_module.Document = Document
        payroll = ModuleType('hrms.api.payroll_input')
        payroll._require_company = lambda company: company
        payroll._require_payroll_master_manager = lambda: None
        self.patcher = patch.dict(sys.modules, {'frappe': self.frappe, 'frappe.utils': utils,
            'frappe.model.document': doc_module, 'hrms.api.payroll_input': payroll})
        self.patcher.start(); self.addCleanup(self.patcher.stop)
        self.permissions = load_module('hrms/payroll/standing_permissions.py', 'standing_permissions')
        sys.modules['hrms.payroll.standing_permissions'] = self.permissions
        self.controller = load_module('hrms/payroll/standing_pay.py', 'standing_controller')
        self.api = load_module('hrms/api/standing_pay.py', 'standing_api')

    def document(self, kind='salary', **values):
        dirname = 'hrms_employee_salary_change' if kind == 'salary' else 'hrms_employee_contribution_change'
        schema = json.loads((ROOT / f'hrms/hr/doctype/{dirname}/{dirname}.json').read_text())
        doc = self.controller.StandingPayDecision(doctype=schema['name'], company='ACME', employee='e1',
            effective_date='2026-07-01', base_salary=3000., function_allowance=200., full_salary=3200.,
            certificate_allowance=0., multi_skill_allowance=0., remarks='业务变更', contribution_type='社保',
            personal_amount=100., company_amount=200., enabled=1, owner='editor', **values)
        doc.meta = SimpleNamespace(fields=[Row(field) for field in schema['fields']])
        return doc

    def pending(self, kind='salary'):
        doc = self.document(kind); doc.validate(); old = Row(deepcopy(dict(doc))); doc._old = old
        return doc

    def test_new_decision_never_accepts_forged_approval(self):
        for kind in ('salary','contribution'):
            doc=self.document(kind, status='已批准', approved_by='forged', submitted_by='forged')
            doc.validate()
            self.assertEqual((doc.status,doc.submitted_by,doc.approved_by),('待审核','editor',None))
            self.assertEqual(str(doc.submitted_on),'2026-09-09 12:34:56')

    def test_second_pending_decision_for_the_same_slot_is_denied(self):
        for kind in ('salary', 'contribution'):
            doc = self.document(kind)
            self.frappe.db.exists.return_value = True
            with self.assertRaisesRegex(ValueError, '已有待审批申请'):
                doc.validate()
            filters = self.frappe.db.exists.call_args.args[1]
            self.assertEqual(filters['effective_date'], '2026-07-01')
            if kind == 'contribution':
                self.assertEqual(filters['contribution_type'], '社保')
            else:
                self.assertEqual(filters['exclude_from_payroll'], 0)
            self.frappe.db.exists.reset_mock(return_value=True)

    def test_salary_api_reuses_exact_retry_but_denies_changed_pending_request(self):
        tree = ast.parse((ROOT / 'hrms/api/payroll_input.py').read_text())
        fn = next(node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name == 'create_employee_salary_change')
        fn.decorator_list = []
        existing = 'pending-salary-request'
        def get_value(doctype, filters, *args, **kwargs):
            if doctype == 'Employee':
                return Row(company='ACME', employee_name='员工甲', custom_employee_code='E1', department='HR')
            # The exact-retry lookup includes the submitted amount.  The
            # single-pending-slot lookup intentionally does not.
            return existing if filters.get('base_salary') == 3000 or 'base_salary' not in filters else None
        self.frappe.db.get_value.side_effect = get_value
        namespace = dict(
            frappe=self.frappe, _=lambda value: value, EMPLOYEE_SALARY_CHANGE_DOCTYPE='salary',
            require_access=lambda company: company, _employee_context=lambda employee: Row(company='ACME', employee_name='员工甲', custom_employee_code='E1', department='HR'),
            _grade_context=lambda grade: {}, _salary_contribution_defaults=lambda context, effective: {'housing_fund_enabled': 0, 'social_insurance_enabled': 0},
            _employee_code=lambda context: context.custom_employee_code, flt=lambda value: float(value or 0),
            _record_payroll_manual_adjustment=lambda **kwargs: None,
        )
        exec(compile(ast.Module(body=[fn], type_ignores=[]), 'payroll_input.py', 'exec'), namespace)
        self.assertEqual(namespace['create_employee_salary_change'](
            company='ACME', employee='e1', effective_date='2026-07-01', base_salary=3000, remarks='调整金额'
        ), existing)
        with self.assertRaisesRegex(ValueError, '已有待审批定薪申请'):
            namespace['create_employee_salary_change'](
                company='ACME', employee='e1', effective_date='2026-07-01', base_salary=3300, remarks='调整金额'
            )
        self.frappe.db.sql.assert_called_once()

    def test_operator_cannot_approve_api_or_direct_document(self):
        self.frappe.get_roles = lambda user: ['薪资经办']
        doc = self.pending()
        self.assertEqual(doc.status, '待审核')
        self.frappe.session.user = 'other-editor'
        doc.status = '已批准'; doc.review_note = 'forged'
        with self.assertRaisesRegex(ValueError, '没有薪资审批权限'):
            doc.validate()
        with self.assertRaisesRegex(ValueError, '没有薪资审批权限'):
            self.api.list_pending('ACME')
        with self.assertRaisesRegex(ValueError, '没有薪资审批权限'):
            self.api.review_decision('ACME', 'HRMS Employee Salary Change', 'missing', '已批准', 'forged')

    def test_self_approval_denied(self):
        for role in ('HR Manager', 'System Manager', '薪资审批'):
            self.frappe.get_roles = lambda user: [role]
            for kind in ('salary', 'contribution'):
                doc=self.pending(kind); doc.status='已批准'; doc.review_note='agree'
                with self.assertRaisesRegex(ValueError,'不能审批自己的'): doc.validate()

    def test_administrator_can_review_own_requests_with_audit(self):
        self.frappe.session.user = 'Administrator'
        for kind in ('salary', 'contribution'):
            for status in ('已批准', '已驳回'):
                doc = self.pending(kind); doc.status = status; doc.review_note = 'checked'
                doc.approved_by = 'forged'
                doc.validate()
                self.assertEqual(doc.submitted_by, 'Administrator')
                self.assertEqual(doc.approved_by, 'Administrator')
                self.assertIsNotNone(doc.approved_on)

    def test_administrator_self_review_preserves_other_guards(self):
        self.frappe.session.user = 'Administrator'
        doc = self.pending(); doc.status = '已批准'
        with self.assertRaisesRegex(ValueError, '审批意见'): doc.validate()
        doc.review_note = 'checked'; self.frappe.db.exists.return_value = True
        with self.assertRaisesRegex(ValueError, '已确认薪资'): doc.validate()
        self.frappe.db.exists.return_value = False; doc.base_salary = 9999
        with self.assertRaisesRegex(ValueError, '不可覆盖'): doc.validate()

    def test_distinct_reviewer_stamped_and_rejection_supported(self):
        for kind in ('salary','contribution'):
            for status in ('已批准','已驳回'):
                self.frappe.session.user='editor';doc=self.pending(kind)
                self.frappe.session.user='reviewer';doc.status=status;doc.review_note='checked';doc.approved_by='forged'
                doc.validate();self.assertEqual(doc.approved_by,'reviewer');self.assertIsNotNone(doc.approved_on)

    def test_editing_payload_during_approval_is_denied(self):
        for field,value in [('base_salary',9999),('effective_date','2026-08-01'),('employee','e2'),('submitted_by','forged')]:
            self.frappe.session.user='editor';doc=self.pending();self.frappe.session.user='reviewer'
            doc[field]=value;doc.status='已批准';doc.review_note='agree'
            with self.assertRaisesRegex(ValueError,'不可覆盖'):doc.validate()

    def test_approved_decision_cannot_be_reapproved_or_deleted(self):
        doc=self.pending();doc._old.status='已批准';doc.status='已批准';doc.review_note='agree'
        with self.assertRaises(ValueError):doc.validate()
        with self.assertRaises(ValueError):doc.on_trash()

    def test_closed_payroll_blocks_backdated_approval(self):
        doc=self.pending();doc.status='已批准';doc.review_note='agree';self.frappe.session.user='reviewer'
        self.frappe.db.exists.return_value=True
        with self.assertRaisesRegex(ValueError,'已确认薪资'):doc.validate()

    def test_stop_is_zero_and_nonfinite_amounts_rejected(self):
        doc=self.document('contribution');doc.enabled=0;doc.validate()
        self.assertEqual((doc.personal_amount,doc.company_amount),(0,0))
        for value in (-1,float('nan'),float('inf')):
            doc=self.document();doc.base_salary=value
            with self.assertRaises(ValueError):doc.validate()

    def selector_store(self):
        self.rows=[]
        def get_all(doctype, filters, **kwargs):
            rows=[]
            for row in self.rows:
                if row.company != filters['company'] or row.status != filters['status']: continue
                if row.effective_date > str(filters['effective_date'][1]): continue
                rows.append(row)
            rows.sort(key=lambda row:(row.effective_date,row.approved_on,row.creation,row.name),reverse=True)
            return rows[:kwargs.get('limit_page_length',100000)]
        self.frappe.get_all=get_all
        def add(name,status='已批准',effective='2026-07-01',approved='2026-07-01',kind='社保',company='ACME',enabled=1):
            self.rows.append(Row(name=name,employee='e1',employee_code='E1',company=company,status=status,
                effective_date=effective,approved_on=approved,creation=approved,contribution_type=kind,
                personal_amount=100,company_amount=200,enabled=enabled))
        return add

    def test_contribution_carries_forward_ignores_pending_rejected_and_other_company(self):
        add=self.selector_store();add('approved');add('pending','待审核','2026-08-01');add('rejected','已驳回','2026-09-01');add('foreign',company='OTHER')
        self.assertEqual([r.name for r in self.api.active_contributions('ACME','2026-09-30')],['approved'])

    def test_same_day_revision_and_future_stop(self):
        add=self.selector_store();add('old');add('same-day',approved='2026-08-01');add('stop',effective='2026-10-01',enabled=0)
        self.assertEqual(self.api.active_contributions('ACME','2026-09-30')[0].name,'same-day')
        self.assertEqual(self.api.active_contributions('ACME','2026-10-31')[0].name,'stop')

    def test_pending_list_includes_only_the_current_pre_change_standard(self):
        pending_salary = Row(name='salary-request', company='ACME', employee='e1', status='待审核', base_salary=3300)
        pending_social = Row(name='social-request', company='ACME', employee='e1', status='待审核', contribution_type='社保', personal_amount=0, company_amount=1256.82, enabled=1)
        approved = [
            Row(name='salary-current', company='ACME', employee='e1', status='已批准', effective_date='2026-09-01', approved_on='2026-09-01', creation='2026-09-01', base_salary=3170),
            Row(name='salary-future', company='ACME', employee='e1', status='已批准', effective_date='2026-10-01', approved_on='2026-09-01', creation='2026-09-01', base_salary=9999),
            Row(name='social-current', company='ACME', employee='e1', status='已批准', effective_date='2026-08-01', approved_on='2026-08-01', creation='2026-08-01', contribution_type='社保', personal_amount=791.45, company_amount=1900.96, enabled=1),
            Row(name='housing-current', company='ACME', employee='e1', status='已批准', effective_date='2026-08-01', approved_on='2026-08-01', creation='2026-08-01', contribution_type='公积金', personal_amount=248, company_amount=248, enabled=1),
        ]
        def get_all(doctype, filters, **kwargs):
            if filters['status'] == '待审核':
                return [pending_salary] if doctype == self.api.SALARY else [pending_social]
            cutoff = str(filters['effective_date'][1])
            rows = [row for row in approved if row.effective_date <= cutoff]
            return sorted(rows, key=lambda row: (row.effective_date, row.approved_on, row.creation, row.name), reverse=True)
        self.frappe.get_all = get_all
        rows = {row.name: row for row in self.api.list_pending('ACME')}
        self.assertEqual(rows['salary-request'].previous_standard.name, 'salary-current')
        self.assertEqual(rows['social-request'].previous_standard.name, 'social-current')
        self.assertNotEqual(rows['salary-request'].previous_standard.name, 'salary-future')

    def test_change_records_exclude_initial_submissions(self):
        row = lambda name, employee, status, stamp, kind='': Row(
            name=name, employee=employee, status=status, submitted_on=stamp,
            creation=stamp, contribution_type=kind)
        salary_rows = [
            row('e1-change-rejected', 'e1', '已驳回', '2026-09-04'),
            row('e1-change-pending', 'e1', '待审核', '2026-09-03'),
            row('e1-initial-approved', 'e1', '已批准', '2026-09-02'),
            row('e1-initial-rejected', 'e1', '已驳回', '2026-09-01'),
            row('e2-initial-pending', 'e2', '待审核', '2026-09-01'),
        ]
        self.assertEqual(
            [item.name for item in self.api._only_change_records(salary_rows, self.api.SALARY)],
            ['e1-change-rejected', 'e1-change-pending'],
        )
        contribution_rows = [
            row('social-change', 'e1', '已批准', '2026-09-03', '社保'),
            row('housing-initial', 'e1', '已批准', '2026-09-02', '公积金'),
            row('social-initial', 'e1', '已批准', '2026-09-01', '社保'),
        ]
        self.assertEqual(
            [item.name for item in self.api._only_change_records(contribution_rows, self.api.CONTRIBUTION)],
            ['social-change'],
        )

    def test_latest_register_action_keeps_one_most_recent_submitter_per_employee(self):
        rows = [
            Row(name='e1-initial', employee='e1', submitted_on='2026-09-01 08:00:00'),
            Row(name='e1-change', employee='e1', submitted_on='2026-09-03 08:00:00'),
            Row(name='e2-initial', employee='e2', submitted_on='2026-09-02 08:00:00'),
        ]
        latest = {row.employee: row.name for row in self.api._latest_action_rows(rows)}
        self.assertEqual(latest, {'e1': 'e1-change', 'e2': 'e2-initial'})

    def test_actor_names_use_full_name_and_preserve_legacy_fallback(self):
        self.frappe.get_all = Mock(return_value=[Row(name='editor@example.com', full_name='张三')])
        rows = [
            Row(submitted_by='editor@example.com', owner='owner@example.com', approved_by='reviewer@example.com'),
            Row(owner='legacy@example.com'),
        ]
        self.api._attach_actor_names(rows)
        self.assertEqual(rows[0].submitted_by_name, '张三')
        self.assertEqual(rows[0].approved_by_name, 'reviewer@example.com')
        self.assertEqual(rows[1].submitted_by_name, 'legacy@example.com')

    def test_change_history_uses_standard_effective_when_request_was_submitted(self):
        baseline = Row(name='baseline', employee='e1', status='已批准', submitted_on='2026-09-01 08:00:00',
            approved_on='2026-09-01 09:00:00', creation='2026-09-01 08:00:00', effective_date='2026-09-01', base_salary=3370)
        first_change = Row(name='first-change', employee='e1', status='已批准', submitted_on='2026-09-10 09:26:24',
            approved_on='2026-09-10 10:03:17', creation='2026-09-10 09:26:24', effective_date='2026-09-10', base_salary=3333)
        parallel_change = Row(name='parallel-change', employee='e1', status='已批准', submitted_on='2026-09-10 09:48:48',
            approved_on='2026-09-10 10:03:14', creation='2026-09-10 09:48:48', effective_date='2026-09-10', base_salary=0)
        later_change = Row(name='later-change', employee='e1', status='待审核', submitted_on='2026-09-10 10:04:00',
            creation='2026-09-10 10:04:00', effective_date='2026-09-10', base_salary=3500)
        rows = [later_change, parallel_change, first_change, baseline]
        self.api._attach_previous_standards(rows, self.api.SALARY)
        self.assertEqual(first_change.previous_standard.name, 'baseline')
        self.assertEqual(parallel_change.previous_standard.name, 'baseline')
        self.assertEqual(later_change.previous_standard.name, 'first-change')
        self.assertIsNone(baseline.previous_standard)

    def test_compensation_export_keeps_opening_snapshot_and_equal_valued_changes(self):
        employee = Row(name='e1', employee_name='员工甲', custom_employee_code='E1',
            department_label='工程课', work_nature='在职·正式')
        salary = lambda name, effective, amount: Row(
            name=name, employee='e1', effective_date=effective, approved_on=effective,
            creation=effective, base_salary=amount, function_allowance=200,
            certificate_allowance=0, multi_skill_allowance=0, full_salary=amount + 200,
            remarks=name)
        contribution = Row(name='social-change', employee='e1', contribution_type='社保',
            effective_date='2026-03-01', approved_on='2026-03-01', creation='2026-03-01',
            enabled=1, personal_amount=120, company_amount=240, remarks='社保调整')
        rows = self.api._build_compensation_export_rows(
            [employee],
            [salary('opening-salary', '2025-12-01', 3000), salary('equal-salary', '2026-02-01', 3000), salary('raise', '2026-04-01', 3200)],
            [contribution],
            '2026-01-01', '2026-08-10',
            ['base_salary', 'social_personal'],
        )
        self.assertEqual([row['change_type'] for row in rows], ['期初沿用', '定薪', '社保', '定薪'])
        self.assertEqual([str(row['period_start']) for row in rows], ['2026-01-01', '2026-02-01', '2026-03-01', '2026-04-01'])
        self.assertEqual([str(row['period_end']) for row in rows], ['2026-01-31', '2026-02-28', '2026-03-31', '2026-08-10'])
        self.assertEqual(rows[0]['base_salary'], rows[1]['base_salary'])
        self.assertEqual(rows[2]['social_personal'], 120)

    def test_compensation_export_only_emits_events_for_selected_content(self):
        employee = Row(name='e1', employee_name='员工甲', employee_code='E1')
        salary = Row(name='salary', employee='e1', effective_date='2026-01-01', approved_on='2026-01-01',
            creation='2026-01-01', base_salary=3000, full_salary=3000)
        social = Row(name='social', employee='e1', contribution_type='社保', effective_date='2026-02-01',
            approved_on='2026-02-01', creation='2026-02-01', enabled=1, personal_amount=100, company_amount=200)
        rows = self.api._build_compensation_export_rows(
            [employee], [salary], [social], '2026-01-01', '2026-03-31', ['base_salary'])
        self.assertEqual([row['change_type'] for row in rows], ['定薪'])
        self.assertEqual(str(rows[0]['period_end']), '2026-03-31')

    def test_compensation_export_creates_typed_filtered_xlsx(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest('openpyxl is unavailable')

        employee = Row(name='e1', employee_name='员工甲', custom_employee_code='E001',
            department='DEP-1', employment_type='Full-time', status='Active')
        salary = Row(name='salary-1', employee='e1', effective_date='2026-01-01',
            approved_on='2026-01-02', creation='2026-01-01', base_salary=3000,
            function_allowance=200, certificate_allowance=0, multi_skill_allowance=0,
            full_salary=3200, remarks='首次定薪')
        department = Row(name='DEP-1', department_name='工程课')
        def get_all(doctype, **kwargs):
            return {'Employee': [employee], 'Department': [department],
                self.api.SALARY: [salary], self.api.CONTRIBUTION: []}.get(doctype, [])
        self.frappe.get_all = get_all
        self.frappe.db.get_all = lambda doctype, **kwargs: [department] if doctype == 'Department' else []
        stored = {}
        class FileDoc(Row):
            def insert(self, **kwargs):
                stored.update(self)
                self.file_url = '/private/files/export.xlsx'
                return self
        self.frappe.get_doc = lambda *args: FileDoc(args[0]) if len(args) == 1 and isinstance(args[0], dict) else SimpleNamespace(check_permission=lambda *a: None)
        payroll = sys.modules['hrms.api.payroll_input']
        payroll._safe_fields = lambda _doctype, fields: fields
        payroll._salary_contribution_defaults = lambda row, effective: {'employment_stage': '正式'}
        watermark = ModuleType('hrms.utils.export_watermark')
        watermark.save_workbook_with_logo_watermark = lambda workbook, output: workbook.save(output)
        with patch.dict(sys.modules, {'hrms.utils.export_watermark': watermark}):
            result = self.api.export_compensation_register(
                'ACME', '2026-01-01', '2026-08-10',
                json.dumps(['base_salary', 'social_personal', 'remarks']), 'DEP-1', 'e1')
        workbook = load_workbook(BytesIO(stored['content']))
        sheet = workbook['工资社保历史']
        self.assertEqual(result['row_count'], 1)
        self.assertEqual([cell.value for cell in sheet[1]], [
            '姓名', '工号', '工作性质', '部门', '生效开始', '生效结束', '记录类型',
            '底薪', '社保个人承担', '变更原因'])
        self.assertEqual(sheet['A2'].value, '员工甲')
        self.assertEqual(sheet['B2'].value, 'E001')
        self.assertEqual(sheet['D2'].value, '工程课')
        self.assertEqual(sheet['H2'].value, 3000)
        self.assertIsInstance(sheet['E2'].value, datetime)
        self.assertEqual(sheet.auto_filter.ref, 'A1:J2')

    def test_compensation_export_whitelist_arguments_are_typed(self):
        parameters = inspect.signature(self.api.export_compensation_register).parameters
        self.assertTrue(all(parameter.annotation is not inspect.Parameter.empty for parameter in parameters.values()))

    def test_salary_selector_uses_approved_only_with_month_end(self):
        add=self.selector_store();add('salary');add('pending','待审核','2026-08-01');add('future',effective='2027-01-01')
        tree=ast.parse((ROOT/'hrms/api/payroll_input.py').read_text())
        fn=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='get_active_salary_change_for_employee');fn.decorator_list=[]
        ns=dict(frappe=self.frappe,EMPLOYEE_SALARY_CHANGE_DOCTYPE='salary',_require_company=lambda c:c,_month_end=lambda m:m+'-30')
        exec(compile(ast.Module(body=[fn],type_ignores=[]),'payroll_input.py','exec'),ns)
        self.assertEqual(ns[fn.name](employee='e1',company='ACME',payroll_month='2026-09').name,'salary')

    def test_payroll_totals_ignore_legacy_monthly_contributions_and_retain_trace(self):
        legacy=Row(name='legacy',company='ACME',employee='e1',variable_type='社保个人',amount=999,attendance_lock_version='',source_sheet='old',source_hash='old')
        bonus=Row(name='bonus',company='ACME',employee='e1',variable_type='其他奖金',amount=50,attendance_lock_version='',source_sheet='bonus',source_hash='bonus')
        standing=Row(name='standing',company='ACME',employee='e1',contribution_type='社保',personal_amount=100,company_amount=200,enabled=1)
        self.frappe.get_all=lambda *a,**k:[legacy,bonus]
        ns=dict(frappe=self.frappe,_require_company=lambda c:c,_doctype_has_field=lambda *a:False,
            VARIABLE_RECORD_DOCTYPE='variables',VARIABLE_BATCH_DOCTYPE='batch',defaultdict=defaultdict,
            CONTRIBUTION_TYPES=self.api.CONTRIBUTION_TYPES,_monthly_variable_scope=lambda m:'monthly:'+m,
            _assert_row_company=lambda *a:None,_employee_identity_key=lambda r:r.employee,
            flt=lambda v:float(v or 0),active_contributions=lambda *a:[standing],_month_end=lambda m:m+'-30',
            _source_trace_hash=lambda row:('json','hash'),_=lambda s:s)
        tree=ast.parse((ROOT/'hrms/api/payroll_input.py').read_text());fn=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='_variable_totals')
        exec(compile(ast.Module(body=[fn],type_ignores=[]),'payroll_input.py','exec'),ns)
        totals,_,sources=ns['_variable_totals']('ACME','2026-09')
        self.assertEqual(dict(totals['e1']),{'其他奖金':50,'社保个人':100,'社保公司':200})
        self.assertEqual([r['name'] for r in sources['e1']],['bonus','standing','standing'])

if __name__=='__main__':unittest.main()
