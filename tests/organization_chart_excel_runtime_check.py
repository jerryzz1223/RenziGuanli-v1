"""Read-only live-tree checks plus synthetic export boundary cases."""
from collections import Counter
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
import json

import frappe
from openpyxl import load_workbook
from hrms.api.organization_chart_export import workbook_bytes, select_tree, _walk, export_excel
from hrms.hr.page.organizational_chart.organizational_chart import get_hybrid_tree


def run():
	frappe.set_user('Administrator')
	root = get_hybrid_tree(company='永新')['root']
	before = deepcopy(root)
	data, sheets = workbook_bytes(root, '2026-09-09 12:00')
	book = load_workbook(BytesIO(data))
	assert root == before
	assert not any(name.startswith('xl/media/') for name in ZipFile(BytesIO(data)).namelist()), 'must use native cells, not screenshots'
	values = [str(c.value) for row in book.worksheets[0] for c in row if c.value is not None]
	text = ''.join(values)
	for node in _walk(root):
		assert str(node.get('name') or '') in text
		for person in node.get('people') or []:
			assert (person.get('employee_name') or person.get('name') or '') in text
	assert '课长（正式）：朱耀辉' in text and '代理人：张付俊' in text
	assert text.count('王传瑞') >= 3, 'retain all three source appointments'
	assert '组长（任职待确认）：王传瑞' in text
	assert '编制/实际：107/82' in ''.join(str(c.value or '') for row in book['连续课'] for c in row)
	for sheet in book:
		by_row = {}
		for merged in sheet.merged_cells.ranges:
			for row in range(merged.min_row, merged.max_row+1):
				by_row.setdefault(row, []).append((merged.min_col, merged.max_col))
		for ranges in by_row.values():
			ranges.sort()
			assert all(left[1] < right[0] for left, right in zip(ranges, ranges[1:])), 'overlapping cards'
		assert sheet.sheet_view.showGridLines is False
		assert sheet.page_setup.scale == 100 and not sheet.sheet_properties.pageSetUpPr.fitToPage
		assert all(c.data_type != 'f' for row in sheet for c in row), 'snapshot text cannot become formulas'
	sample = {'node_id':'root','name':'测试公司','children':[
		{'node_id':'dept','name':'测试课','organization_node_type':'课','children':[
			{'node_id':'staff','name':'人员','people':[{'name':f'职员{i:03d}','role':'代理任职'} for i in range(82)],'card_content':'=HYPERLINK("https://example.com")','children':[]}
		]}, {'node_id':'other','name':'其他','children':[]} ]}
	current = select_tree(sample, 'current', ['root','dept','other'])
	assert len(list(_walk(current))) == 3 and current['children'][0]['export_collapsed'] == 1
	assert len(list(_walk(select_tree(sample,'current',['dept','staff'])))) == 2
	for ids in (['invalid'], ['dept','other']):
		try: select_tree(sample, 'current', ids)
		except ValueError: pass
		else: raise AssertionError('must reject changed or disconnected scope')
	synthetic, _ = workbook_bytes(sample, 'test')
	test_book = load_workbook(BytesIO(synthetic))
	cells = [c for row in test_book.worksheets[0] for c in row if c.value is not None]
	joined = ''.join(str(c.value) for c in cells)
	assert all(f'职员{i:03d}' in joined for i in range(82))
	assert all(c.data_type == 's' for c in cells)
	frappe.set_user('Guest')
	try:
		try: export_excel('永新')
		except frappe.PermissionError: pass
		else: raise AssertionError('guest must not export personnel')
	finally: frappe.set_user('Administrator')
	print(json.dumps({'pass':True,'nodes':len(list(_walk(root))),'sheets':sheets,'bytes':len(data),'checks':'people, multiple roles, no image, no formulas, no overlaps, 82-person list, current scope, permissions'},ensure_ascii=False))
