import importlib.util
import io
import sys
import unittest
from pathlib import Path
from types import ModuleType
from unittest.mock import patch

from openpyxl import load_workbook


class OrganizationChartExportTest(unittest.TestCase):
	@classmethod
	def setUpClass(cls):
		frappe = ModuleType("frappe")
		frappe.whitelist = lambda *args, **kwargs: (lambda fn: fn)
		utils = ModuleType("frappe.utils")
		utils.now_datetime = lambda: None
		frappe.utils = utils
		path = Path(__file__).parents[1] / "hrms/api/organization_chart_export.py"
		with patch.dict(sys.modules, {"frappe": frappe, "frappe.utils": utils}):
			spec = importlib.util.spec_from_file_location("organization_chart_export_test", path)
			cls.module = importlib.util.module_from_spec(spec)
			spec.loader.exec_module(cls.module)

	def test_report_is_appended_below_chart_on_main_sheet(self):
		root = {
			"node_id": "root",
			"name": "测试公司",
			"children": [{
				"node_id": "department",
				"name": "品保课",
				"organization_node_type": "课",
				"planned_headcount": 14,
				"current_headcount": 12,
				"vacancy_count": 2,
				"children": [],
			}],
		}
		report = {
			"title": "测试公司组织报表",
			"columns": ["部门/课别", "编制人数", "现有人数", "空缺人数", "岗位满足率", "备注"],
			"rows": [{
				"department": "品保课",
				"planned_headcount": 14,
				"current_headcount": 12,
				"vacancy_count": 2,
				"fulfillment_rate": 12 / 14,
				"vacancy_notes": "招聘计划：2人",
			}],
			"total": {
				"planned_headcount": 14,
				"current_headcount": 12,
				"vacancy_count": 2,
				"fulfillment_rate": 12 / 14,
			},
		}
		data, sheet_names = self.module.workbook_bytes(
			root,
			"2026-09-15 12:00",
			split_departments=False,
			report=report,
		)
		book = load_workbook(io.BytesIO(data))
		sheet = book.worksheets[0]
		self.assertEqual(sheet_names, ["当前层级架构图"])

		locations = {}
		for row in sheet.iter_rows():
			for cell in row:
				if cell.value is not None:
					locations.setdefault(str(cell.value), []).append(cell)
		report_title = locations[report["title"]][0]
		chart_department = locations["品保课"][0]
		self.assertGreater(report_title.row, chart_department.row)

		header_row = locations["部门/课别"][0].row
		self.assertEqual(
			[cell.value for cell in sheet[header_row] if cell.value is not None],
			report["columns"],
		)
		data_row = header_row + 1
		values = [cell.value for cell in sheet[data_row] if cell.value is not None]
		self.assertEqual(values[:4], ["品保课", 14, 12, 2])
		self.assertAlmostEqual(values[4], 12 / 14)
		self.assertEqual(values[5], "招聘计划：2人")
		fulfillment_cell = next(cell for cell in sheet[data_row] if cell.value == 12 / 14)
		self.assertEqual(fulfillment_cell.number_format, "0%")

		total_row = header_row + 2
		self.assertEqual([cell.value for cell in sheet[total_row] if cell.value is not None], ["汇总", 14, 12, 2, 12 / 14, "-"])
		approval = locations["批准："][0]
		self.assertGreater(approval.row, total_row)
		self.assertGreaterEqual(sheet.max_row, approval.row)
		self.assertTrue(sheet.print_area)
		self.assertFalse(any(cell.data_type == "f" for row in sheet for cell in row))


if __name__ == "__main__":
	unittest.main()
