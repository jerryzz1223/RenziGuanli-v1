from __future__ import annotations

import importlib.util
import sys
import types
from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import patch
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile


MODULE_PATH = Path(__file__).resolve().parents[1] / "hrms" / "utils" / "export_watermark.py"
SPEC = importlib.util.spec_from_file_location("export_watermark", MODULE_PATH)
WATERMARK = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
SPEC.loader.exec_module(WATERMARK)
_add_logo_watermark = WATERMARK._add_logo_watermark


def _minimal_workbook():
	content = BytesIO()
	with ZipFile(content, "w", ZIP_DEFLATED) as archive:
		archive.writestr(
			"[Content_Types].xml",
			'<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"></Types>',
		)
		archive.writestr(
			"xl/worksheets/sheet1.xml",
			'<?xml version="1.0"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData/></worksheet>',
		)
		archive.writestr(
			"xl/worksheets/sheet2.xml",
			'<?xml version="1.0"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData/></worksheet>',
		)
	return content.getvalue()


class TestExportWatermark(unittest.TestCase):
	def test_configured_export_and_print_are_independent(self):
		from openpyxl import Workbook

		settings = types.ModuleType("hrms.utils.export_watermark_settings")
		options = {"export": False, "print": False, "text": "测试水印", "opacity": 12}
		settings.get_export_watermark_options = lambda key: options.copy()
		with patch.dict(sys.modules, {"hrms.utils.export_watermark_settings": settings}):
			plain = BytesIO()
			WATERMARK.save_workbook_with_logo_watermark(Workbook(), plain, export_key="roster")
			with ZipFile(BytesIO(plain.getvalue())) as archive:
				self.assertNotIn(WATERMARK._WATERMARK_MEDIA_PATH, archive.namelist())
				self.assertNotIn(b"picture", archive.read("xl/worksheets/sheet1.xml"))

			options.update(export=True, print=False)
			screen_only = BytesIO()
			WATERMARK.save_workbook_with_logo_watermark(Workbook(), screen_only, export_key="roster")
			with ZipFile(BytesIO(screen_only.getvalue())) as archive:
				self.assertIn(WATERMARK._WATERMARK_MEDIA_PATH, archive.namelist())
				self.assertNotIn(b"<headerFooter>", archive.read("xl/worksheets/sheet1.xml"))

			options.update(export=True, print=True)
			marked = BytesIO()
			WATERMARK.save_workbook_with_logo_watermark(Workbook(), marked, export_key="roster")
			with ZipFile(BytesIO(marked.getvalue())) as archive:
				self.assertIn(WATERMARK._WATERMARK_MEDIA_PATH, archive.namelist())
				sheet = archive.read("xl/worksheets/sheet1.xml")
				self.assertIn(b"picture", sheet)
				header = ElementTree.fromstring(sheet).find(f"{{{WATERMARK._SPREADSHEET_NAMESPACE}}}headerFooter/{{{WATERMARK._SPREADSHEET_NAMESPACE}}}oddHeader")
				self.assertIn("测试水印", header.text)

	def test_print_header_escapes_excel_format_character(self):
		from openpyxl import Workbook

		settings = types.ModuleType("hrms.utils.export_watermark_settings")
		settings.get_export_watermark_options = lambda key: {"export": False, "print": True, "text": "A&B", "opacity": 20}
		with patch.dict(sys.modules, {"hrms.utils.export_watermark_settings": settings}):
			output = BytesIO()
			WATERMARK.save_workbook_with_logo_watermark(Workbook(), output, export_key="roster")
			with ZipFile(BytesIO(output.getvalue())) as archive:
				sheet = archive.read("xl/worksheets/sheet1.xml")
				self.assertIn(b"A&amp;&amp;B", sheet)

	def test_background_relationship_is_valid_with_header_comments(self):
		namespace = WATERMARK._DOCUMENT_RELATIONSHIPS_NAMESPACE
		content = (
			f'<worksheet xmlns="{WATERMARK._SPREADSHEET_NAMESPACE}"><sheetData/>'
			f'<legacyDrawing xmlns:r="{namespace}" r:id="rId1"/></worksheet>'
		).encode()
		root = ElementTree.fromstring(WATERMARK._add_background_picture(content, "rId2"))
		picture = root.find(f"{{{WATERMARK._SPREADSHEET_NAMESPACE}}}picture")
		self.assertEqual(picture.attrib[f"{{{namespace}}}id"], "rId2")
		self.assertEqual(root.find(f"{{{WATERMARK._SPREADSHEET_NAMESPACE}}}legacyDrawing").attrib[f"{{{namespace}}}id"], "rId1")

	def test_export_uses_the_pale_watermark_asset(self):
		asset_path = Path(WATERMARK.__file__).resolve().parents[1] / "public" / "images" / "yongxin-brand-watermark.png"
		self.assertTrue(asset_path.is_file())
		self.assertNotEqual(
			asset_path.read_bytes(),
			asset_path.with_name("yongxin-brand-mark.png").read_bytes(),
		)

	def test_export_workbook_has_yongxin_logo_background_on_each_sheet(self):
		watermarked = _add_logo_watermark(_minimal_workbook())
		with ZipFile(BytesIO(watermarked)) as workbook:
			for sheet_number in (1, 2):
				sheet = workbook.read(f"xl/worksheets/sheet{sheet_number}.xml")
				relationships = workbook.read(f"xl/worksheets/_rels/sheet{sheet_number}.xml.rels")
				self.assertIn(b"<picture r:id=", sheet)
				self.assertIn(b"hrms-yongxin-watermark.png", relationships)
			self.assertIn("xl/media/hrms-yongxin-watermark.png", workbook.namelist())

	def test_export_watermark_preserves_a_workbook_that_already_has_one(self):
		once = _add_logo_watermark(_minimal_workbook())
		twice = _add_logo_watermark(once)
		with ZipFile(BytesIO(twice)) as workbook:
			self.assertEqual(workbook.read("xl/worksheets/sheet1.xml").count(b"<picture r:id="), 1)

	def test_export_watermark_uses_next_available_single_quoted_relationship_id(self):
		relationships, relationship_id = WATERMARK._add_image_relationship(
			b"<Relationships><Relationship Id='rId12'/></Relationships>"
		)

		self.assertEqual(relationship_id, "rId13")
		self.assertIn(b'Id="rId13"', relationships)
