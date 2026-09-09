"""Portable organization configuration. HR master records are never written.

The workbook uses stable node keys, company-scoped department labels and employee
business codes. A preview is bound to the exact workbook and target snapshot.
"""
import hashlib
import io
import json
import zipfile
from collections import defaultdict

import frappe
from frappe.utils import cstr, now_datetime
from hrms.utils.organization_roles import binding_assignment_type, ASSIGNMENT_TYPES, base_role

SCHEMA = "HRMS-ORGANIZATION-2"
NODE_COLUMNS = {
	"节点编号": "portable_id", "上级节点编号": "parent", "组织名称": "display_name",
	"节点类型": "node_kind", "关联部门": "department", "关联岗位": "designation",
	"职级编码": "chart_grade_code", "花名册职级引用": "grade", "部门内分组": "roster_subset",
	"人员自动更新": "roster_auto_sync", "图中职务": "role_title", "任职方式": "assignment_mode",
	"已设置编制": "planned_headcount_set", "编制人数": "planned_headcount",
	"管理层读取花名册": "leadership_from_roster", "保留原表负责人": "template_leadership",
	"来源单元格": "template_source_cell", "负责人来源单元格": "template_leadership_cell", "来源文档": "template_source_document",
	"原表空缺": "template_source_vacancies", "归属待设置": "reporting_scope_pending",
	"合并关联部门": "roster_department_alias_labels", "自动生成节点": "roster_generated",
	"分管显示姓名": "manager_name", "使用原表人员规则": "has_template_bindings", "逐人确认任职": "assignment_rules_manual",
	"原表职级标签": "source_grade_tags", "原表职级来源": "source_grade_reference", "原表职级确认状态": "source_grade_status",
}
OPTIONAL_NODE_FIELDS = {"source_grade_tags", "source_grade_reference", "source_grade_status"}
PERSON_COLUMNS = {
	"节点编号": "node", "引用类型": "type", "工号": "code", "姓名": "name",
	"图中职务": "role", "负责人位置": "slot", "仅展示代理": "display_only", "人工确认": "manual_confirmed", "原表职务": "source_role", "任职性质": "assignment_type",
}
GRADE_COLUMNS = {"职级编码": "code", "职级名称": "label", "等级顺序": "rank", "上级职级编码": "parent"}
BOOL_FIELDS = {"roster_subset", "roster_auto_sync", "planned_headcount_set", "leadership_from_roster",
	"template_leadership", "reporting_scope_pending", "roster_generated", "has_template_bindings", "assignment_rules_manual", "manual_confirmed"}
REFERENCES = {"分管人员": "manager_employee", "员工节点": "employee", "任职人": "primary_employee",
	"代理人": "proxy_employee", "岗位成员": "assigned_employees", "原表人员": "template_bindings"}


def chart_module():
	from hrms.hr.page.organizational_chart import organizational_chart
	return organizational_chart


def authorize(company):
	frappe.only_for("System Manager")
	frappe.get_doc("Company", company).check_permission("read")
	for action in ("read", "create", "write"):
		frappe.has_permission("Organization Node", action, throw=True)


def version_options(version):
	try:
		value = json.loads(frappe.db.get_value("Organization Structure Version", version, "notes") or "{}")
		return value if isinstance(value, dict) else {}
	except ValueError:
		return {}


def target_state(company):
	manual = chart_module()._get_manual_organization_records(company)
	from hrms.api.organization_source_grades import source_grade_evidence
	return {
		"manual": manual,
		"source_grades": source_grade_evidence(company, manual["nodes"]),
		"employees": frappe.get_all("Employee", filters={"company": company}, fields=["name", "custom_employee_code", "employee_name", "department", "designation", "status"], order_by="name"),
		"departments": frappe.get_all("Department", filters={"company": company, "disabled": 0}, fields=["name", "department_name"], order_by="name"),
		"designations": frappe.get_all("Designation", pluck="name", order_by="name"),
		"grades": frappe.get_all("Employee Grade", pluck="name", order_by="name"),
		"options": version_options(manual["version"]) if manual["version"] else {},
	}


def workbook_bytes(company, state=None):
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill, Alignment
	from openpyxl.utils import get_column_letter
	state = state if state is not None else target_state(company)
	nodes = state["manual"]["nodes"]
	ids = {n.name: n.manual_config.get("portable_id") or n.node_code for n in nodes}
	employees = {e.name: e for e in state["employees"]}
	departments = {d.name: d.department_name for d in state["departments"]}
	rows, people = [], []
	for node in nodes:
		cfg = {**state.get("source_grades", {}).get(node.name, {}), **node.manual_config}
		row = {**cfg, "portable_id": ids[node.name], "parent": ids.get(node.parent_node, ""),
			"display_name": node.display_name, "planned_headcount": node.planned_headcount or 0,
			"department": departments.get(cfg.get("department"), cfg.get("department") or ""),
			"has_template_bindings": "template_bindings" in cfg,
			"roster_department_alias_labels": "\n".join(cfg.get("roster_department_alias_labels", []))}
		rows.append(row)
		for label, field in REFERENCES.items():
			if field == "assigned_employees" and cfg.get("node_kind") != "岗位" and not cfg.get("roster_subset"): continue
			values = cfg.get(field, []) if field in {"assigned_employees", "template_bindings"} else [cfg.get(field)]
			for value in values:
				if not value:
					continue
				binding = value if isinstance(value, dict) else {}
				person = employees.get(binding.get("employee") if binding else value)
				code = cstr(person.custom_employee_code).strip() if person else cstr(binding.get("source_code")).strip()
				if person and not code:
					frappe.throw(f"员工 {person.employee_name} 缺少工号，请先补齐再导出。")
				people.append({"node": ids[node.name], "type": label, "code": code,
					"name": person.employee_name if person else binding.get("source_name", ""),
					"role": binding.get("role", ""), "slot": binding.get("slot", ""), "manual_confirmed": bool(binding.get("manual_confirmed")), "source_role": binding.get("source_role", ""), "assignment_type": binding_assignment_type(binding) if binding else "",
					"display_only": bool(binding.get("display_only") or (field == "proxy_employee" and (
						cfg.get("portable_proxy_display_only") or any(b.get("employee") == value and b.get("display_only") for b in cfg.get("template_bindings", [])))))})
	book = Workbook()
	info = book.active
	info.title = "说明"
	for row in [
		["格式版本", SCHEMA], ["来源公司", company],
		["公司根节点", "导入时选定的公司自动作为根节点；上级节点编号留空表示公司直属。"],
		["上下级关系", "节点编号必须唯一且保持不变；上级节点编号决定汇报关系，不按职位名称猜测。"],
		["人员匹配", "工号按文本填写，保留前导零；仅按目标公司工号匹配。原表人员可不填工号，作为待确认注释。"],
		["自动更新", "自动岗位按部门和岗位读取花名册；原表组线按已绑定工号更新。手动节点保留人员安排。"],
		["职级", "职级编码引用职级定义；等级顺序越小越高，允许留空。职级不代替上级节点关系，也不修改员工档案职级。"],
		["原表职级标签", "保留原组织草稿中与来源文件、工作表及单元格唯一对应的标签和确认状态。标签不代表高低顺序，不写入员工档案职级。"],
		["花名册职级引用", "现有花名册职级的引用，和图中职级编码分开；需在目标服务器已存在。"],
		["重复导入", "按节点编号更新，未包含节点保留；预览通过后才能导入。"],
		["代理", "引用类型代理人保留代理身份；原表人员的负责人位置填 primary 或 proxy。跨部门代理须勾选仅展示代理。"],
		["多人任职", "同一工号可在多个节点或同节点不同职务出现，汇总人数去重；唯一正式职位需人工确认，同一公司不允许多个已确认正式职位。未确认的原表任职保留待确认。图中职务保留（代）、（兼）；代理任职不等于代理人。人工确认与逐人确认任职随文件保留。仅展示代理也适用于明确的跨部门兼任。"],
		["布尔值", "所有开关填写 1 或 0。合并关联部门每行一个部门名称。"],
		["部署", "目标服务器需要安装包含组织配置导入功能的系统版本，并先导入花名册、部门和岗位。"],
	]: info.append(row)
	for title, columns, data in [("组织层级", NODE_COLUMNS, rows), ("人员任职", PERSON_COLUMNS, people),
		("职级定义", GRADE_COLUMNS, state["options"].get("chart_grades", []))]:
		sheet = book.create_sheet(title)
		sheet.append(list(columns))
		for row in data:
			sheet.append([int(bool(row.get(key))) if key in BOOL_FIELDS or key == "display_only" else row.get(key, "") for key in columns.values()])
		sheet.auto_filter.ref = sheet.dimensions
	for sheet in book:
		sheet.freeze_panes = "A2"
		for cell in sheet[1]:
			cell.font = Font(bold=True, color="FFFFFF")
			cell.fill = PatternFill("solid", fgColor="315C8C")
		for row in sheet:
			for cell in row:
				if isinstance(cell.value, str):
					cell.data_type = "s"  # Treat leading '=' as text, never executable spreadsheet formulas.
					cell.number_format = "@"
				cell.alignment = Alignment(vertical="top", wrap_text=True)
		for col in range(1, sheet.max_column + 1):
			sheet.column_dimensions[get_column_letter(col)].width = 24
	info.column_dimensions["B"].width = 110
	stream = io.BytesIO()
	book.save(stream)
	return stream.getvalue()


def save_private(content, prefix, company):
	from frappe.utils.file_manager import save_file
	filename = f"{prefix}-{now_datetime().strftime('%Y%m%d-%H%M%S')}-{frappe.generate_hash(length=6)}.xlsx"
	file = save_file(filename, content, None, None, is_private=1)
	return {"file_url": file.file_url, "file_name": file.file_name}


@frappe.whitelist(methods=["POST"])
def export_configuration(company: str):
	authorize(company)
	state = target_state(company)
	content = workbook_bytes(company, state)
	plan = prepare(company, read_content(content), state=state)
	return {**save_private(content, "组织配置", company), "completeness": plan["completeness"],
		"warnings": plan["warnings"], "errors": plan["errors"]}


def read_package(file_url):
	if not cstr(file_url).startswith("/private/files/"):
		frappe.throw("请上传私有的 .xlsx 组织配置文件。")
	name = frappe.db.get_value("File", {"file_url": file_url, "is_private": 1}, "name")
	if not name:
		frappe.throw("找不到组织配置文件。")
	file = frappe.get_doc("File", name)
	file.check_permission("read")
	content = file.get_content()
	return read_content(content)


def read_content(content):
	from openpyxl import load_workbook
	if len(content) > 10 * 1024 * 1024:
		frappe.throw("组织配置文件不能超过 10 MB。")
	try:
		with zipfile.ZipFile(io.BytesIO(content)) as archive:
			if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
				frappe.throw("组织配置解压后过大。")
		book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
	except (ValueError, zipfile.BadZipFile, KeyError):
		frappe.throw("无法读取文件，请使用系统导出的 .xlsx 模板。")
	try:
		if "说明" not in book.sheetnames or book["说明"]["B1"].value not in {SCHEMA, "HRMS-ORGANIZATION-1"}:
			frappe.throw("组织配置格式版本不匹配，请使用系统导出的模板。")
		package = {}
		for title, columns in [("组织层级", NODE_COLUMNS), ("人员任职", PERSON_COLUMNS), ("职级定义", GRADE_COLUMNS)]:
			if title not in book.sheetnames:
				frappe.throw(f"缺少工作表：{title}")
			sheet = book[title]
			if sheet.max_row > 10001 or sheet.max_column > 50:
				frappe.throw(f"{title} 超出 10000 行或 50 列限制。")
			rows = sheet.iter_rows()
			headers = [cstr(c.value).strip() for c in next(rows)]
			if any(headers.count(h) > 1 or (headers.count(h) == 0 and h not in {"逐人确认任职", "人工确认", "原表职务", "负责人来源单元格", "任职性质", "原表职级标签", "原表职级来源", "原表职级确认状态"}) for h in columns):
				frappe.throw(f"{title} 表头缺失或重复，请保留模板表头。")
			data = []
			for index, cells in enumerate(rows, 2):
				if any(c.data_type == "f" for c in cells):
					frappe.throw(f"{title} 第 {index} 行含公式，请粘贴为值。")
				row = {key: cstr(cells[headers.index(label)].value).strip() if label in headers and headers.index(label) < len(cells) else "" for label, key in columns.items()}
				for label, key in columns.items():
					if key in OPTIONAL_NODE_FIELDS and label not in headers:
						row.pop(key, None)
				if any(row.values()):
					data.append({**row, "_row": index})
			package[title] = data
		return package
	finally:
		book.close()


def prepare(company, package, state=None):
	"""Pure planning against a read-only target snapshot; collect all actionable errors."""
	chart = chart_module()
	state = state if state is not None else target_state(company)
	errors, warnings = [], []
	def problem(row, message): errors.append(f"第 {row.get('_row', '?')} 行：{message}")
	def boolean(row, field):
		value = row.get(field, "")
		if value not in ("", "0", "1"):
			problem(row, f"{field} 必须为 0 或 1")
		return value == "1"
	def integer(row, field):
		value = row.get(field, "")
		if value and (not value.isdigit() or int(value) > 1000000):
			problem(row, f"{field} 必须是 0 至 1000000 的整数")
			return 0
		return int(value or 0)
	def index(items, field):
		result = defaultdict(list)
		for item in items: result[cstr(item.get(field)).strip()].append(item)
		return result
	departments = index(state["departments"], "department_name")
	staff = index(state["employees"], "custom_employee_code")
	current = {}
	for node in state["manual"]["nodes"]:
		key = node.manual_config.get("portable_id") or node.node_code
		if key in current: errors.append(f"现有节点编号重复：{key}")
		current[key] = node
	grades = {g["code"]: g for g in state["options"].get("chart_grades", [])}
	input_grades = set()
	for row in package["职级定义"]:
		code = row["code"]
		if not code or code in input_grades or not row["label"]:
			problem(row, "职级编码必须唯一且名称不能为空")
		input_grades.add(code)
		grades[code] = {**row, "rank": integer(row, "rank") if row["rank"] else None}
	def order_graph(items, parent_field, label):
		order, visiting, done = [], set(), set()
		def visit(key, depth=0):
			if key in done: return
			if key in visiting or depth > 100:
				errors.append(f"{label} 存在循环或超过 100 层：{key}")
				return
			visiting.add(key)
			parent = items[key].get(parent_field)
			if parent:
				if parent not in items: errors.append(f"{label} {key} 的上级不存在：{parent}")
				else: visit(parent, depth + 1)
			visiting.remove(key)
			done.add(key)
			order.append(key)
		for key in items: visit(key)
		return order
	order_graph(grades, "parent", "职级")
	for grade in grades.values():
		parent = grades.get(grade["parent"])
		if parent and grade["rank"] is not None and parent["rank"] is not None and grade["rank"] <= parent["rank"]:
			problem(grade, "下级职级的等级顺序必须大于上级")
	plans = {}
	for row in package["组织层级"]:
		key, kind = row["portable_id"], row["node_kind"]
		if not key or key in plans:
			problem(row, "节点编号不能为空或重复")
		if not row["display_name"] or kind not in chart.MANUAL_ORGANIZATION_NODE_KINDS:
			problem(row, "组织名称或节点类型无效")
		cfg = {k: row.get(k, "") for k in NODE_COLUMNS.values() if k not in {"parent", "display_name", "planned_headcount", "has_template_bindings"}}
		previous = current.get(key)
		for field in OPTIONAL_NODE_FIELDS:
			if field not in row and previous:
				cfg[field] = previous.manual_config.get(field, state.get("source_grades", {}).get(previous.name, {}).get(field, ""))
		if cfg["source_grade_status"] not in {"", "待确认", "已确认"}:
			problem(row, "原表职级确认状态必须为待确认或已确认")
		cfg.update({k: boolean(row, k) for k in BOOL_FIELDS if k not in {"has_template_bindings", "manual_confirmed"}})
		cfg.update(manual_organization=True, framework=True, assigned_employees=[])
		cfg["template_source_vacancies"] = integer(row, "template_source_vacancies")
		cfg["roster_department_alias_labels"] = list(dict.fromkeys(filter(None, row["roster_department_alias_labels"].splitlines())))
		if boolean(row, "has_template_bindings"): cfg["template_bindings"] = []
		matches = departments.get(row["department"], [])
		if row["department"] and len(matches) != 1 or kind in chart.MANUAL_ORGANIZATION_DEPARTMENT_KINDS and not row["department"]:
			problem(row, f"关联部门未唯一匹配：{row['department'] or '未填写'}；请先导入目标花名册和部门")
		cfg["department"] = matches[0].name if len(matches) == 1 else None
		cfg["roster_department"] = cfg["department"]
		if row["designation"] and row["designation"] not in state["designations"]:
			problem(row, f"关联岗位不存在：{row['designation']}")
		if kind == "岗位" and not cfg["roster_subset"] and not row["designation"]:
			problem(row, "岗位节点必须填写关联岗位")
		if cfg["roster_subset"] and kind not in {"室", "组", "线", "岗位"}:
			problem(row, "该节点类型不能作为部门内分组")
		if cfg["roster_subset"] and cfg["roster_auto_sync"] and "template_bindings" not in cfg:
			problem(row, "自动分组必须有原表人员规则，不能从职位猜测组线归属")
		if cfg["assignment_mode"] not in {"自动", "正式", "代理"}: problem(row, "任职方式无效")
		if cfg["grade"] and cfg["grade"] not in state["grades"]: problem(row, f"花名册职级不存在：{cfg['grade']}")
		if cfg["chart_grade_code"] and cfg["chart_grade_code"] not in grades: problem(row, f"职级定义缺少：{cfg['chart_grade_code']}")
		cfg["chart_grade"] = {k: v for k, v in grades.get(cfg["chart_grade_code"], {}).items() if k != "_row"}
		plans[key] = {"config": cfg, "parent": row["parent"], "name": row["display_name"], "planned": integer(row, "planned_headcount"), "_row": row["_row"]}
	if not plans: errors.append("组织层级不能为空")
	# Include retained nodes when checking cycles and conflicting whole-department identities.
	existing_ids = {n.name: key for key, n in current.items()}
	graph = {key: {"parent": existing_ids.get(n.parent_node, ""), "config": n.manual_config} for key, n in current.items()}
	graph.update(plans)
	order = order_graph(graph, "parent", "组织节点")
	units, aliases, auto_positions = {}, {}, {}
	for key, plan in graph.items():
		cfg = plan["config"]
		parent = graph.get(plan["parent"])
		if parent and parent["config"].get("node_kind") == "员工": errors.append(f"员工节点不能有下级：{key}")
		dept = cfg.get("department")
		if chart.whole_department(cfg) and dept:
			if dept in units: errors.append(f"同一部门存在两个组织节点：{dept}；请先合并")
			units[dept] = key
		if cfg.get("node_kind") == "岗位" and not cfg.get("roster_subset") and cfg.get("roster_auto_sync"):
			position = (dept, cfg.get("designation"))
			if position in auto_positions: errors.append(f"同一部门岗位存在重复自动节点：{position}")
			auto_positions[position] = key
		for label in cfg.get("roster_department_alias_labels", []):
			if label in aliases and aliases[label] != key: errors.append(f"合并部门重复关联：{label}")
			aliases[label] = key
	for label, key in aliases.items():
		matches = departments.get(label, [])
		if len(matches) > 1: errors.append(f"合并部门未唯一匹配：{label}")
		if matches and matches[0].name in units and units[matches[0].name] != key: errors.append(f"合并部门仍有独立组织节点：{label}")
		if not matches: warnings.append(f"合并部门 {label} 尚未创建；保留关联名称，花名册导入后再匹配。")
	from hrms.utils.organization_scope import department_scopes
	scopes = department_scopes(graph, state["departments"])
	department_labels = {d.name: d.department_name for d in state["departments"]}
	seen_refs = set()
	for row in package["人员任职"]:
		if row["node"] not in plans or row["type"] not in REFERENCES:
			problem(row, "人员任职的节点编号或引用类型无效")
			continue
		cfg = plans[row["node"]]["config"]
		field = REFERENCES[row["type"]]
		kind = cfg["node_kind"]
		if (field == "manager_employee" and kind not in {"管理层", "分管"}
			or field == "employee" and kind != "员工"
			or field == "primary_employee" and kind not in {"室", "课", "组", "线", "岗位"}
			or field == "proxy_employee" and kind == "员工"
			or field == "assigned_employees" and kind != "岗位" and not cfg.get("roster_subset")):
			problem(row, f"{row['type']} 不适用于 {kind} 节点")
		key = (row["node"], field, row["code"] or row["name"], row["role"] if field == "template_bindings" else "", row["slot"] if field == "template_bindings" else "")
		if key in seen_refs: problem(row, "同一节点的人员引用重复")
		seen_refs.add(key)
		matches = staff.get(row["code"], []) if row["code"] else []
		person = matches[0] if len(matches) == 1 else None
		if row["code"] and (not person or person.status != "Active"):
			problem(row, f"工号 {row['code']} 未唯一匹配目标公司在职员工")
		elif not row["code"] and (field != "template_bindings" or not row["name"]):
			problem(row, "人员必须填写工号；仅原表人员注释允许无工号")
		elif not row["code"]:
			warnings.append(f"原表人员 {row['name']} 无工号，仅保留待确认注释，不计人数。")
		if person and row["name"] and row["name"] != person.employee_name:
			warnings.append(f"工号 {row['code']} 姓名不同，采用服务器花名册姓名 {person.employee_name}。")
		display_only = boolean(row, "display_only")
		if row["slot"] not in {"", "primary", "proxy"}: problem(row, "负责人位置只能填 primary 或 proxy")
		if display_only and field not in {"proxy_employee", "template_bindings"}: problem(row, "仅展示代理只能用于代理或原表人员引用")
		if person and cfg.get("department") and person.department not in scopes[row["node"]] and not display_only:
			allowed = "、".join(sorted(department_labels.get(d, d) for d in scopes[row["node"]]))
			problem(row, f"工号 {row['code']} 花名册部门与节点不一致：花名册为 {department_labels.get(person.department, person.department) or '未填写'}，节点 {plans[row['node']]['name']} 允许部门为 {allowed}；请核对归属")
		if field == "template_bindings":
			binding = {"source_name": row["name"] or (person.employee_name if person else ""), "source_code": row["code"],
				"role": row["role"], "slot": row["slot"], "display_only": display_only, "no_auto_match": not bool(row["code"]),
				"manual_confirmed": boolean(row, "manual_confirmed") and bool(person), "source_role": row.get("source_role", "")}
			assignment_type = row.get("assignment_type", "")
			if assignment_type and assignment_type not in ASSIGNMENT_TYPES: problem(row, "任职性质只能填写待确认、正式、代理任职、兼任或代理人")
			if assignment_type == "待确认" and binding["manual_confirmed"]: problem(row, "待确认任职不能同时勾选人工确认")
			if assignment_type:
				binding["assignment_type"] = assignment_type
				if assignment_type == "代理人" and row["slot"] != "proxy": problem(row, "代理人的负责人位置必须填 proxy")
				if assignment_type != "代理人" and row["slot"] == "proxy": problem(row, "proxy 位置的任职性质必须为代理人")
				if assignment_type in {"代理任职", "兼任", "正式"}:
					binding["role"] = base_role(row["role"]) + ("（代）" if assignment_type == "代理任职" else "（兼）" if assignment_type == "兼任" else "")
			if display_only and binding_assignment_type(binding) not in {"兼任", "代理人"}: problem(row, "跨部门展示只能用于明确的兼任或代理人")
			if binding["manual_confirmed"] and not binding["role"]: problem(row, "已确认任职必须填写图中职务")
			if person: binding["employee"] = person.name
			else: binding["issue"] = "未提供工号，待确认"
			cfg.setdefault("template_bindings", []).append(binding)
		elif field == "assigned_employees":
			if person: cfg[field].append(person.name)
		else:
			if cfg.get(field): problem(row, f"{row['type']} 只能有一人")
			cfg[field] = person.name if person else None
			if field == "manager_employee" and person: cfg["manager_name"] = person.employee_name
			if field == "proxy_employee" and display_only and person:
				cfg["portable_proxy_display_only"] = True
	# Validate the resulting graph including nodes retained on the target server.
	from hrms.api.organization_assignment_review import formal_conflicts
	resulting_nodes = [n for key, n in current.items() if key not in plans]
	resulting_nodes.extend(frappe._dict(name=key, display_name=plan["name"], manual_config=plan["config"]) for key, plan in plans.items())
	errors.extend(formal_conflicts(resulting_nodes))
	for key, plan in plans.items():
		cfg = plan["config"]
		if cfg.get("node_kind") == "员工" and not cfg.get("employee"): errors.append(f"员工节点 {key} 必须填写员工节点任职工号")
		if cfg.get("node_kind") == "员工" and cfg.get("employee"):
			parent, visited = plan["parent"], set()
			department = None
			while parent in graph and parent not in visited:
				visited.add(parent)
				department = graph[parent]["config"].get("department")
				if department: break
				parent = graph[parent]["parent"]
			person = next(e for e in state["employees"] if e.name == cfg["employee"])
			if not department or person.department not in scopes.get(parent, set()): errors.append(f"员工节点 {key} 的上级部门与花名册不一致")
			cfg["roster_department"] = department
		if cfg.get("proxy_employee") and cfg["proxy_employee"] in {cfg.get("primary_employee"), cfg.get("manager_employee")}:
			errors.append(f"节点 {key} 任职人与代理人不能相同")
	completeness = {
		"nodes": len(plans), "grade_definitions": len(package["职级定义"]),
		"graded_nodes": sum(bool(p["config"].get("chart_grade_code")) for p in plans.values()),
		"source_grade_nodes": sum(bool(p["config"].get("source_grade_tags")) for p in plans.values()),
		"person_references": len(package["人员任职"]),
		"unbound_references": sum(not r["code"] for r in package["人员任职"]),
	}
	if not completeness["grade_definitions"] and not completeness["graded_nodes"]:
		warnings.append("文件未配置图中职级及等级顺序；原表职级标签单独保留，不自动推断高低级别。")
	fingerprint = hashlib.sha256(json.dumps([package, state], ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()
	return {"errors": errors, "warnings": list(dict.fromkeys(warnings)), "fingerprint": fingerprint,
		"completeness": completeness,
		"create_count": len(set(plans) - set(current)), "update_count": len(set(plans) & set(current)),
		"retained_count": len(set(current) - set(plans)), "person_rows": len(package["人员任职"]),
		"grades": [{k: v for k, v in g.items() if k != "_row"} for g in grades.values()],
		"preview": [{"id": key, "name": plans[key]["name"], "parent": plans[key]["parent"],
			"parent_name": plans.get(plans[key]["parent"], {}).get("name") or (current[plans[key]["parent"]].display_name if plans[key]["parent"] in current else "公司"),
			"node_kind": plans[key]["config"].get("node_kind", ""),
			"role": plans[key]["config"].get("role_title") or plans[key]["config"].get("designation", ""),
			"source_grade_tags": plans[key]["config"].get("source_grade_tags", ""),
			"source_grade_status": plans[key]["config"].get("source_grade_status", ""),
			"source_grade_reference": plans[key]["config"].get("source_grade_reference", ""),
			"grade": plans[key]["config"].get("chart_grade", {}).get("label", "")} for key in order if key in plans],
		"_plans": plans, "_current": current, "_order": order, "_state": state}


@frappe.whitelist(methods=["POST"])
def preview_configuration(company: str, file_url: str):
	authorize(company)
	return {k: v for k, v in prepare(company, read_package(file_url)).items() if not k.startswith("_")}


@frappe.whitelist(methods=["POST"])
def import_configuration(company: str, file_url: str, fingerprint: str, auto_sync: int = 1):
	authorize(company)
	# Permission queries can open a repeatable-read snapshot. Begin a fresh
	# transaction before waiting on the same locks used by automatic updates.
	frappe.db.commit()
	return apply_configuration(company, file_url, fingerprint, auto_sync)


def apply_configuration(company, file_url, fingerprint, auto_sync=1):
	# The company lock also serializes first-time imports, before a version exists.
	frappe.db.sql("select name from tabCompany where name=%s for update", company)
	chart = chart_module()
	versions = frappe.db.sql("select name from `tabOrganization Structure Version` where company=%s and source_reference=%s and status!='已归档' for update", (company, chart._manual_organization_reference(company)))
	version = versions[0][0] if versions else None
	plan = prepare(company, read_package(file_url))
	if plan["errors"]: frappe.throw("配置校验未通过：" + "；".join(plan["errors"][:10]))
	if plan["fingerprint"] != fingerprint: frappe.throw("花名册、组织或文件已变化，请重新预览后导入。")
	backup = save_private(workbook_bytes(company), "组织配置导入前备份", company)
	version = version or chart._ensure_manual_organization_version(company)
	ids = {key: node.name for key, node in plan["_current"].items()}
	# Detach only imported nodes first; this permits valid hierarchy reversals.
	for key in plan["_plans"]:
		if key in ids: frappe.db.set_value("Organization Node", ids[key], "parent_node", None, update_modified=False)
	for key in plan["_order"]:
		if key not in plan["_plans"]: continue
		row = plan["_plans"][key]
		doc = frappe.get_doc("Organization Node", ids[key]) if key in ids else frappe.new_doc("Organization Node")
		if doc.is_new(): doc.node_code = "MANUAL-" + frappe.generate_hash(length=10).upper()
		doc.update({"structure_version": version, "display_name": row["name"], "parent_node": ids.get(row["parent"]),
			"node_type": "部门" if row["config"]["node_kind"] in {"室", "课", "组", "线"} else "其他",
			"planned_headcount": row["planned"], "confirmation_status": "已确认",
			"source_text": json.dumps(row["config"], ensure_ascii=False, sort_keys=True)})
		doc.save()
		ids[key] = doc.name
	options = {**plan["_state"]["options"], "chart_grades": plan["grades"], "roster_auto_sync_enabled": bool(int(auto_sync))}
	doc = frappe.get_doc("Organization Structure Version", version)
	doc.notes = json.dumps(options, ensure_ascii=False)
	doc.save()
	return {"created": plan["create_count"], "updated": plan["update_count"], "backup": backup,
		"warnings": plan["warnings"], "auto_sync": bool(int(auto_sync))}
