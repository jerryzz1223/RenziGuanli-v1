"""Persistent salary and contribution registers, independent of payroll months."""
import json
from datetime import timedelta
from io import BytesIO

import frappe
from frappe import _
from frappe.utils import flt, getdate, nowdate

SALARY = 'HRMS Employee Salary Change'
CONTRIBUTION = 'HRMS Employee Contribution Change'
CONTRIBUTION_TYPES = ('社保个人', '社保公司', '公积金个人', '公积金公司')

COMPENSATION_EXPORT_FIELDS = {
	'base_salary': ('底薪', 'salary'),
	'function_allowance': ('职能津贴', 'salary'),
	'certificate_allowance': ('证书津贴', 'salary'),
	'multi_skill_allowance': ('多能工津贴', 'salary'),
	'full_salary': ('薪资小计', 'salary'),
	'social_enabled': ('社保缴纳状态', 'social'),
	'social_personal': ('社保个人承担', 'social'),
	'social_company': ('社保公司承担', 'social'),
	'housing_enabled': ('公积金缴纳状态', 'housing'),
	'housing_personal': ('公积金个人承担', 'housing'),
	'housing_company': ('公积金公司承担', 'housing'),
	'remarks': ('变更原因', 'audit'),
}


def _row_value(row, key, default=None):
	return row.get(key, default) if hasattr(row, 'get') else getattr(row, key, default)


def _export_event_order(event):
	row = event['row']
	return (
		str(_row_value(row, 'effective_date') or ''),
		str(_row_value(row, 'approved_on') or _row_value(row, 'submitted_on') or _row_value(row, 'creation') or ''),
		str(_row_value(row, 'name') or ''),
		event['kind'],
	)


def _apply_compensation_event(state, event):
	row, kind = event['row'], event['kind']
	if kind == 'salary':
		for field in ('base_salary', 'function_allowance', 'certificate_allowance', 'multi_skill_allowance', 'full_salary'):
			state[field] = flt(_row_value(row, field))
		return
	prefix = 'social' if kind == 'social' else 'housing'
	enabled = int(bool(_row_value(row, 'enabled')))
	state[f'{prefix}_enabled'] = enabled
	state[f'{prefix}_personal'] = flt(_row_value(row, 'personal_amount')) if enabled else 0
	state[f'{prefix}_company'] = flt(_row_value(row, 'company_amount')) if enabled else 0


def _build_compensation_export_rows(employees, salary_rows, contribution_rows, start_date, end_date, selected_fields):
	"""Build effective-period snapshots without collapsing equal-valued decisions."""
	start, end = getdate(start_date), getdate(end_date)
	selected_kinds = {COMPENSATION_EXPORT_FIELDS[field][1] for field in selected_fields if field in COMPENSATION_EXPORT_FIELDS}
	selected_kinds.discard('audit')
	events_by_employee = {}
	for row in salary_rows:
		events_by_employee.setdefault(_row_value(row, 'employee'), []).append({'kind': 'salary', 'row': row})
	for row in contribution_rows:
		kind = 'social' if _row_value(row, 'contribution_type') == '社保' else 'housing'
		events_by_employee.setdefault(_row_value(row, 'employee'), []).append({'kind': kind, 'row': row})

	result = []
	for employee in employees:
		employee_id = _row_value(employee, 'name') or _row_value(employee, 'employee')
		events = sorted(events_by_employee.get(employee_id, []), key=_export_event_order)
		state, output_events = {}, []
		for event in events:
			effective = getdate(_row_value(event['row'], 'effective_date'))
			if effective < start:
				_apply_compensation_event(state, event)
				continue
			if effective > end:
				break
			if event['kind'] in selected_kinds:
				output_events.append(event)

		# A standard already in force at the start of the requested period must be
		# represented even when the employee had no changes inside that period.
		if state and not any(getdate(_row_value(event['row'], 'effective_date')) == start for event in output_events):
			output_events.insert(0, {'kind': 'opening', 'row': {'effective_date': start}})

		for index, event in enumerate(output_events):
			if event['kind'] != 'opening':
				_apply_compensation_event(state, event)
			period_start = max(start, getdate(_row_value(event['row'], 'effective_date')))
			next_start = getdate(_row_value(output_events[index + 1]['row'], 'effective_date')) if index + 1 < len(output_events) else None
			period_end = min(end, next_start - timedelta(days=1)) if next_start and next_start > period_start else (period_start if next_start else end)
			result.append({
				'employee': employee_id,
				'employee_name': _row_value(employee, 'employee_name') or employee_id,
				'employee_code': _row_value(employee, 'custom_employee_code') or _row_value(employee, 'employee_code') or employee_id,
				'department': _row_value(employee, 'department_label') or _row_value(employee, 'department') or '',
				'employment_type': _row_value(employee, 'work_nature') or _row_value(employee, 'employment_type') or '',
				'period_start': period_start,
				'period_end': period_end,
				'change_type': {'salary': '定薪', 'social': '社保', 'housing': '公积金', 'opening': '期初沿用'}[event['kind']],
				'remarks': '' if event['kind'] == 'opening' else (_row_value(event['row'], 'remarks') or _row_value(event['row'], 'change_reason') or ''),
				**state,
			})
	return result


def _access(company, approve=False):
	from hrms.payroll.standing_permissions import require_access
	return require_access(company, approve=approve)


@frappe.whitelist()
def permissions(company: str):
	company = _access(company)
	from hrms.payroll.standing_permissions import can_approve
	return {'company': company, 'can_submit': True, 'can_approve': can_approve()}


def _register_key(row, doctype):
	return (row.employee, row.get('contribution_type') if doctype == CONTRIBUTION else '')


def _request_order(row):
	return (str(row.get('submitted_on') or row.get('creation') or ''), str(row.name))


def _latest_action_rows(rows):
	"""Keep the most recently submitted salary/contribution action per employee."""
	latest = {}
	for row in sorted(rows, key=_request_order, reverse=True):
		latest.setdefault(row.employee, row)
	return list(latest.values())


def _attach_actor_names(rows):
	"""Expose user full names while preserving the immutable user ids in audit rows."""
	users = {
		row.get(field)
		for row in rows
		for field in ('submitted_by', 'owner', 'approved_by')
		if row.get(field)
	}
	full_names = {row.name: row.full_name or row.name for row in frappe.get_all(
		'User', filters={'name': ['in', sorted(users or {'__none__'})]},
		fields=['name', 'full_name'], limit_page_length=100000)}
	for row in rows:
		submitter = row.get('submitted_by') or row.get('owner')
		row['submitted_by_name'] = full_names.get(submitter, submitter) or ''
		row['approved_by_name'] = full_names.get(row.get('approved_by'), row.get('approved_by')) or ''
	return rows


def _attach_previous_standards(rows, doctype):
	"""Attach the approved standard that was effective when each request was submitted."""
	groups = {}
	for row in rows:
		groups.setdefault(_register_key(row, doctype), []).append(row)
	for group in groups.values():
		approved = [row for row in group if row.status == '已批准']
		for row in group:
			request_stamp = _request_order(row)[0]
			request_date = request_stamp[:10]
			candidates = [candidate for candidate in approved
				if candidate.name != row.name
				and str(candidate.get('approved_on') or candidate.get('submitted_on') or candidate.get('creation') or '') <= request_stamp
				and str(candidate.get('effective_date') or '') <= request_date]
			row['previous_standard'] = max(candidates, key=lambda candidate: (
				str(candidate.get('effective_date') or ''),
				str(candidate.get('approved_on') or candidate.get('submitted_on') or candidate.get('creation') or ''),
				str(candidate.name),
			)) if candidates else None
	return rows


def _only_change_records(rows, doctype):
	"""Return requests made after the first approved standard for each register."""
	first_approved = {}
	for row in sorted(rows, key=_request_order):
		key = _register_key(row, doctype)
		if row.status == '已批准' and key not in first_approved:
			first_approved[key] = _request_order(row)
	return [row for row in rows
		if _register_key(row, doctype) in first_approved
		and _request_order(row) > first_approved[_register_key(row, doctype)]]


@frappe.whitelist()
def list_change_records(company: str, kind: str):
	company = _access(company)
	if kind not in ('salary', 'contribution'):
		frappe.throw(_('无效的档案类型。'))
	doctype = SALARY if kind == 'salary' else CONTRIBUTION
	filters = {'company': company}
	if doctype == SALARY:
		filters['exclude_from_payroll'] = 0
	rows = frappe.get_all(doctype, filters=filters, fields=['*'],
		order_by='submitted_on desc, creation desc', limit_page_length=10000)
	_attach_previous_standards(rows, doctype)
	rows = _only_change_records(rows, doctype)
	departments = {row.name: row.department_name for row in frappe.get_all('Department',
		filters={'company': company}, fields=['name', 'department_name'], limit_page_length=100000)}
	for row in rows:
		row['decision_doctype'] = doctype
		row['department'] = departments.get(row.department, row.department) or ''
	return _attach_actor_names(rows)


def active_contributions(company, as_of=None, employee=None):
	"""Latest approved revision per employee/type. A stopped revision is explicit zero."""
	filters = {'company': company, 'status': '已批准', 'effective_date': ['<=', as_of or nowdate()]}
	if employee:
		filters['employee'] = employee
	rows = frappe.get_all(CONTRIBUTION,
		filters=filters,
		fields=['*'], order_by='effective_date desc, approved_on desc, creation desc, name desc',
		limit_page_length=100000)
	latest = {}
	for row in rows:
		latest.setdefault((row.employee, row.contribution_type), row)
	return list(latest.values())


@frappe.whitelist()
def list_register(company: str):
	company = _access(company)
	contributions = active_contributions(company)
	registered = {row.employee for row in contributions}
	rows = frappe.get_all('Employee', filters={'company': company},
		fields=['name', 'employee_name', 'custom_employee_code', 'department', 'status'],
		order_by='employee_name asc', limit_page_length=100000)
	departments = {row.name: row.department_name for row in frappe.get_all('Department',
		filters={'company': company}, fields=['name', 'department_name'], limit_page_length=100000)}
	employees = [dict(employee=row.name, employee_name=row.employee_name,
		employee_code=row.custom_employee_code or row.name,
		department=departments.get(row.department, row.department) or '', employee_status=row.status)
		for row in rows if row.status == 'Active' or row.name in registered]
	initialized = frappe.get_all(CONTRIBUTION, filters={'company': company, 'status': ['in', ['待审核', '已批准']]},
		fields=['employee', 'contribution_type'], limit_page_length=100000)
	actions = []
	for doctype in (SALARY, CONTRIBUTION):
		filters = {'company': company}
		if doctype == SALARY:
			filters['exclude_from_payroll'] = 0
		for row in frappe.get_all(doctype, filters=filters,
			fields=['name', 'employee', 'status', 'owner', 'submitted_by', 'submitted_on', 'creation'],
			order_by='submitted_on desc, creation desc, name desc', limit_page_length=100000):
			actions.append(row)
	latest_actions = _attach_actor_names(_latest_action_rows(actions))
	return {'employees': employees, 'contributions': contributions, 'initialized': initialized,
		'latest_actions': latest_actions,
		'pending_employees': sorted({row.employee for row in actions if row.status == '待审核'})}


@frappe.whitelist()
def initial_salary_rows(company: str):
	company = _access(company)
	from hrms.api.payroll_input import list_employee_salary_change_grid
	rows = list_employee_salary_change_grid(company, payroll_month='', page_length=100000)['rows']
	initialized = set(frappe.get_all(SALARY, filters={'company': company,
		'exclude_from_payroll': 0, 'status': ['in', ['待审核', '已批准']]}, pluck='employee', limit_page_length=100000))
	return {'rows': [row for row in rows if row['employee'] not in initialized]}


@frappe.whitelist()
def submit_contribution(company: str, employee: str, contribution_type: str, effective_date: str, personal_amount: float = 0,
	company_amount: float = 0, enabled: int = 1, remarks: str = '', request_mode: str = 'change'):
	company = _access(company)
	if request_mode == 'initial' and frappe.db.exists(CONTRIBUTION, {
		'company': company, 'employee': employee, 'contribution_type': contribution_type,
		'status': ['in', ['待审核', '已批准']]}):
		frappe.throw(_('该缴费项目已录入或正在审批，请通过申请修改处理。'))
	if not str(remarks).strip():
		frappe.throw(_('请填写修改原因。'))
	if not effective_date:
		frappe.throw(_('请填写生效日期。'))
	context = frappe.db.get_value('Employee', employee,
		['company', 'employee_name', 'custom_employee_code', 'department'], as_dict=True) or {}
	if context.get('company') != company:
		frappe.throw(_('员工不属于当前公司。'))
	values = dict(company=company, employee=employee, employee_code=context.get('custom_employee_code') or employee,
		employee_name=context.get('employee_name'), department=context.get('department'),
		contribution_type=contribution_type, effective_date=getdate(effective_date),
		personal_amount=flt(personal_amount), company_amount=flt(company_amount),
		enabled=int(bool(int(enabled))), remarks=remarks)
	# Exact retries of a pending request reuse the request; changed payloads are new revisions.
	existing = frappe.db.get_value(CONTRIBUTION,
		{**values, 'status': '待审核', 'submitted_by': frappe.session.user}, 'name')
	if existing:
		return existing
	# Keep one pending decision per employee, contribution type and effective
	# date.  The Employee-row lock also prevents concurrent submissions from
	# both observing an empty pending slot.
	frappe.db.sql('SELECT name FROM `tabEmployee` WHERE name=%s FOR UPDATE', (employee,))
	pending = frappe.db.get_value(CONTRIBUTION, {
		'company': company, 'employee': employee, 'contribution_type': contribution_type,
		'effective_date': values['effective_date'], 'status': '待审核',
	}, 'name')
	if pending:
		frappe.throw(_('该员工在此生效日期已有待审批的缴费申请，请先等待审批结果或驳回后再提交。'))
	doc = frappe.get_doc({'doctype': CONTRIBUTION, **values}).insert(ignore_permissions=True)
	return doc.name


@frappe.whitelist()
def list_pending(company: str):
	company = _access(company, approve=True)
	result = []
	for doctype in (SALARY, CONTRIBUTION):
		filters = {'company': company, 'status': '待审核'}
		if doctype == SALARY:
			filters['exclude_from_payroll'] = 0
		pending = frappe.get_all(doctype, filters=filters,
			fields=['*'], order_by='creation desc', limit_page_length=100000)
		# Compare against the approved standard in force when the reviewer opens
		# the page. Pending proposals and future approvals are not current pay.
		current = {}
		if pending:
			approved_filters = {**filters, 'status': '已批准', 'effective_date': ['<=', nowdate()],
				'employee': ['in', list({row.employee for row in pending})]}
			for approved in frappe.get_all(doctype, filters=approved_filters, fields=['*'],
				order_by='effective_date desc, approved_on desc, creation desc, name desc', limit_page_length=100000):
				current.setdefault((approved.employee, approved.get('contribution_type')), approved)
		for row in pending:
			row['decision_doctype'] = doctype
			row['previous_standard'] = current.get((row.employee, row.get('contribution_type')))
			result.append(row)
	return result


@frappe.whitelist()
def employee_history(company: str, employee: str):
	company = _access(company)
	if frappe.db.get_value('Employee', employee, 'company') != company:
		frappe.throw(_('员工不属于当前公司。'))
	result = []
	for doctype in (SALARY, CONTRIBUTION):
		filters = {'company': company, 'employee': employee}
		if doctype == SALARY:
			filters['exclude_from_payroll'] = 0
		for row in frappe.get_all(doctype, filters=filters,
			fields=['*'], order_by='effective_date desc, creation desc', limit_page_length=100000):
			row['decision_doctype'] = doctype
			result.append(row)
	result = sorted(result, key=lambda row: str(row.get('submitted_on') or row.get('creation') or ''), reverse=True)
	return _attach_actor_names(result)


def _selected_compensation_export_fields(selected_fields):
	if isinstance(selected_fields, str):
		try:
			selected_fields = json.loads(selected_fields)
		except (TypeError, ValueError):
			selected_fields = []
	selected = []
	for field in selected_fields or []:
		if field in COMPENSATION_EXPORT_FIELDS and field not in selected:
			selected.append(field)
	if not selected or not any(COMPENSATION_EXPORT_FIELDS[field][1] != 'audit' for field in selected):
		frappe.throw(_('请至少选择一项薪资、社保或公积金内容。'))
	return selected


@frappe.whitelist()
def export_compensation_register(company: str, start_date: str, end_date: str, selected_fields: str | None = None,
	department: str = '', employee: str = ''):
	"""Export approved effective-dated standards as a real XLSX workbook."""
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter
	from hrms.api.payroll_input import _safe_fields, _salary_contribution_defaults
	from hrms.utils.export_watermark import save_workbook_with_logo_watermark

	company = _access(company)
	start, end = getdate(start_date), getdate(end_date)
	if end < start:
		frappe.throw(_('结束日期不能早于开始日期。'))
	selected = _selected_compensation_export_fields(selected_fields)
	employee_fields = _safe_fields('Employee', [
		'name', 'employee_name', 'custom_employee_code', 'employee_number', 'department',
		'employment_type', 'status', 'custom_is_confirmed', 'final_confirmation_date', 'confirmation_date', 'company',
	])
	employee_filters = {'company': company}
	if employee:
		employee_filters['name'] = employee
	employees = frappe.get_all('Employee', filters=employee_filters, fields=employee_fields,
		order_by='employee_name asc, name asc', limit_page_length=100000)

	department_ids = {row.get('department') for row in employees if row.get('department')}
	# Salary operators can read the compensation register without having the
	# standalone Department doctype permission. Access was already checked above;
	# read only the two labels needed to reproduce the authorized register filter.
	department_labels = {row.name: row.department_name for row in frappe.db.get_all('Department',
		filters={'name': ['in', sorted(department_ids or {'__none__'})]},
		fields=['name', 'department_name'], limit_page_length=100000)}
	for row in employees:
		row['department_label'] = department_labels.get(row.get('department'), row.get('department')) or ''
		defaults = _salary_contribution_defaults(row, str(end))
		row['work_nature'] = _('在职·{0}').format(defaults['employment_stage'])
	if department:
		employees = [row for row in employees if department in (row.get('department'), row.get('department_label'))]
	if not employees:
		frappe.throw(_('当前筛选条件下没有员工。'))

	employee_ids = [row.name for row in employees]
	salary_rows = frappe.get_all(SALARY, filters={
		'company': company, 'employee': ['in', employee_ids], 'status': '已批准',
		'exclude_from_payroll': 0, 'effective_date': ['<=', end],
	}, fields=['*'], order_by='effective_date asc, approved_on asc, creation asc, name asc', limit_page_length=100000)
	contribution_rows = frappe.get_all(CONTRIBUTION, filters={
		'company': company, 'employee': ['in', employee_ids], 'status': '已批准', 'effective_date': ['<=', end],
	}, fields=['*'], order_by='effective_date asc, approved_on asc, creation asc, name asc', limit_page_length=100000)
	rows = _build_compensation_export_rows(employees, salary_rows, contribution_rows, start, end, selected)
	if not rows:
		frappe.throw(_('所选日期和内容范围内没有已批准且生效的工资社保记录。'))

	fixed_columns = [
		('姓名', 'employee_name'), ('工号', 'employee_code'), ('工作性质', 'employment_type'), ('部门', 'department'),
		('生效开始', 'period_start'), ('生效结束', 'period_end'), ('记录类型', 'change_type'),
	]
	columns = fixed_columns + [(COMPENSATION_EXPORT_FIELDS[field][0], field) for field in selected]
	workbook = Workbook()
	sheet = workbook.active
	sheet.title = '工资社保历史'
	sheet.append([label for label, _field in columns])
	header_fill = PatternFill('solid', fgColor='DDEBF7')
	thin = Side(style='thin', color='B7C9D6')
	border = Border(left=thin, right=thin, top=thin, bottom=thin)
	for cell in sheet[1]:
		cell.fill = header_fill
		cell.font = Font(name='Microsoft YaHei', size=10, bold=True)
		cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
		cell.border = border
	for row in rows:
		values = []
		for _label, field in columns:
			value = row.get(field, '')
			if field in ('social_enabled', 'housing_enabled'):
				value = '' if field not in row else ('缴纳' if value else '停缴')
			values.append(value)
		sheet.append(values)
		for column_index, (_label, field) in enumerate(columns, start=1):
			cell = sheet.cell(sheet.max_row, column_index)
			cell.border = border
			cell.alignment = Alignment(vertical='center', wrap_text=True)
			if field in ('period_start', 'period_end'):
				cell.number_format = 'yyyy-mm-dd'
			elif field in COMPENSATION_EXPORT_FIELDS and COMPENSATION_EXPORT_FIELDS[field][1] in ('salary', 'social', 'housing') and not field.endswith('_enabled'):
				cell.number_format = '#,##0.00'
	widths = {'employee_name': 14, 'employee_code': 14, 'employment_type': 14, 'department': 18,
		'period_start': 13, 'period_end': 13, 'change_type': 12, 'remarks': 28}
	for index, (_label, field) in enumerate(columns, start=1):
		sheet.column_dimensions[get_column_letter(index)].width = widths.get(field, 16)
	sheet.freeze_panes = 'A2'
	sheet.auto_filter.ref = f'A1:{get_column_letter(len(columns))}{sheet.max_row}'
	sheet.sheet_view.showGridLines = False

	output = BytesIO()
	save_workbook_with_logo_watermark(workbook, output)
	filename = f'工资社保历史_{start:%Y%m%d}-{end:%Y%m%d}.xlsx'
	file_doc = frappe.get_doc({'doctype': 'File', 'file_name': filename, 'content': output.getvalue(), 'is_private': 1}).insert(ignore_permissions=True)
	return {'file_url': file_doc.file_url, 'file_name': filename, 'row_count': len(rows)}


@frappe.whitelist()
def review_decision(company: str, decision_doctype: str, name: str, decision: str, review_note: str):
	company = _access(company, approve=True)
	if decision_doctype not in (SALARY, CONTRIBUTION) or decision not in ('已批准', '已驳回'):
		frappe.throw(_('无效的审批操作。'))
	# Serialize competing reviewers; the document controller checks status and self-approval.
	frappe.db.sql(f'SELECT name FROM `tab{decision_doctype}` WHERE name=%s FOR UPDATE', (name,))
	doc = frappe.get_doc(decision_doctype, name)
	if doc.company != company:
		frappe.throw(_('不能审批其他公司的记录。'))
	doc.status, doc.review_note = decision, review_note
	doc.save(ignore_permissions=True)
	return {'name': doc.name, 'status': doc.status, 'approved_by': doc.approved_by, 'approved_on': doc.approved_on}


@frappe.whitelist()
def import_contribution_workbook(company: str, file_url: str, effective_date: str, preview: int = 1, request_mode: str = 'change'):
	"""Reuse the existing source parsers; the independent register owns the approval."""
	company = _access(company)
	from hrms.api.payroll_input import (_load_workbook, _raw_payroll_source_kind, _raw_payroll_source_rows,
		_first, _payroll_employee_lookup, _raw_variable_entries)
	if not effective_date:
		frappe.throw(_('请选择本次变更的生效日期。'))
	workbook = _load_workbook(file_url)
	rows, errors, seen = [], [], set()
	for sheet in workbook:
		kind = _raw_payroll_source_kind(sheet)
		if kind not in ('social_insurance', 'housing_fund'):
			continue
		for index, raw in enumerate(_raw_payroll_source_rows(sheet, kind), 1):
			code, name = _first(raw, '工号', '员工编号'), _first(raw, '姓名')
			if not code and not name:
				continue
			try:
				employee = _payroll_employee_lookup(code, name)
				if not employee or frappe.db.get_value('Employee', employee, 'company') != company:
					raise ValueError('工号、姓名或公司不匹配')
				if request_mode == 'initial' and frappe.db.exists(CONTRIBUTION, {'company': company, 'employee': employee,
					'contribution_type': '社保' if kind == 'social_insurance' else '公积金', 'status': ['in', ['待审核', '已批准']]}):
					raise ValueError('已录入或正在审批，请通过申请修改处理')
				key = (employee, kind)
				if key in seen:
					raise ValueError('同一员工同一缴费类型重复，请先合并源文件')
				seen.add(key)
				amounts = _raw_variable_entries(kind, raw)
				if any(amount < 0 for _, amount in amounts):
					raise ValueError('缴费金额不能为负数')
				rows.append(dict(employee=employee, employee_name=name or employee,
					contribution_type='社保' if kind == 'social_insurance' else '公积金',
					personal_amount=amounts[0][1], company_amount=amounts[1][1]))
			except Exception as exc:
				errors.append(f'{sheet.title} / {index} / {code or name}: {exc}')
	if int(preview):
		return {'rows': rows, 'errors': errors}
	if errors or not rows:
		frappe.throw(_('请先修正预览中的异常；文件必须包含社保或公积金来源表。'))
	for row in rows:
		submit_contribution(company, row['employee'], row['contribution_type'], effective_date,
			row['personal_amount'], row['company_amount'], remarks=f'Excel 导入：{file_url}', request_mode=request_mode)
	return {'submitted': len(rows)}
