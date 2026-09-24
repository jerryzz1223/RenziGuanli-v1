# Copyright (c) 2026, Frappe Technologies Pvt. Ltd. and contributors
# For license information, please see license.txt

"""Guided write APIs for the manual training-learning workflow.

The workbook importer remains the audit-preserving bulk path.  These endpoints
cover day-to-day courses that are created, delivered and recorded directly in
HRMS, while resolving every participant back to the current company's durable
employee code.
"""

from io import BytesIO

import frappe
from frappe import _
from frappe.utils import cint, flt, get_datetime

from hrms.access_control import require_hrms_capability


def _values(payload):
	if isinstance(payload, str):
		payload = frappe.parse_json(payload)
	if not isinstance(payload, dict):
		frappe.throw(_("培训数据格式不正确。"))
	return payload


def _company(company):
	company = str(company or "").strip()
	if not company:
		frappe.throw(_("请先选择当前公司。"))
	if not frappe.has_permission("Company", "read", company):
		frappe.throw(_("无权使用该公司的培训数据。"), frappe.PermissionError)
	return company


def _text(value):
	return str(value or "").strip()


def _require(value, label):
	value = _text(value)
	if not value:
		frappe.throw(_("请填写{0}。").format(label))
	return value


def _optional_number(value, label, minimum=0, maximum=None):
	"""Validate an optional numeric draft value without turning a blank into zero."""
	if value in (None, ""):
		return None
	try:
		number = float(value)
	except (TypeError, ValueError):
		frappe.throw(_("{0}必须是数字。").format(label))
	if number < minimum or (maximum is not None and number > maximum):
		if maximum is None:
			frappe.throw(_("{0}不能小于 {1}。").format(label, minimum))
		frappe.throw(_("{0}必须介于 {1} 到 {2}。").format(label, minimum, maximum))
	return number


def _employee_fields():
	if not frappe.db.has_column("Employee", "custom_employee_code"):
		frappe.throw(_("员工主档尚未配置公司工号字段，不能登记培训人员。"))
	return ["name", "custom_employee_code", "employee_name", "department", "status", "company"]


def _employee_rows(company, participants, allow_empty=False):
	"""Resolve UI selections and re-check company code/company on the server."""
	if isinstance(participants, str):
		participants = frappe.parse_json(participants)
	if not isinstance(participants, list) or (not participants and not allow_empty):
		frappe.throw(_("请至少选择一名参训员工。"))

	resolved = []
	seen = set()
	for participant in participants:
		participant = participant or {}
		employee = _text(participant.get("employee"))
		employee_code = _text(participant.get("employee_code"))
		filters = {"company": company}
		if employee:
			filters["name"] = employee
		elif employee_code:
			filters["custom_employee_code"] = employee_code
		else:
			frappe.throw(_("参训人员中存在未选择员工的空行。"))
		row = frappe.db.get_value("Employee", filters, _employee_fields(), as_dict=True)
		if not row:
			frappe.throw(_("员工 {0} 不属于当前公司或公司工号无效。").format(employee_code or employee))
		code = _text(row.custom_employee_code)
		if not code:
			frappe.throw(_("员工 {0} 尚未维护公司工号，不能写入培训档案。").format(row.employee_name or row.name))
		if frappe.db.count("Employee", {"company": company, "custom_employee_code": code}) != 1:
			frappe.throw(_("公司工号 {0} 在当前公司不是唯一值，请先修正员工主档。").format(code))
		if code in seen:
			frappe.throw(_("公司工号 {0} 在参训名单中重复。").format(code))
		seen.add(code)
		resolved.append({**participant, "employee": row.name, "employee_code": code, "employee_name": row.employee_name, "department": row.department})
	return resolved


def _draft_participant_values(row):
	hours = _optional_number(row.get("hours"), _("{0} 的学时").format(row["employee_code"]))
	score = _optional_number(row.get("score"), _("{0} 的成绩").format(row["employee_code"]), maximum=100)
	attendance = row.get("attendance") or "Present"
	if attendance not in ("Present", "Absent"):
		frappe.throw(_("工号 {0} 的出席状态无效。").format(row["employee_code"]))
	return {
		"employee": row["employee"], "employee_code": row["employee_code"],
		"employee_name": row["employee_name"], "department": row["department"],
		"status": "Invited", "attendance": attendance,
		"is_mandatory": cint(row.get("is_mandatory")),
		"draft_hours": "" if hours is None else str(hours),
		"draft_score": "" if score is None else str(score),
		"draft_grade": _text(row.get("grade")),
		"draft_needs_retraining": cint(row.get("needs_retraining")),
		"draft_comments": _text(row.get("comments")),
	}


def _unique_event_name(course, start_time):
	stamp = get_datetime(start_time).strftime("%Y-%m-%d %H:%M")
	base = f"{course} · {stamp}"
	name = base
	index = 2
	while frappe.db.exists("Training Event", name):
		name = f"{base} · {index}"
		index += 1
	return name


ROSTER_TEMPLATE_HEADERS = ["公司工号", "姓名", "学时", "成绩", "出席状态", "等级/结果", "是否补训", "备注"]


def _roster_file_content(file_url):
	file_name = frappe.db.get_value("File", {"file_url": file_url}, "name")
	if not file_name:
		frappe.throw(_("未找到上传的参训名单文件。"))
	file_doc = frappe.get_doc("File", file_name)
	file_doc.check_permission("read")
	content = file_doc.get_content()
	return content.encode() if isinstance(content, str) else content


def _training_roster_rows(content):
	"""Read the basic roster template without coupling matching to spreadsheet row order."""
	from openpyxl import load_workbook

	workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
	worksheet = workbook["参训名单"] if "参训名单" in workbook.sheetnames else workbook.active
	values = list(worksheet.iter_rows(values_only=True))
	if not values:
		frappe.throw(_("参训名单为空。"))
	header_index = next(
		(index for index, row in enumerate(values[:10]) if "公司工号" in row or "姓名" in row),
		None,
	)
	if header_index is None:
		frappe.throw(_("未找到表头，请使用参训名单基础模板。"))
	headers = [_text(value) for value in values[header_index]]
	missing = [header for header in ("公司工号", "姓名") if header not in headers]
	if missing:
		frappe.throw(_("模板缺少字段：{0}").format("、".join(missing)))
	rows = []
	for row_number, values_row in enumerate(values[header_index + 1 :], start=header_index + 2):
		item = {header: values_row[index] if index < len(values_row) else None for index, header in enumerate(headers) if header}
		if not any(item.get(header) not in (None, "") for header in ROSTER_TEMPLATE_HEADERS):
			continue
		rows.append({"row_number": row_number, **item})
	return rows


def _import_attendance(value):
	value = _text(value)
	if value in ("", "出席", "Present"):
		return "Present"
	if value in ("缺席", "Absent"):
		return "Absent"
	return ""


@frappe.whitelist()
def download_training_roster_template():
	"""Download a minimal Excel template; the instruction sheet is not imported."""
	require_hrms_capability("training_submit", legacy_roles=("HR Manager",))
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill

	workbook = Workbook()
	worksheet = workbook.active
	worksheet.title = "参训名单"
	worksheet.append(ROSTER_TEMPLATE_HEADERS)
	for cell in worksheet[1]:
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = PatternFill("solid", fgColor="4472C4")
	for column, width in enumerate((16, 18, 12, 12, 14, 16, 14, 28), start=1):
		worksheet.column_dimensions[chr(64 + column)].width = width
	instructions = workbook.create_sheet("填写说明")
	instructions.append(["填写说明"])
	instructions.append(["公司工号或姓名至少填写一项；推荐填写公司工号。姓名存在重名时不会自动匹配。"])
	instructions.append(["学时、成绩、等级/结果、是否补训、备注均可留空，保存上课后仍可继续补录。"])
	instructions.append(["出席状态可填写：出席、缺席；是否补训可填写：是、否。"])
	output = BytesIO()
	workbook.save(output)
	frappe.response["filename"] = "参训员工导入基础模板.xlsx"
	frappe.response["filecontent"] = output.getvalue()
	frappe.response["type"] = "binary"


@frappe.whitelist()
def preview_training_roster_import(file_url: str, company: str):
	"""Match one-event roster rows by company code first, then by an unambiguous exact name."""
	require_hrms_capability("training_submit", legacy_roles=("HR Manager",))
	if not frappe.has_permission("Employee", "read"):
		frappe.throw(_("无权读取员工主档。"), frappe.PermissionError)
	company = _company(company)
	rows = _training_roster_rows(_roster_file_content(file_url))
	employees = frappe.get_all(
		"Employee",
		filters={"company": company, "status": "Active"},
		fields=_employee_fields(),
		limit_page_length=0,
	)
	by_code = {}
	by_name = {}
	for employee in employees:
		code = _text(employee.custom_employee_code)
		if code:
			by_code.setdefault(code, []).append(employee)
			by_name.setdefault(_text(employee.employee_name), []).append(employee)

	participants = []
	issues = []
	seen = set()
	for source in rows:
		code = _text(source.get("公司工号"))
		name = _text(source.get("姓名"))
		matches = by_code.get(code, []) if code else by_name.get(name, [])
		reason = ""
		if not code and not name:
			reason = _("公司工号和姓名不能同时为空")
		elif not matches:
			reason = _("未找到在职员工")
		elif len(matches) > 1:
			reason = _("匹配到多名员工，请填写唯一公司工号")
		elif code and name and _text(matches[0].employee_name) != name:
			reason = _("公司工号与姓名不一致")
		elif _text(matches[0].custom_employee_code) in seen:
			reason = _("名单中公司工号重复")
		attendance = _import_attendance(source.get("出席状态"))
		if not reason and not attendance:
			reason = _("出席状态只能填写出席或缺席")
		try:
			hours = _optional_number(source.get("学时"), _("第 {0} 行学时").format(source["row_number"]))
			score = _optional_number(source.get("成绩"), _("第 {0} 行成绩").format(source["row_number"]), maximum=100)
		except Exception as error:
			hours = score = None
			reason = _text(error)
		if reason:
			issues.append({"row_number": source["row_number"], "employee_code": code, "employee_name": name, "reason": reason})
			continue
		employee = matches[0]
		resolved_code = _text(employee.custom_employee_code)
		seen.add(resolved_code)
		participants.append(
			{
				"employee": employee.name, "employee_code": resolved_code,
				"employee_name": employee.employee_name, "department": employee.department,
				"attendance": attendance, "hours": "" if hours is None else hours,
				"score": "" if score is None else score, "grade": _text(source.get("等级/结果")),
				"needs_retraining": cint(_text(source.get("是否补训")) in ("是", "1", "Yes", "yes")),
				"comments": _text(source.get("备注")), "match_basis": "company_code" if code else "unique_name",
			}
		)
	return {"participants": participants, "issues": issues, "row_count": len(rows), "matched_count": len(participants)}


@frappe.whitelist()
def search_training_employees(
	doctype: str, txt: str, searchfield: str, start: int, page_len: int, filters: str | dict
):
	"""Link-query used by participant grids; results lead with company code."""
	if not frappe.has_permission("Employee", "read"):
		frappe.throw(_("无权读取员工主档。"), frappe.PermissionError)
	if isinstance(filters, str):
		filters = frappe.parse_json(filters)
	company = _company((filters or {}).get("company"))
	query = f"%{_text(txt)}%"
	limit = min(max(cint(page_len) or 20, 1), 100)
	return frappe.db.sql(
		"""
			select employee.name, employee.custom_employee_code, employee.employee_name, employee.department
			from `tabEmployee` employee
			where employee.company = %(company)s
				and employee.status = 'Active'
				and coalesce(employee.custom_employee_code, '') != ''
				and (
					employee.custom_employee_code like %(query)s
					or employee.employee_name like %(query)s
					or employee.department like %(query)s
				)
			order by employee.custom_employee_code asc
			limit %(start)s, %(limit)s
		""",
		{"company": company, "query": query, "start": max(cint(start), 0), "limit": limit},
	)


@frappe.whitelist()
def find_training_employees(company: str, query: str = "", limit: int = 20):
	"""Company-code-first search for the custom training roster editor."""
	if not frappe.has_permission("Employee", "read"):
		frappe.throw(_("无权读取员工主档。"), frappe.PermissionError)
	company = _company(company)
	search = f"%{_text(query)}%"
	return frappe.db.sql(
		"""
			select employee.name as employee, employee.custom_employee_code as employee_code,
				employee.employee_name, employee.department
			from `tabEmployee` employee
			where employee.company = %(company)s and employee.status = 'Active'
				and coalesce(employee.custom_employee_code, '') != ''
				and (employee.custom_employee_code like %(query)s
					or employee.employee_name like %(query)s or employee.department like %(query)s)
			order by case when employee.custom_employee_code = %(exact)s then 0 else 1 end,
				employee.custom_employee_code asc
			limit %(limit)s
		""",
		{"company": company, "query": search, "exact": _text(query), "limit": min(max(cint(limit), 1), 50)},
		as_dict=True,
	)


@frappe.whitelist()
def create_training_course(payload: str):
	require_hrms_capability("training_submit", legacy_roles=("HR Manager",))
	data = _values(payload)
	company = _company(data.get("company"))
	course = _require(data.get("course_name"), _("课程名称"))
	period = _text(data.get("plan_period"))
	name = course
	if frappe.db.exists("Training Program", name) and period:
		name = f"{course}（{period}）"
	if frappe.db.exists("Training Program", name):
		frappe.throw(_("课程 {0} 已存在，请直接安排上课或使用不同的计划期间。").format(name))

	doc = frappe.get_doc(
		{
			"doctype": "Training Program",
			"training_program": name,
			"company": company,
			"owner_department": data.get("owner_department"),
			"plan_period": period,
			"planned_month": data.get("planned_month"),
			"planned_hours": flt(data.get("planned_hours")) or None,
			"planned_location": data.get("planned_location"),
			"target_audience": data.get("target_audience"),
			"approval_status": "Draft",
			"status": "Scheduled",
			"training_category": data.get("training_category"),
			"training_mode": data.get("training_mode"),
			"is_mandatory": cint(data.get("is_mandatory")),
			"objective": data.get("objective"),
			"trainer_name": data.get("trainer_name"),
			"description": _require(data.get("description") or data.get("objective"), _("课程说明或培训目标")),
		}
	)
	doc.insert()
	return {"name": doc.name, "course_name": course, "approval_status": doc.approval_status}


@frappe.whitelist()
def create_training_activity(payload: str):
	"""Create a scheduled draft; actual attendees are confirmed after class."""
	require_hrms_capability("training_submit", legacy_roles=("HR Manager",))
	data = _values(payload)
	company = _company(data.get("company"))
	program_name = _require(data.get("training_program"), _("计划课程"))
	program = frappe.get_doc("Training Program", program_name)
	if program.company != company:
		frappe.throw(_("计划课程不属于当前公司。"))

	start_time = get_datetime(_require(data.get("start_time"), _("开始时间")))
	end_time = get_datetime(_require(data.get("end_time"), _("结束时间")))
	if end_time <= start_time:
		frappe.throw(_("结束时间必须晚于开始时间。"))
	course = _text(data.get("course")) or _text(program.get("source_content")) or program.training_program
	participants = data.get("participants") or []
	resolved = _employee_rows(company, participants) if participants else []
	duration = round((end_time - start_time).total_seconds() / 3600, 2)
	assessment_required = cint(data.get("assessment_required"))
	passing_score = flt(data.get("passing_score"))
	if assessment_required and not 0 < passing_score <= 100:
		frappe.throw(_("需要考核的课程必须设置 1 到 100 之间的合格分数。"))
	doc = frappe.get_doc(
		{
			"doctype": "Training Event",
			"event_name": _unique_event_name(course, start_time),
			"training_program": program.name,
			"event_status": "Scheduled",
			"type": data.get("type") or "Theory",
			"company": company,
			"owner_department": data.get("owner_department") or program.owner_department,
			"training_category": data.get("training_category") or program.training_category,
			"training_mode": data.get("training_mode") or program.training_mode,
			"course": course,
			"course_hours": flt(data.get("course_hours")) or duration,
			"delivery_method": data.get("delivery_method"),
			"target_audience": data.get("target_audience") or program.target_audience,
			"location": _require(data.get("location") or program.planned_location, _("上课地点")),
			"start_time": start_time,
			"end_time": end_time,
			"trainer_name": data.get("trainer_name") or program.trainer_name,
			"introduction": data.get("introduction") or program.description,
			"assessment_required": assessment_required,
			"passing_score": passing_score or None,
			"retraining_due_on": data.get("retraining_due_on"),
			"qualification_gate": data.get("qualification_gate"),
			"plan_match_status": "已匹配计划",
			"plan_match_basis": "人工创建并关联计划",
			"plan_match_score": 100,
			"employees": [_draft_participant_values(row) for row in resolved],
		}
	)
	doc.insert()
	if program.approval_status == "Approved":
		program.db_set("approval_status", "In Progress")
	return {"name": doc.name, "event_name": doc.event_name, "participant_count": len(resolved), "docstatus": doc.docstatus}


@frappe.whitelist()
def get_training_activity_roster(event_name: str):
	event = frappe.get_doc("Training Event", event_name)
	if not frappe.has_permission("Training Event", "read", doc=event):
		frappe.throw(_("无权读取该培训活动。"), frappe.PermissionError)
	rows = []
	for row in event.employees:
		draft_hours = _optional_number(row.draft_hours, _("{0} 的学时").format(row.employee_code))
		draft_score = _optional_number(row.draft_score, _("{0} 的成绩").format(row.employee_code), maximum=100)
		rows.append(
			{
				"employee": row.employee, "employee_code": row.employee_code,
				"employee_name": row.employee_name, "department": row.department,
				"attendance": row.attendance or "Present", "hours": draft_hours,
				"score": draft_score, "grade": row.draft_grade or "",
				"needs_retraining": cint(row.draft_needs_retraining), "comments": row.draft_comments or "",
			}
		)
	return {
		"name": event.name, "company": event.company, "course": event.course or event.event_name,
		"course_hours": event.course_hours, "assessment_required": event.assessment_required,
		"passing_score": event.passing_score, "docstatus": event.docstatus,
		"has_submitted_result": bool(frappe.db.exists("Training Result", {"training_event": event.name, "docstatus": 1})),
		"participants": rows,
	}


@frappe.whitelist()
def save_training_activity_roster(payload: str):
	"""Save a draft roster without completing the event or writing employee history."""
	require_hrms_capability("training_submit", legacy_roles=("HR Manager",))
	data = _values(payload)
	event = frappe.get_doc("Training Event", _require(data.get("training_event"), _("培训活动")))
	event.check_permission("write")
	if event.docstatus != 0:
		frappe.throw(_("该培训活动已提交，只能查看已归档结果。"))
	participants = _employee_rows(_company(event.company), data.get("participants") or [], allow_empty=True)
	event.set("employees", [_draft_participant_values(row) for row in participants])
	event.save()
	return {"training_event": event.name, "participant_count": len(participants)}


@frappe.whitelist()
def record_training_completion(payload: str):
	"""Submit the event and result, then let Training Result update employee history."""
	require_hrms_capability("training_submit", legacy_roles=("HR Manager",))
	data = _values(payload)
	event = frappe.get_doc("Training Event", _require(data.get("training_event"), _("培训活动")))
	event.check_permission("submit" if event.docstatus == 0 else "write")
	company = _company(event.company)
	if event.event_status == "Cancelled" or event.docstatus == 2:
		frappe.throw(_("已取消的培训活动不能登记课后结果。"))
	if frappe.db.exists("Training Result", {"training_event": event.name, "docstatus": 1}):
		frappe.throw(_("该培训活动已有已提交结果，不能重复登记。"))

	participants = _employee_rows(company, data.get("participants"))
	default_hours = flt(event.course_hours) or round((get_datetime(event.end_time) - get_datetime(event.start_time)).total_seconds() / 3600, 2)
	if event.assessment_required:
		if not 0 < flt(event.passing_score) <= 100:
			frappe.throw(_("该培训活动需要考核，但尚未设置有效的合格分数。"))
		missing_scores = [row["employee_code"] for row in participants if row.get("attendance", "Present") != "Absent" and row.get("score") in (None, "")]
		if missing_scores:
			frappe.throw(_("需要考核的课程必须填写成绩；未填写工号：{0}").format("、".join(missing_scores)))

	event_rows = []
	result_rows = []
	for row in participants:
		attendance = row.get("attendance") or "Present"
		if attendance not in ("Present", "Absent"):
			frappe.throw(_("工号 {0} 的出席状态无效。").format(row["employee_code"]))
		hours = 0 if attendance == "Absent" else flt(row.get("hours") if row.get("hours") not in (None, "") else default_hours)
		if hours < 0:
			frappe.throw(_("工号 {0} 的课时不能小于 0。").format(row["employee_code"]))
		score = row.get("score")
		if score not in (None, "") and not 0 <= flt(score) <= 100:
			frappe.throw(_("工号 {0} 的分数必须介于 0 到 100。").format(row["employee_code"]))
		assessment = "Absent" if attendance == "Absent" else "Pending" if event.assessment_required else "Pass"
		event_rows.append({
			"employee": row["employee"], "employee_code": row["employee_code"],
			"employee_name": row["employee_name"], "department": row["department"],
			"attendance": attendance, "status": "Open" if attendance == "Absent" else "Completed",
			"draft_hours": str(hours), "draft_score": "" if score in (None, "") else str(flt(score)),
			"draft_grade": _text(row.get("grade")),
			"draft_needs_retraining": cint(row.get("needs_retraining")),
			"draft_comments": _text(row.get("comments")),
		})
		result_rows.append({
			"employee": row["employee"], "employee_code": row["employee_code"],
			"employee_name": row["employee_name"], "department": row["department"],
			"hours": hours, "score": flt(score) if score not in (None, "") else None,
			"grade": row.get("grade"), "assessment_result": assessment,
			"needs_retraining": cint(row.get("needs_retraining")) or attendance == "Absent",
			"comments": row.get("comments"),
		})

	if event.docstatus == 0:
		event.set("employees", event_rows)
		event.submit()
	elif event.docstatus != 1:
		frappe.throw(_("培训活动状态不允许登记结果。"))

	draft_name = frappe.db.get_value("Training Result", {"training_event": event.name, "docstatus": 0}, "name")
	if draft_name:
		result = frappe.get_doc("Training Result", draft_name)
		result.set("employees", result_rows)
		result.save()
	else:
		result = frappe.get_doc({"doctype": "Training Result", "training_event": event.name, "employees": result_rows})
		result.insert()
	result.submit()
	return {
		"training_event": event.name, "training_result": result.name,
		"participant_count": len(result_rows),
		"passed_count": sum(row.assessment_result == "Pass" for row in result.employees),
		"retraining_count": sum(cint(row.needs_retraining) for row in result.employees),
	}
