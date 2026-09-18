"""Announcement workflow, directory data and blank Word/Excel templates."""

import html
import io
import json
import re
import zipfile
from datetime import date, datetime, time
from xml.etree import ElementTree
from typing import Any

import frappe
from frappe import _
from frappe.utils import get_datetime, now_datetime


DOCTYPE = "HRMS Announcement"
APPROVED_STATUSES = ("审核通过", "待上传签字", "已签字")
SUBMISSION_STATUSES = ("草稿", "待审核", "审核驳回")
SUBMISSION_RECORD_STATUSES = ("待审核", "审核通过", "审核驳回", "待上传签字", "已签字")
FILE_FIELDS = {"source_attachment", "signed_attachment", "supplementary_attachment"}
MAX_ANNOUNCEMENT_PREVIEW_BYTES = 10 * 1024 * 1024
_UNSAFE_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|\x00-\x1f]')

ANNOUNCEMENT_PAGE_DEFINITIONS = (
	("announcement-directory", "公告目录", "list", ()),
	("announcement-submit", "提交公告", "edit", ()),
	("announcement-submission-records", "提交记录", "list", ()),
	("announcement-approval", "公告审批", "check", ()),
	("announcement-approval-records", "公告审批记录", "list", ()),
	("announcement-signed-upload", "上传签字版", "upload", ()),
	("announcement-signed-records", "签字版记录", "list", ()),
)


def ensure_announcement_pages():
	"""Ensure the announcement Desk pages exist after migration on existing sites."""
	created = []
	updated = []
	for page_name, title, icon, roles in ANNOUNCEMENT_PAGE_DEFINITIONS:
		values = {
			"doctype": "Page",
			"name": page_name,
			"page_name": page_name,
			"title": title,
			"module": "HR",
			"icon": icon,
			"standard": "Yes",
			"system_page": 0,
		}
		existing_name = frappe.db.exists("Page", page_name) or frappe.db.get_value(
			"Page", {"page_name": page_name}, "name"
		)
		if not existing_name:
			page_doc = frappe.get_doc(values)
			page_doc.set("roles", [{"role": role} for role in roles])
			page_doc.insert(ignore_permissions=True)
			created.append(page_name)
			continue
		page_doc = frappe.get_doc("Page", existing_name)
		changed = False
		for fieldname in ("page_name", "title", "module", "icon", "standard", "system_page"):
			if getattr(page_doc, fieldname, None) != values[fieldname]:
				page_doc.set(fieldname, values[fieldname])
				changed = True
		desired_roles = [role for role in roles]
		if [row.role for row in page_doc.roles] != desired_roles:
			page_doc.set("roles", [{"role": role} for role in desired_roles])
			changed = True
		if changed:
			page_doc.save(ignore_permissions=True)
			updated.append(page_name)
	if created or updated:
		frappe.clear_cache()
	return {"created": created, "updated": updated}


def _require(key, legacy_roles=()):
	from hrms.access_control import require_hrms_capability

	require_hrms_capability(key, legacy_roles=legacy_roles)


def _has(key, legacy_roles=()):
	from hrms.access_control import has_hrms_capability

	return has_hrms_capability(key, legacy_roles=legacy_roles)


def _payload(payload):
	if isinstance(payload, str):
		try:
			return json.loads(payload)
		except ValueError:
			frappe.throw(_("公告内容格式不正确。"))
	return payload or {}


def _user_name(user):
	return frappe.db.get_value("User", user, "full_name") or user


def _visible(doc):
	if _has("announcement_approve", legacy_roles=("HR Manager",)):
		return True
	if _has("announcement_sign_upload", legacy_roles=("HR Manager",)) and doc.status in {"审核通过", "待上传签字", "已签字"}:
		return True
	if doc.owner == frappe.session.user or doc.submitted_by == frappe.session.user:
		return _has("announcement_submit", legacy_roles=("HR Manager",))
	return doc.status in APPROVED_STATUSES and _has(
		"announcement_view", legacy_roles=("Employee", "HR Manager")
	)


def _get_visible(name):
	doc = frappe.get_doc(DOCTYPE, name)
	if not _visible(doc):
		frappe.throw(_("当前账户无权查看此公告。"), frappe.PermissionError)
	return doc


def _announcement_event_datetime(doc):
	"""Use signature upload time after signing, otherwise the approval time."""
	if doc.signed_attachment and doc.signed_on:
		return get_datetime(doc.signed_on)
	return get_datetime(doc.reviewed_on) if doc.reviewed_on else None


def _announcement_event_date(doc):
	event_datetime = _announcement_event_datetime(doc)
	if not event_datetime:
		return ""
	return f"{event_datetime.year}年{event_datetime.month}月{event_datetime.day}日"


def _announcement_event_date_key(doc):
	event_datetime = _announcement_event_datetime(doc)
	return event_datetime.strftime("%Y-%m-%d %H:%M:%S") if event_datetime else ""


def _serialise(doc, include_content=False):
	files = frappe.get_all(
		"File",
		filters={"attached_to_doctype": DOCTYPE, "attached_to_name": doc.name},
		fields=["name", "file_name", "file_url", "attached_to_doctype", "attached_to_name", "attached_to_field", "creation"],
		order_by="creation asc",
		limit_page_length=0,
		ignore_permissions=True,
	)
	versions = _archive_versions(files)
	return {
		"name": doc.name,
		"announcement_number": doc.announcement_number or "",
		"announcement_year": doc.announcement_year or "",
		"announcement_month": doc.announcement_month or "",
		"sequence": doc.sequence or "",
		"created_on": str(doc.creation or ""),
		"announcement_date": _announcement_event_date(doc),
		"announcement_date_key": _announcement_event_date_key(doc),
		"status": doc.status or "草稿",
		"signature_status": "已签字" if doc.signed_attachment else "未签字",
		"subject": doc.subject or "",
		"content": doc.content or "" if include_content else "",
		"receiving_units": doc.receiving_units or "",
		"cc_units": doc.cc_units or "",
		"issuing_unit": doc.issuing_unit or "",
		"issuing_unit_name": frappe.db.get_value("Department", doc.issuing_unit, "department_name") or doc.issuing_unit or "",
		"issuer_name": doc.issuer_name or "",
		"issuer_user": doc.issuer_user or "",
		"reviewer_name": doc.reviewer_name or "",
		"approver_name": doc.approver_name or "",
		"approval_comment": doc.approval_comment or "",
		"submitted_by": doc.submitted_by or "",
		"submitted_by_name": _user_name(doc.submitted_by) if doc.submitted_by else "",
		"submitted_on": str(doc.submitted_on or ""),
		"reviewed_on": str(doc.reviewed_on or ""),
		"signed_by": doc.signed_by or "",
		"signed_by_name": _user_name(doc.signed_by) if doc.signed_by else "",
		"signed_on": str(doc.signed_on or ""),
		"source_attachment": doc.source_attachment or "",
		"signed_attachment": doc.signed_attachment or "",
		"word_template": doc.word_template or "",
		"excel_template": doc.excel_template or "",
		"files": files,
		"versions": versions,
	}


def _archive_versions(files):
	"""Group retained files by archive version, independently of workflow status."""
	signed_files = [file for file in files if file.get("attached_to_field") == "signed_attachment"]
	unsigned_files = [file for file in files if file.get("attached_to_field") != "signed_attachment"]
	versions = [{"key": "unsigned", "label": "未签字版", "signature_status": "未签字", "files": unsigned_files}]
	if signed_files:
		# Supplementary files uploaded with/after the signed copy belong to the
		# signed archive; before signing they remain visible with the unsigned one.
		signed_files.extend(file for file in unsigned_files if file.get("attached_to_field") == "supplementary_attachment")
		versions[0]["files"] = [file for file in unsigned_files if file.get("attached_to_field") != "supplementary_attachment"]
		versions.append({"key": "signed", "label": "签字版", "signature_status": "已签字", "files": signed_files})
	return versions


def _expand_directory_versions(row):
	"""Return one directory row per retained archive version.

	The announcement remains one business record and keeps one number. The
	directory nevertheless needs two records after signing so users can find
	the unsigned and signed files independently instead of seeing only a
	workflow/signature status change on one row.
	"""
	return [
		{
			**row,
			"archive_version": version.get("key") or "unsigned",
			"version_label": version.get("label") or "未签字版",
			"signature_status": version.get("signature_status") or "未签字",
			"has_signed_version": any(item.get("key") == "signed" for item in row.get("versions") or []),
			"files": version.get("files") or [],
			"versions": [version],
		}
		for version in row.get("versions") or []
	]


@frappe.whitelist()
def list_announcements(
	search: str = "",
	view: str = "directory",
	issuer: str = "",
	announcement_number: str = "",
	signature_status: str = "",
	sort_field: str = "announcement_date",
	sort_order: str = "desc",
):
	"""Return workflow rows, or one directory row for each archive version."""
	if view == "submit":
		_require("announcement_submit", legacy_roles=("HR Manager",))
		# 提交窗口只保留创建人仍可处理或追踪的未完成审批记录。
		filters = {"owner": frappe.session.user, "status": ["in", SUBMISSION_STATUSES]}
	elif view == "submission_records":
		_require("announcement_submit", legacy_roles=("HR Manager",))
		# 提交记录从提交动作完成后即展示，包含待审核和后续处理状态。
		filters = {"owner": frappe.session.user, "status": ["in", SUBMISSION_RECORD_STATUSES]}
	elif view == "approval":
		_require("announcement_approve", legacy_roles=("HR Manager",))
		# 审批窗口只处理尚未审核的公告；审核通过/驳回的记录进入后续窗口或目录。
		filters = {"status": "待审核"}
	elif view == "approval_records":
		_require("announcement_approve", legacy_roles=("HR Manager",))
		# 审批记录只读展示已完成审批的公告，兼容旧的签字流程状态。
		filters = {"status": ["in", ["审核通过", "审核驳回", "待上传签字", "已签字"]]}
	elif view == "sign":
		_require("announcement_sign_upload", legacy_roles=("HR Manager",))
		# 待上传列表只保留尚未归档签字版的记录；已归档记录进入 signed_records。
		filters = {"status": ["in", ["审核通过", "待上传签字", "已签字"]]}
	elif view == "signed_records":
		_require("announcement_sign_upload", legacy_roles=("HR Manager",))
		# 签字版记录是独立的只读归档列表，不与待上传列表混合。
		filters = {"status": ["in", ["审核通过", "待上传签字", "已签字"]]}
	else:
		_require("announcement_view", legacy_roles=("Employee", "HR Manager"))
		filters = {"status": ["in", APPROVED_STATUSES]}
	if search:
		search = str(search).strip()
		filters["subject"] = ["like", f"%{search}%"]
	if issuer:
		filters["issuer_name"] = ["like", f"%{str(issuer).strip()}%"]
	if announcement_number:
		filters["announcement_number"] = ["like", f"%{str(announcement_number).strip()}%"]
	rows = frappe.get_all(
		DOCTYPE,
		filters=filters,
		fields=[
			"name", "announcement_number", "announcement_year", "announcement_month", "sequence",
			"status", "subject", "issuer_name", "issuing_unit", "reviewer_name", "reviewed_on", "approval_comment", "signed_attachment", "signed_by", "signed_on", "submitted_by", "submitted_on", "creation",
		],
		order_by="creation desc",
		limit_page_length=0,
		ignore_permissions=True,
	)
	for row in rows:
		row["created_on"] = str(row.get("creation") or "")
		row["signature_status"] = "已签字" if row.get("signed_attachment") else "未签字"
		row["signed_by_name"] = _user_name(row.get("signed_by")) if row.get("signed_by") else ""
		event_datetime = get_datetime(row.get("signed_on")) if row.get("signed_attachment") and row.get("signed_on") else get_datetime(row.get("reviewed_on")) if row.get("reviewed_on") else None
		row["announcement_date"] = f"{event_datetime.year}年{event_datetime.month}月{event_datetime.day}日" if event_datetime else ""
		row["announcement_date_key"] = event_datetime.strftime("%Y-%m-%d %H:%M:%S") if event_datetime else ""
		row["issuing_unit_name"] = frappe.db.get_value("Department", row.get("issuing_unit"), "department_name") or row.get("issuing_unit") or ""
		row["files"] = frappe.get_all(
			"File",
			filters={"attached_to_doctype": DOCTYPE, "attached_to_name": row.name},
			fields=["file_name", "file_url", "attached_to_doctype", "attached_to_name", "attached_to_field"],
			order_by="creation asc",
			limit_page_length=0,
			ignore_permissions=True,
		)
		row["versions"] = _archive_versions(row["files"])
	if view == "submit":
		# 已提交记录统一进入提交记录，提交页只保留仍可处理的草稿/驳回项。
		rows = [row for row in rows if row.get("status") in {"草稿", "审核驳回"}]
	if view == "sign":
		rows = [row for row in rows if not row.get("signed_attachment")]
	elif view == "signed_records":
		rows = [row for row in rows if row.get("signed_attachment")]
	if view == "directory":
		rows = [version_row for row in rows for version_row in _expand_directory_versions(row)]
	if signature_status in {"已签字", "未签字"}:
		rows = [row for row in rows if row["signature_status"] == signature_status]

	sort_fields = {
		"announcement_date": "announcement_date_key",
		"sequence": "sequence",
		"issuer_name": "issuer_name",
		"subject": "subject",
		"announcement_number": "announcement_number",
		"created_on": "created_on",
		"reviewed_on": "reviewed_on",
		"signed_on": "signed_on",
		"reviewer_name": "reviewer_name",
		"status": "status",
		"signature_status": "signature_status",
		"version_label": "version_label",
	}
	sort_key = sort_fields.get(sort_field, "announcement_date_key")
	reverse = str(sort_order).lower() == "desc"
	rows.sort(
		key=lambda row: int(row.get(sort_key) or 0) if sort_key == "sequence" else str(row.get(sort_key) or "").lower(),
		reverse=reverse,
	)
	# Keep records without a date/name at the bottom in either direction.
	rows.sort(key=lambda row: row.get(sort_key) in (None, ""))
	return rows


def _directory_export_filters(current_filters):
	if not current_filters:
		return {}
	try:
		filters = json.loads(current_filters) if isinstance(current_filters, str) else current_filters
	except (TypeError, ValueError, json.JSONDecodeError):
		frappe.throw(_("公告目录筛选条件格式不正确。"))
	return filters if isinstance(filters, dict) else {}


def _directory_export_matches(row, filters):
	for key, query in filters.items():
		query = str(query or "").strip().casefold()
		if not query or key in {"files", "actions"}:
			continue
		value = row.get(key, "")
		if isinstance(value, (list, tuple, dict)):
			value = ""
		if query not in str(value or "").strip().casefold():
			return False
	return True


def _directory_export_date_range(start_date, end_date):
	"""Validate and normalise the inclusive announcement-date range."""
	start = str(start_date or "").strip()[:10]
	end = str(end_date or "").strip()[:10]
	for label, value in (("开始日期", start), ("结束日期", end)):
		if not value:
			continue
		try:
			date.fromisoformat(value)
		except ValueError:
			frappe.throw(_("{0}格式不正确，请选择有效日期。").format(label))
	if start and end and start > end:
		frappe.throw(_("开始日期不能晚于结束日期。"))
	return start, end


def _directory_export_matches_date(row, start_date, end_date):
	date_key = str(row.get("announcement_date_key") or "")[:10]
	if start_date and (not date_key or date_key < start_date):
		return False
	if end_date and (not date_key or date_key > end_date):
		return False
	return True


@frappe.whitelist()
def download_announcement_directory_export(
	current_filters: str = "{}",
	sort_field: str = "announcement_date",
	sort_order: str = "desc",
	start_date: str = "",
	end_date: str = "",
):
	"""Download the announcement directory using filters, order and date range."""
	_require("announcement_view", legacy_roles=("Employee", "HR Manager"))
	filters = _directory_export_filters(current_filters)
	start_date, end_date = _directory_export_date_range(start_date, end_date)
	rows = list_announcements(view="directory", sort_field=sort_field, sort_order=sort_order)
	rows = [row for row in rows if _directory_export_matches(row, filters)]
	rows = [row for row in rows if _directory_export_matches_date(row, start_date, end_date)]

	from frappe.desk.utils import provide_binary_file
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter

	columns = [
		("公告日期", "announcement_date"),
		("序号", "sequence"),
		("发文者", "issuer_name"),
		("主旨", "subject"),
		("编号", "announcement_number"),
		("版本", "version_label"),
		("创建时间", "created_on"),
		("审批时间", "reviewed_on"),
		("签字版上传时间", "signed_on"),
		("审核人", "reviewer_name"),
		("审核状态", "status"),
		("签字状态", "signature_status"),
		("附属文件", "files"),
	]
	workbook = Workbook()
	sheet = workbook.active
	sheet.title = "公告目录"
	header_fill = PatternFill("solid", fgColor="DDEBF7")
	thin = Side(style="thin", color="B7C9D6")
	border = Border(left=thin, right=thin, top=thin, bottom=thin)
	sheet.append([label for label, _field in columns])
	for cell in sheet[1]:
		cell.fill = header_fill
		cell.font = Font(name="Microsoft YaHei", size=10, bold=True)
		cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
		cell.border = border

	for row in rows:
		file_names = "\n".join(
			file.get("file_name") or file.get("file_url") or ""
			for file in row.get("files") or []
		)
		values = dict(row)
		values["files"] = file_names
		sheet.append([values.get(field, "") for _label, field in columns])
		for cell in sheet[sheet.max_row]:
			if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
				cell.value = "'" + cell.value
			cell.alignment = Alignment(vertical="top", wrap_text=True)
			cell.border = border

	for index, (label, field) in enumerate(columns, start=1):
		values = [str((row.get(field) or "") if field != "files" else "") for row in rows[:500]]
		if field == "files":
			values = [
				"\n".join(file.get("file_name") or file.get("file_url") or "" for file in row.get("files") or [])
				for row in rows[:500]
			]
		sheet.column_dimensions[get_column_letter(index)].width = min(
			42, max(12, max([len(label), *(len(value) for value in values)], default=12) + 2)
		)
	sheet.freeze_panes = "A2"
	sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{sheet.max_row}"
	sheet.sheet_view.showGridLines = False

	output = io.BytesIO()
	workbook.save(output)
	provide_binary_file("公告目录", "xlsx", output.getvalue())


@frappe.whitelist()
def get_announcement(name: str):
	return _serialise(_get_visible(name), include_content=True)


@frappe.whitelist()
def preview_announcement_file(name: str, file_url: str):
	"""Return a safe, read-only preview model for an announcement attachment."""
	doc = _get_visible(name)
	file_name = frappe.db.get_value(
		"File",
		{
			"file_url": file_url,
			"attached_to_doctype": DOCTYPE,
			"attached_to_name": doc.name,
		},
		"name",
	)
	if not file_name:
		frappe.throw(_("该文件不是此公告的附件。"), frappe.PermissionError)
	file_doc = frappe.get_doc("File", file_name)
	content = file_doc.get_content()
	if len(content or b"") > MAX_ANNOUNCEMENT_PREVIEW_BYTES:
		return {
			"file_name": file_doc.file_name,
			"file_url": file_doc.file_url,
			"kind": "unsupported",
			"message": _("文件超过 10 MB，暂不生成在线预览，请下载原文件查看。"),
		}

	extension = _file_extension(file_doc.file_name or file_doc.file_url)
	try:
		if extension == ".docx":
			return {
				"file_name": file_doc.file_name,
				"file_url": file_doc.file_url,
				"kind": "docx",
				"paragraphs": _preview_docx(content),
			}
		if extension == ".xlsx":
			return {
				"file_name": file_doc.file_name,
				"file_url": file_doc.file_url,
				"kind": "xlsx",
				"sheets": _preview_xlsx(content),
			}
		if extension in {".jpg", ".jpeg", ".png", ".webp", ".gif"}:
			return {
				"file_name": file_doc.file_name,
				"file_url": file_doc.file_url,
				"kind": "image",
			}
	except (ElementTree.ParseError, KeyError, ValueError, zipfile.BadZipFile, OSError) as exc:
		frappe.log_error(f"公告附件预览失败: {file_doc.file_name}: {exc}", "Announcement attachment preview")
		return {
			"file_name": file_doc.file_name,
			"file_url": file_doc.file_url,
			"kind": "unsupported",
			"message": _("文件内容无法解析，请下载原文件查看。"),
		}

	return {
		"file_name": file_doc.file_name,
		"file_url": file_doc.file_url,
		"kind": "unsupported",
		"message": _("此格式暂不支持在线预览，请下载原文件查看。"),
	}


def _file_extension(file_name):
	return "." + str(file_name or "").rsplit(".", 1)[-1].lower() if "." in str(file_name or "") else ""


def _preview_docx(content):
	with zipfile.ZipFile(io.BytesIO(content)) as archive:
		root = ElementTree.fromstring(archive.read("word/document.xml"))
	paragraphs = []
	for paragraph in root.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
		parts = []
		for node in paragraph.iter():
			if node.tag.endswith("}t") and node.text:
				parts.append(node.text)
			elif node.tag.endswith("}tab"):
				parts.append("\t")
			elif node.tag.endswith("}br"):
				parts.append("\n")
		paragraphs.append("".join(parts))
	return paragraphs[:2000]


def _preview_xlsx(content):
	from openpyxl import load_workbook
	from openpyxl.utils import get_column_letter

	workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
	sheets = []
	try:
		for worksheet in workbook.worksheets[:20]:
			max_row = min(worksheet.max_row or 1, 100)
			max_column = min(worksheet.max_column or 1, 30)
			rows = []
			for row in worksheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_column, values_only=True):
				values = [_preview_cell(value) for value in row]
				while values and values[-1] == "":
					values.pop()
				if values:
					rows.append(values)
			if rows:
				width = min(max(len(row) for row in rows), 30)
			else:
				width = max_column
			sheets.append({
				"name": worksheet.title,
				"columns": [get_column_letter(index) for index in range(1, width + 1)],
				"rows": [row[:width] for row in rows],
			})
	finally:
		workbook.close()
	return sheets


def _preview_cell(value):
	if value is None:
		return ""
	if isinstance(value, (datetime, date, time)):
		return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
	return str(value)


@frappe.whitelist()
def save_announcement(payload: Any = None, name: str = ""):
	_require("announcement_submit", legacy_roles=("HR Manager",))
	values = _payload(payload)
	if name:
		doc = frappe.get_doc(DOCTYPE, name)
		if doc.owner != frappe.session.user or doc.status != "草稿":
			frappe.throw(_("只能修改本人未提交的公告草稿。"), frappe.PermissionError)
		is_new = False
	else:
		doc = frappe.new_doc(DOCTYPE)
		doc.status = "草稿"
		is_new = True
	for field in ("subject", "content", "receiving_units", "cc_units", "issuing_unit", "issuer_name", "approver_name"):
		if field in values:
			doc.set(field, values.get(field))
	doc.validate()
	doc.flags.ignore_permissions = True
	if is_new:
		doc.insert(ignore_permissions=True)
	else:
		doc.save(ignore_permissions=True)
	return _serialise(doc, include_content=True)


@frappe.whitelist()
def submit_announcement(name: str):
	_require("announcement_submit", legacy_roles=("HR Manager",))
	doc = frappe.get_doc(DOCTYPE, name)
	if doc.owner != frappe.session.user or doc.status != "草稿":
		frappe.throw(_("只能提交本人未提交的公告草稿。"), frappe.PermissionError)
	doc.prepare_submission()
	doc.flags.ignore_permissions = True
	doc.save(ignore_permissions=True)
	_generate_templates(doc)
	doc.reload()
	return _serialise(doc, include_content=True)


@frappe.whitelist()
def generate_announcement_templates(name: str):
	_require("announcement_submit", legacy_roles=("HR Manager",))
	doc = frappe.get_doc(DOCTYPE, name)
	if doc.owner != frappe.session.user and not _has("announcement_approve", legacy_roles=("HR Manager",)):
		frappe.throw(_("只有公告提交人或审批人可以生成模板。"), frappe.PermissionError)
	if not doc.announcement_number:
		frappe.throw(_("公告提交后才能生成带编号的 Word/Excel 模板。"))
	_generate_templates(doc)
	doc.reload()
	return _serialise(doc, include_content=True)


def _generate_templates(doc):
	context = _template_context(doc)
	word_filename = _announcement_filename(doc, ".docx")
	excel_filename = _announcement_filename(doc, ".xlsx")
	if not doc.word_template:
		word = _create_file(
			doc,
			word_filename,
			_build_docx(context),
			"word_template",
		)
		doc.db_set("word_template", word.file_url)
	else:
		_rename_file(doc.word_template, word_filename)
	if not doc.excel_template:
		excel = _create_file(
			doc,
			excel_filename,
			_build_xlsx(context),
			"excel_template",
		)
		doc.db_set("excel_template", excel.file_url)
	else:
		_rename_file(doc.excel_template, excel_filename)


def _announcement_filename(doc, extension):
	"""Use the announcement subject as the human-readable file-name suffix."""
	subject = _UNSAFE_FILENAME_CHARS.sub("_", str(doc.subject or "").strip())
	subject = re.sub(r"\s+", " ", subject).rstrip(" .")[:80].rstrip(" .")
	return f"{doc.announcement_number}-{subject or '未填写主旨'}{extension}"


def _rename_file(file_url, filename):
	"""Rename an already attached file without changing its stored URL/content."""
	file_name = frappe.db.get_value("File", {"file_url": file_url}, "name")
	if not file_name:
		return
	file_doc = frappe.get_doc("File", file_name)
	if file_doc.file_name != filename:
		file_doc.db_set("file_name", filename)


def _template_context(doc):
	company = frappe.db.get_value("Department", doc.issuing_unit, "company")
	company_name = frappe.db.get_value("Company", company, "company_name") or company or "永新电子(常熟)有限公司"
	department_name = frappe.db.get_value("Department", doc.issuing_unit, "department_name") or doc.issuing_unit or ""
	return {
		"company": company_name,
		"number": doc.announcement_number or "待提交自动生成",
		"subject": doc.subject or "",
		"content": doc.content or "",
		"receiving": doc.receiving_units or "",
		"cc": doc.cc_units or "",
		"issuing": department_name,
		"issuer": doc.issuer_name or "",
		"reviewer": doc.reviewer_name or "",
		"approver": doc.approver_name or "",
		"date": str(doc.submitted_on or now_datetime())[:10],
	}


def _create_file(doc, filename, content, fieldname):
	file_doc = frappe.get_doc(
		{
			"doctype": "File",
			"file_name": filename,
			"content": content,
			"is_private": 1,
			"attached_to_doctype": DOCTYPE,
			"attached_to_name": doc.name,
			"attached_to_field": fieldname,
		}
	).insert(ignore_permissions=True)
	return file_doc


@frappe.whitelist()
def review_announcement(name: str, decision: str, approval_comment: str = "", approver_name: str = ""):
	_require("announcement_approve", legacy_roles=("HR Manager",))
	if decision not in {"approve", "reject"}:
		frappe.throw(_("审核结果只能是批准或驳回。"))
	doc = frappe.get_doc(DOCTYPE, name)
	if doc.status != "待审核":
		frappe.throw(_("只有待审核公告可以处理。"))
	values = {
		"status": "审核通过" if decision == "approve" else "审核驳回",
		"reviewer_user": frappe.session.user,
		"reviewer_name": _user_name(frappe.session.user),
		"reviewed_on": now_datetime(),
		"approval_comment": str(approval_comment or "").strip(),
	}
	if str(approver_name or "").strip():
		values["approver_name"] = str(approver_name).strip()
	doc.db_set(values)
	doc.reload()
	return _serialise(doc, include_content=True)


@frappe.whitelist()
def attach_announcement_file(name: str, file_url: str, fieldname: str = "source_attachment"):
	"""Attach a FileUploader result and keep the field boundary explicit."""
	if fieldname not in FILE_FIELDS:
		frappe.throw(_("不支持的公告附件类型。"))
	doc = _get_visible(name)
	if fieldname == "source_attachment":
		_require("announcement_submit", legacy_roles=("HR Manager",))
		if doc.owner != frappe.session.user or doc.status != "草稿":
			frappe.throw(_("只有公告草稿提交人可以上传提交附件。"), frappe.PermissionError)
	elif fieldname == "signed_attachment":
		_require("announcement_sign_upload", legacy_roles=("HR Manager",))
		if doc.signed_attachment or doc.status == "已签字":
			frappe.throw(_("该公告已完成签字归档，不能重复修改。"))
		if doc.status not in {"审核通过", "待上传签字", "已签字"}:
			frappe.throw(_("公告审批通过后才能上传签字版。"))
	elif fieldname == "supplementary_attachment":
		_require("announcement_sign_upload", legacy_roles=("HR Manager",))
		if doc.status not in {"审核通过", "待上传签字", "已签字"}:
			frappe.throw(_("公告审批通过后才能上传附属文件。"))
	file_doc = _find_owned_file(file_url, doc.name)
	file_values = {
		"attached_to_doctype": DOCTYPE,
		"attached_to_name": doc.name,
		"attached_to_field": fieldname,
	}
	if fieldname == "signed_attachment":
		file_extension = _file_extension(file_doc.file_name or file_doc.file_url) or ".docx"
		file_values["file_name"] = _announcement_filename(doc, file_extension)
	file_doc.db_set(file_values)
	if fieldname == "source_attachment":
		doc.db_set("source_attachment", file_url)
	return {"file_url": file_url, "file_name": file_doc.file_name}


@frappe.whitelist()
def upload_signed_announcement(name: str, signed_file_url: str, supplementary_file_urls: Any = None):
	_require("announcement_sign_upload", legacy_roles=("HR Manager",))
	doc = frappe.get_doc(DOCTYPE, name)
	if doc.status not in {"审核通过", "待上传签字", "已签字"}:
		frappe.throw(_("只有审批通过的公告才能上传签字版。"))
	if doc.signed_attachment:
		frappe.throw(_("该公告已完成签字归档，不能重复修改。"))
	file_doc = _find_owned_file(signed_file_url, doc.name)
	file_extension = _file_extension(file_doc.file_name or file_doc.file_url) or ".docx"
	file_doc.db_set(
		{
			"file_name": _announcement_filename(doc, file_extension),
			"attached_to_doctype": DOCTYPE,
			"attached_to_name": doc.name,
			"attached_to_field": "signed_attachment",
		}
	)
	now = now_datetime()
	# The signed file is a retained archive version. It does not replace the
	# unsigned files and does not mutate the announcement workflow status.
	doc.db_set({"signed_attachment": signed_file_url, "signed_by": frappe.session.user, "signed_on": now})
	urls = _payload(supplementary_file_urls)
	if isinstance(urls, str):
		urls = [urls]
	for url in urls or []:
		if url:
			additional = _find_owned_file(url, doc.name)
			additional.db_set({"attached_to_doctype": DOCTYPE, "attached_to_name": doc.name, "attached_to_field": "supplementary_attachment"})
	doc.reload()
	return _serialise(doc, include_content=True)


def _find_owned_file(file_url, announcement_name):
	file_name = frappe.db.get_value("File", {"file_url": file_url}, "name")
	if not file_name:
		frappe.throw(_("附件不存在，请重新上传。"))
	file_doc = frappe.get_doc("File", file_name)
	if file_doc.owner != frappe.session.user and file_doc.attached_to_name != announcement_name:
		frappe.throw(_("不能使用其他账户上传的附件。"), frappe.PermissionError)
	return file_doc


def _plain_text(value):
	value = str(value or "")
	value = re.sub(r"<br\s*/?>", "\n", value, flags=re.I)
	value = re.sub(r"</(?:p|div|li|h[1-6]|tr)\s*>", "\n", value, flags=re.I)
	value = re.sub(r"<li\b[^>]*>", "• ", value, flags=re.I)
	value = html.unescape(re.sub(r"<[^>]+>", "", value))
	value = value.replace("\r\n", "\n").replace("\r", "\n")
	return "\n".join(line.rstrip() for line in value.split("\n")).strip()


def _docx_paragraph(text="", align="left", bold=False, size=24):
	text = html.escape(str(text or ""))
	align_xml = f'<w:jc w:val="{align}"/>' if align else ""
	bold_xml = "<w:b/>" if bold else ""
	return f'<w:p><w:pPr>{align_xml}</w:pPr><w:r><w:rPr>{bold_xml}<w:sz w:val="{size}"/></w:rPr><w:t xml:space="preserve">{text}</w:t></w:r></w:p>'


def _build_docx(context):
	body = _plain_text(context.get("content"))
	body_lines = body.split("\n") if body else [""]
	lines = [
		_docx_paragraph("公 告", "center", True, 40),
		_docx_paragraph(context["company"], "center", True, 28),
		_docx_paragraph(f"收文单位：{context['receiving']}"),
		_docx_paragraph(f"附本抄送：{context['cc']}    发文单位：{context['issuing']}    编号：{context['number']}"),
		_docx_paragraph(f"发文者：{context['issuer']}    审核：{context['reviewer']}    核准：{context['approver']}"),
		_docx_paragraph(f"主旨：{context['subject']}", bold=True, size=28),
		*[_docx_paragraph(line) for line in body_lines],
		_docx_paragraph(f"以上特此公告！                                      {context['issuing']}{context['date']}"),
	]
	document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>{''.join(lines)}<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/></w:sectPr></w:body></w:document>'''
	content_types = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'''
	rels = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'''
	buffer = io.BytesIO()
	with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
		archive.writestr("[Content_Types].xml", content_types)
		archive.writestr("_rels/.rels", rels)
		archive.writestr("word/document.xml", document)
	return buffer.getvalue()


def _build_xlsx(context):
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Border, Font, Side

	book = Workbook()
	sheet = book.active
	sheet.title = "模板"
	for column in range(2, 22):
		sheet.column_dimensions[chr(64 + column)].width = 5.8
	for row in range(1, 41):
		sheet.row_dimensions[row].height = 25
	thin = Side(style="thin", color="808080")
	medium = Side(style="medium", color="000000")
	base = Font(name="宋体", size=14)
	for row in range(1, 19):
		for column in range(2, 22):
			cell = sheet.cell(row, column)
			cell.font = base
			cell.alignment = Alignment(vertical="center", wrap_text=True)
			cell.border = Border(bottom=thin)
	for rng in ("B1:U1", "B2:U2", "B3:U3", "B4:U4", "B5:U5", "B6:U6", "B7:U7", "B8:U17", "B18:U18"):
		sheet.merge_cells(rng)
	sheet["B1"] = "公  告"
	sheet["B1"].font = Font(name="宋体", size=28, bold=True)
	sheet["B1"].alignment = Alignment(horizontal="center", vertical="center")
	sheet["B2"] = context["company"]
	sheet["B2"].font = Font(name="宋体", size=18, bold=True)
	sheet["B2"].alignment = Alignment(horizontal="center", vertical="center")
	sheet["B3"] = f"收文单位：{context['receiving']}"
	sheet["B4"] = f"附本抄送：{context['cc']}    发文单位：{context['issuing']}    编号：{context['number']}"
	sheet["B5"] = f"发文者：{context['issuer']}    审核：{context['reviewer']}    核准：{context['approver']}"
	sheet["B6"] = f"主旨：{context['subject']}"
	sheet["B6"].font = Font(name="宋体", size=16, bold=True)
	sheet["B8"] = _plain_text(context.get("content"))
	sheet["B8"].alignment = Alignment(vertical="top", wrap_text=True)
	sheet["B18"] = f"以上特此公告！                                      {context['issuing']}{context['date']}"
	for column in range(2, 22):
		sheet.cell(6, column).border = Border(bottom=medium)
	sheet.sheet_view.showGridLines = False
	sheet.page_setup.orientation = "portrait"
	sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
	sheet.page_margins.left = 0.35
	sheet.page_margins.right = 0.35
	sheet.page_margins.top = 0.35
	sheet.page_margins.bottom = 0.35
	sheet.print_area = "B1:U18"
	buffer = io.BytesIO()
	book.save(buffer)
	return buffer.getvalue()
