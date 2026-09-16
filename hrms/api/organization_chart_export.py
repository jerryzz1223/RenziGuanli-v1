"""Editable cell-and-border diagrams from the same live tree used by the page.

Runs inside Frappe with its existing openpyxl dependency; no desktop renderer or
original template file is required on the deployment server.
"""
from copy import copy, deepcopy
from io import BytesIO
import re
import unicodedata

import frappe
from frappe.utils import now_datetime


def _walk(node):
	yield node
	for child in node.get("children") or []:
		yield from _walk(child)


def _text(value):
	return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value if value is not None else ""))


def _wrap(value, limit=28):
	rows = []
	for paragraph in _text(value).splitlines() or [""]:
		line, width = "", 0
		for char in paragraph:
			size = 2 if unicodedata.east_asian_width(char) in "WF" else 1
			if line and width + size > limit:
				rows.append(line)
				line, width = "", 0
			line += char
			width += size
		rows.append(line)
	return rows


def node_lines(node):
	"""Preserve card roles and every displayed person, including pending entries."""
	rows = []
	def add(value, kind="body"):
		if value is not None and value != "":
			rows.extend((line, kind) for line in _wrap(value))
	add(node.get("name"), "heading")
	add(node.get("title"), "type")
	if node.get("reporting_scope_pending"):
		add("分管待设置", "muted")
	grade = node.get("chart_grade") or {}
	if grade.get("label"):
		add(f"职级：{grade['label']}")
		if grade.get("rank") is not None:
			add(f"等级：{grade['rank']}")
	if node.get("source_grade_tags"):
		labels = "、".join(node["source_grade_tags"].splitlines())
		add(f"原表职级：{labels}（{node.get('source_grade_status') or '待确认'}）", "muted")
	if node.get("people"):
		for person in node["people"]:
			add("：".join(str(v) for v in (person.get("role"), person.get("employee_name") or person.get("name")) if v))
			if person.get("source_reference"):
				add(f"原表 · {person.get('match_status') or '待核对'}", "muted")
	else:
		for line in node.get("employee_names") or node.get("lines") or []:
			add(line)
	add(node.get("card_content"))
	if node.get("node_type") != "organization_person":
		planned = "未设置" if node.get("has_staffing_plan") is False else node.get("planned_headcount", 0) or 0
		vacant = "未设置" if node.get("has_staffing_plan") is False else node.get("vacancy_count", 0) or 0
		add(f"编制/实际：{planned}/{node.get('current_headcount') or 0}", "metric")
		add(f"空缺：{vacant}", "metric")
	if node.get("template_source_vacancies"):
		add(f"原表空缺：{node['template_source_vacancies']}", "vacancy")
	if node.get("export_collapsed"):
		add(f"已收起 {node['export_collapsed']} 个下级", "muted")
	return rows


def select_tree(root, scope, node_ids=None, search=""):
	if scope == "complete":
		return deepcopy(root)
	if scope != "current":
		raise ValueError("不支持的导出范围")
	ids = set(node_ids or [])
	if not ids or len(ids) > 10000:
		raise ValueError("当前没有可导出的节点，请刷新架构图后重试")
	known = {n["node_id"] for n in _walk(root)}
	if not ids <= known:
		raise ValueError("组织节点已变化，请刷新架构图后重试")
	def prune(node):
		if node["node_id"] not in ids:
			return None
		result = deepcopy({k: v for k, v in node.items() if k != "children"})
		result["children"] = [value for child in node.get("children", []) if (value := prune(child))]
		if node.get("children") and not result["children"]:
			result["export_collapsed"] = len(node["children"])
		if search and result.get("people"):
			result["people"] = [p for p in result["people"] if search.lower() in " ".join(str(p.get(k) or "") for k in ("name", "employee_name", "employee_code", "department", "designation", "grade", "role")).lower()]
			# Empty search results must not fall back to all names.
			result["employee_names"] = result["lines"] = []
		return result
	for node in _walk(root):
		if node["node_id"] in ids:
			selected = prune(node)
			if {n["node_id"] for n in _walk(selected)} != ids:
				raise ValueError("导出节点必须属于同一棵连续的组织树")
			return selected


def _report_cell(sheet, row, start_col, end_col, value, font, border, alignment, number_format=None):
	from openpyxl.styles import Alignment

	sheet.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
	cell = sheet.cell(row, start_col)
	cell.value = _text(value) if isinstance(value, str) else value
	if isinstance(value, str):
		cell.data_type = "s"
	cell.font = font
	cell.alignment = alignment
	if number_format:
		cell.number_format = number_format
	for column in range(start_col, end_col + 1):
		merged_cell = sheet.cell(row, column)
		merged_cell.border = border
		merged_cell.alignment = Alignment(
			horizontal=alignment.horizontal,
			vertical=alignment.vertical,
			wrap_text=alignment.wrap_text,
		)
	return cell


def _append_organization_report(sheet, report, diagram_end_row, diagram_end_col, font_name):
	"""Place the page's organization report below the editable chart."""
	from openpyxl.styles import Alignment, Border, Font, Side
	from openpyxl.utils import get_column_letter

	columns = report.get("columns") or ["部门/课别", "编制人数", "现有人数", "空缺人数", "岗位满足率", "备注"]
	keys = ("department", "planned_headcount", "current_headcount", "vacancy_count", "fulfillment_rate", "vacancy_notes")
	spans = (6, 4, 4, 4, 5, 8)
	table_width = sum(spans)
	left = max(1, (max(diagram_end_col, table_width) - table_width) // 2 + 1)
	right = left + table_width - 1
	for column in range(diagram_end_col + 1, right + 1):
		sheet.column_dimensions[get_column_letter(column)].width = 5.5

	title_row = diagram_end_row + 3
	header_row = title_row + 2
	line = Side(style="thin", color="333333")
	border = Border(left=line, right=line, top=line, bottom=line)
	title_cell = sheet.cell(title_row, left, _text(report.get("title") or "组织报表"))
	title_cell.data_type = "s"
	title_cell.font = Font(name=font_name, size=14, bold=True, color="111111")
	title_cell.alignment = Alignment(horizontal="left", vertical="center")
	sheet.merge_cells(start_row=title_row, start_column=left, end_row=title_row, end_column=right)
	sheet.row_dimensions[title_row].height = 24

	column_ranges = []
	next_column = left
	for span in spans:
		column_ranges.append((next_column, next_column + span - 1))
		next_column += span
	header_font = Font(name=font_name, size=10, bold=True, color="111111")
	body_font = Font(name=font_name, size=10, color="111111")
	for label, (start_col, end_col) in zip(columns, column_ranges):
		_report_cell(
			sheet,
			header_row,
			start_col,
			end_col,
			label,
			header_font,
			border,
			Alignment(horizontal="center", vertical="center", wrap_text=True),
		)
	sheet.row_dimensions[header_row].height = 24

	rows = list(report.get("rows") or [])
	total = {"department": "汇总", **(report.get("total") or {}), "vacancy_notes": "-"}
	for offset, data in enumerate(rows + [total], start=1):
		row = header_row + offset
		is_total = offset == len(rows) + 1
		for key, (start_col, end_col) in zip(keys, column_ranges):
			value = data.get(key)
			if key == "vacancy_notes":
				value = value or "-"
			if key == "fulfillment_rate" and value is None:
				value = "-"
			_report_cell(
				sheet,
				row,
				start_col,
				end_col,
				value,
				Font(name=font_name, size=10, bold=is_total, color="111111") if is_total else body_font,
				border,
				Alignment(
					horizontal="left" if key in {"department", "vacancy_notes"} else "center",
					vertical="center",
					wrap_text=True,
				),
				"0%" if key == "fulfillment_rate" and isinstance(value, (int, float)) else None,
			)
		sheet.row_dimensions[row].height = 24
	approval_row = header_row + len(rows) + 3
	approval_cell = sheet.cell(approval_row, left, "批准：")
	approval_cell.data_type = "s"
	approval_cell.font = Font(name=font_name, size=10, bold=True, color="111111")
	approval_cell.alignment = Alignment(horizontal="left", vertical="center")
	return approval_row, right


def _diagram(book, root, title, stamp, report=None):
	from openpyxl.styles import Alignment, Border, Font, Side
	from openpyxl.utils import get_column_letter
	from openpyxl.worksheet.page import PageMargins

	name = re.sub(r"[\\/*?:\[\]]", "_", title)[:31] or "架构图"
	base, suffix = name, 2
	while name in book.sheetnames:
		name = f"{base[:26]} ({suffix})"
		suffix += 1
	sheet = book.create_sheet(name)
	sheet.sheet_view.showGridLines = False
	sheet.sheet_view.zoomScale = 100
	sheet.sheet_properties.pageSetUpPr.fitToPage = False
	sheet.page_setup.orientation = "landscape"
	sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
	sheet.page_setup.scale = 100
	sheet.page_margins = PageMargins(left=.25, right=.25, top=.4, bottom=.4, header=.2, footer=.2)
	sheet.oddFooter.center.text = "第 &P 页"
	font = "等线"
	line = Side(style="thin", color="333333")
	nodes, levels = [], []
	def layout(node, depth=0):
		rows = node_lines(node)
		item = {"node": node, "rows": rows, "depth": depth, "height": len(rows) + 2}
		nodes.append(item)
		if depth == len(levels):
			levels.append(0)
		levels[depth] = max(levels[depth], item["height"])
		item["children"] = [layout(c, depth+1) for c in node.get("children") or []]
		item["width"] = max(6, sum(c["width"] for c in item["children"]) + max(0, len(item["children"])-1)*2)
		return item
	graph = layout(root)
	if graph["width"] + 3 > 16384 or sum(levels) + len(levels)*4 > 1048500:
		raise ValueError("架构超出 Excel 工作表大小，请选择局部层级导出")
	y = [6]
	for height in levels[:-1]:
		y.append(y[-1] + height + 4)
	def place(item, left):
		item["center"] = left + item["width"]//2
		item["left"] = item["center"]-3
		item["top"] = y[item["depth"]]
		width = sum(c["width"] for c in item["children"]) + max(0,len(item["children"])-1)*2
		next_left = left + (item["width"]-width)//2
		for child in item["children"]:
			place(child,next_left)
			next_left += child["width"]+2
	place(graph,2)
	header_col = graph['left'] if graph['width'] > 80 else 2
	end_col = max(graph['width']+3,header_col+20)
	for col in range(1,end_col+1):
		sheet.column_dimensions[get_column_letter(col)].width = 5.5
	end_row = y[-1]+levels[-1]+1
	for row in range(1,end_row+1):
		sheet.row_dimensions[row].height = 17
	def write(row,col,text,bold=False,size=11,color="111111"):
		cell=sheet.cell(row,col)
		cell.value=_text(text)
		cell.data_type="s"  # Treat names/notes starting with '=' as literal text.
		cell.font=Font(name=font,size=size,bold=bold,color=color)
		cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=False,shrink_to_fit=False)
		return cell
	if graph['width'] > 80:
		sheet.sheet_view.topLeftCell = f"{get_column_letter(max(1,header_col-5))}1"
	for row,text,size in [(2,f"{root.get('name','')}组织架构图",16),(3,f"来源：系统组织架构　导出时间：{stamp}",10)]:
		sheet.merge_cells(start_row=row,start_column=header_col,end_row=row,end_column=header_col+20)
		cell=write(row,header_col,text,row==2,size)
		cell.alignment=Alignment(horizontal="left",vertical="center")
	sheet.row_dimensions[2].height=26
	def edge(row,col,side):
		cell=sheet.cell(row,col); border=copy(cell.border); setattr(border,side,line);cell.border=border
	for item in nodes:
		x,top,h=item['left'],item['top'],item['height']
		for offset in range(h):
			r=top+offset
			sheet.merge_cells(start_row=r,start_column=x,end_row=r,end_column=x+5)
			if 0<offset<h-1:
				text,kind=item['rows'][offset-1]
				write(r,x,text,kind=='heading',11,'008833' if kind=='vacancy' else '111111')
			for c in range(x,x+6):
				if offset==0:edge(r,c,'top')
				if offset==h-1:edge(r,c,'bottom')
				edge(r,c,'left') if c==x else None
				edge(r,c,'right') if c==x+5 else None
		if item['children']:
			bus=y[item['depth']+1]-3
			for r in range(top+h,bus+1):edge(r,item['center']-1,'right')
			first,last=item['children'][0]['center'],item['children'][-1]['center']
			for c in range(first,last):edge(bus,c,'bottom')
			for child in item['children']:
				for r in range(bus+1,child['top']):edge(r,child['center']-1,'right')
	if report is not None:
		end_row, report_end_col = _append_organization_report(sheet, report, end_row, end_col, font)
		end_col = max(end_col, report_end_col)
	sheet.print_options.horizontalCentered=True
	sheet.print_area=f"A1:{get_column_letter(end_col)}{end_row}"
	return sheet


def workbook_bytes(root, stamp, split_departments=True, report=None):
	from openpyxl import Workbook
	book=Workbook();book.remove(book.active)
	_diagram(book,root,"完整架构图" if split_departments else "当前层级架构图",stamp,report=report)
	if split_departments:
		def departments(node):
			if node.get("organization_node_type") in ("课","室"):
				yield node
				return
			for child in node.get("children") or []:
				yield from departments(child)
		for node in departments(root):
			_diagram(book,node,node.get("name") or "课室",stamp)
	output=BytesIO();book.save(output)
	return output.getvalue(),book.sheetnames


@frappe.whitelist(methods=["POST"])
def export_excel(company: str, scope: str = "complete", node_ids: str | None = None, search: str = ""):
	from frappe.utils.file_manager import save_file
	from hrms.hr.page.organizational_chart.organizational_chart import get_hybrid_tree
	from hrms.utils.business_department import organization_report_from_tree
	# Whole-company personnel export is restricted to HR administrators.
	frappe.only_for(("System Manager", "HR Manager"))
	frappe.get_doc("Company",company).check_permission("read")
	frappe.has_permission("Employee","read",throw=True)
	payload=get_hybrid_tree(company=company)
	try:
		ids=frappe.parse_json(node_ids) if node_ids else []
		if not isinstance(ids,list) or not all(isinstance(v,str) for v in ids):
			raise ValueError("导出节点格式错误")
		root=select_tree(payload.get("root") or {},scope,ids,search.strip() if scope=='current' else '')
		if not root or not root.get("node_id"):
			raise ValueError("暂无可导出的组织架构")
		report = {
			"title": f"{payload.get('root', {}).get('name') or company}组织报表",
			"columns": ["部门/课别", "编制人数", "现有人数", "空缺人数", "岗位满足率", "备注"],
			**organization_report_from_tree(root),
		}
		data,sheets=workbook_bytes(root,now_datetime().strftime("%Y-%m-%d %H:%M"),scope=='complete',report=report)
	except ValueError as error:
		frappe.throw(str(error))
	label=re.sub(r'[\\/:*?"<>|]', '_',root.get('name') or company)
	file=save_file(f"{label}-Excel组织架构图及报表-{now_datetime():%Y%m%d-%H%M%S}.xlsx",data,None,None,is_private=1)
	return {"file_url":file.file_url,"file_name":file.file_name,"sheet_count":len(sheets),"node_count":sum(1 for _ in _walk(root))}
