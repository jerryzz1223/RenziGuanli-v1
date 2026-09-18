import re
from datetime import date, datetime
from io import BytesIO

import frappe
from frappe import _
from frappe.utils import cint, get_datetime, getdate


@frappe.whitelist()
def get_separation_records(
	company: str | None = None,
	search: str | None = None,
	department: str | None = None,
	start_date: str | None = None,
	end_date: str | None = None,
	year: str | None = None,
	month: str | None = None,
	reason: str | None = None,
	start: int = 0,
	page_length: int = 50,
) -> dict:
	if not frappe.has_permission("Employee", ptype="read"):
		frappe.throw(_("您没有查看离职记录的权限。"), frappe.PermissionError)

	rows, can_read_separations = _collect_records(company, include_separation_details=True)
	filter_options = _build_filter_options(rows)
	rows = _filter_records(
		rows,
		search=search,
		department=department,
		start_date=start_date,
		end_date=end_date,
		year=year,
		month=month,
		reason=reason,
	)
	total = len(rows)
	start = max(cint(start), 0)
	page_length = min(max(cint(page_length) or 50, 1), 100)
	return {
		"rows": rows[start : start + page_length],
		"total": total,
		"start": start,
		"page_length": page_length,
		"can_read_separations": can_read_separations,
		"filter_options": filter_options,
	}


@frappe.whitelist()
def export_separation_records(
	company: str | None = None,
	search: str | None = None,
	department: str | None = None,
	start_date: str | None = None,
	end_date: str | None = None,
	year: str | None = None,
	month: str | None = None,
	reason: str | None = None,
):
	"""Download every field currently available in the separation-record detail view."""
	if not frappe.has_permission("Employee", ptype="read"):
		frappe.throw(_("您没有导出离职记录的权限。"), frappe.PermissionError)
	rows, _can_read_separations = _collect_records(company, include_separation_details=True)
	rows = _filter_records(
		rows,
		search=search,
		department=department,
		start_date=start_date,
		end_date=end_date,
		year=year,
		month=month,
		reason=reason,
	)
	if not rows:
		frappe.throw(_("当前筛选条件下没有可导出的离职记录。"))

	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter

	columns = _export_columns()
	workbook = Workbook()
	sheet = workbook.active
	sheet.title = "离职记录"
	header_fill = PatternFill("solid", fgColor="DDEBF7")
	thin = Side(style="thin", color="B7C9D6")
	border = Border(left=thin, right=thin, top=thin, bottom=thin)
	sheet.append([label for label, _field in columns])
	for cell in sheet[1]:
		cell.fill = header_fill
		cell.font = Font(name="Microsoft YaHei", size=10, bold=True)
		cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
		cell.border = border
	for row in rows:
		sheet.append([_excel_value(row.get(field), field) for _label, field in columns])
		for cell, (_label, field) in zip(sheet[sheet.max_row], columns):
			cell.alignment = Alignment(vertical="top", wrap_text=True)
			cell.border = border
			if field in {"departure_date", "planned_departure_date"} and cell.value:
				cell.number_format = "yyyy-mm-dd"
			elif field in {"application_time", "approval_time", "actual_departure_time", "modified"} and cell.value:
				cell.number_format = "yyyy-mm-dd hh:mm:ss"
	for index, (_label, field) in enumerate(columns, start=1):
		values = [str(_excel_value(row.get(field), field) or "") for row in rows[:500]]
		sheet.column_dimensions[get_column_letter(index)].width = min(
			42, max(12, max([len(_label), *(len(value) for value in values)], default=12) + 2)
		)
	sheet.freeze_panes = "A2"
	sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{sheet.max_row}"
	sheet.sheet_view.showGridLines = False

	output = BytesIO()
	workbook.save(output)
	period = _export_period_label(start_date, end_date, year, month)
	filename = re.sub(r'[\\/:*?"<>|]', "_", f"离职记录_{period}.xlsx")
	frappe.local.response.filename = filename
	frappe.local.response.filecontent = output.getvalue()
	frappe.local.response.type = "binary"


def _collect_records(company=None, include_separation_details=True):
	# Employee is the source of truth for departed staff: scheduled or pending
	# actual-departure times must not appear in the historical record yet.
	employees = _get_departed_employees(company)
	can_read_separations = include_separation_details and frappe.has_permission(
		"Employee Separation", ptype="read"
	)
	department_names = _get_department_display_names(
		[employee.get("department") for employee in employees]
	)

	separations = {}
	if can_read_separations:
		try:
			separations = _get_latest_separations([row.name for row in employees])
		except Exception:
			frappe.log_error(title="Employee separation detail lookup failed")
	rows = [
		_build_record(employee, separations.get(employee.name), department_names)
		for employee in employees
	]
	rows.sort(
		key=lambda row: (str(row.departure_date or ""), str(row.modified or "")),
		reverse=True,
	)
	return rows, bool(can_read_separations)


def _filter_records(rows, search=None, department=None, start_date=None, end_date=None, year=None, month=None, reason=None):
	needle = str(search or "").strip().casefold()
	department = _normalize_filter_text(department)
	start_date = str(start_date or "").strip()[:10]
	end_date = str(end_date or "").strip()[:10]
	year = str(year or "").strip()
	month = str(month or "").strip().zfill(2) if str(month or "").strip() else ""
	reason = _normalize_filter_text(reason)
	filtered = []
	for row in rows:
		departure_date = _record_filter_date(row)
		department_values = {
			_normalize_filter_text(row.department),
			_normalize_filter_text(row.department_display),
			_normalize_filter_text(_strip_department_company_suffix(row.department)),
			_normalize_filter_text(_strip_department_company_suffix(row.department_display)),
		}
		if department and department not in department_values:
			continue
		if start_date and (not departure_date or departure_date < start_date):
			continue
		if end_date and (not departure_date or departure_date > end_date):
			continue
		if year and (not departure_date or not departure_date.startswith(f"{year}-")):
			continue
		if month and (not departure_date or departure_date[5:7] != month):
			continue
		if reason and not _matches_reason(row, reason):
			continue
		if needle and not _matches_search(row, needle):
			continue
		filtered.append(row)
	return filtered


def _matches_reason(row, needle):
	return any(
		needle == _normalize_filter_text(value)
		for value in (
			row.separation_reason_type,
			row.separation_reason,
			row.custom_separation_reason,
			row.separation_reason_display,
			row.separation_reason_detail,
			row.approver_reason_type,
			row.approver_reason,
			row.approver_custom_reason,
			row.approver_reason_display,
			row.approver_reason_detail,
		)
	)


def _normalize_filter_text(value):
	return " ".join(str(value or "").strip().casefold().split())


def _record_filter_date(row):
	"""Return the single date shared by the UI filters and the Excel export."""
	for value in (row.actual_departure_time, row.departure_date, row.planned_departure_date):
		date_text = _date_text(value)
		if date_text:
			return date_text
	return ""


def _date_text(value):
	if not value:
		return ""
	if isinstance(value, (date, datetime)):
		return value.isoformat()[:10]
	return str(value).strip()[:10]


def _build_filter_options(rows):
	departments = {}
	reasons = set()
	years = set()
	for row in rows:
		if row.department:
			departments[str(row.department)] = row.department_display or row.department
		departure_date = _record_filter_date(row)
		if len(departure_date) >= 4 and departure_date[:4].isdigit():
			years.add(departure_date[:4])
		for value in (
			row.separation_reason_type,
			row.separation_reason,
			row.custom_separation_reason,
			row.separation_reason_display,
			row.approver_reason_type,
			row.approver_reason,
			row.approver_custom_reason,
			row.approver_reason_display,
		):
			if value:
				reasons.add(str(value))
	return {
		"departments": [{"value": value, "label": label} for value, label in sorted(departments.items(), key=lambda item: item[1])],
		"reasons": sorted(reasons),
		"years": sorted(years, reverse=True),
	}


def _export_columns():
	return [
		("员工内部编号", "employee"),
		("员工姓名", "employee_name"),
		("工号", "employee_code"),
		("公司", "company"),
		("部门", "department_display"),
		("岗位", "designation"),
		("实际离职日期", "departure_date"),
		("拟离职日期", "planned_departure_date"),
		("离职申请时间", "application_time"),
		("申请操作人", "application_operator"),
		("离职审批时间", "approval_time"),
		("审批操作人", "approval_operator"),
		("实际离职时间", "actual_departure_time"),
		("实际离职操作人", "actual_departure_operator"),
		("离职单号", "separation_name"),
		("离职单状态", "separation_status"),
		("员工自述原因分类", "separation_reason_type"),
		("员工自述离职原因", "separation_reason_display"),
		("员工自述原因原值", "separation_reason"),
		("员工自述自定义原因", "custom_separation_reason"),
		("员工自述详细原因", "separation_reason_detail"),
		("审批确认原因分类", "approver_reason_type"),
		("审批确认离职原因", "approver_reason_display"),
		("审批确认原因原值", "approver_reason"),
		("审批确认自定义原因", "approver_custom_reason"),
		("审批确认详细原因", "approver_reason_detail"),
		("离职面谈", "exit_interview"),
		("记录更新时间", "modified"),
	]


def _excel_value(value, field=None):
	if value is None:
		return ""
	if field in {"departure_date", "planned_departure_date"}:
		if isinstance(value, datetime):
			return value.date()
		if isinstance(value, date):
			return value
		try:
			return getdate(value)
		except Exception:
			return str(value)
	if field in {"application_time", "approval_time", "actual_departure_time", "modified"}:
		if isinstance(value, datetime):
			return value
		try:
			return get_datetime(value)
		except Exception:
			return str(value)
	if isinstance(value, (str, int, float, bool)):
		return value
	return str(value)


def _export_period_label(start_date, end_date, year, month):
	if year and month:
		return f"{year}年{str(month).zfill(2)}月"
	if year:
		return f"{year}年"
	if start_date or end_date:
		return f"{start_date or '起始'}-{end_date or '截至'}"
	return "全部"


def _employee_fields():
	meta = frappe.get_meta("Employee")
	candidates = (
		"name",
		"employee_name",
		"company",
		"department",
		"designation",
		"relieving_date",
		"status",
		"modified",
		"custom_employee_code",
	)
	return [fieldname for fieldname in candidates if _meta_has_field(meta, fieldname)]


def _get_departed_employees(company=None):
	fields = _employee_fields()

	filters = {}
	if company and "company" in fields:
		filters["company"] = company
	order_fields = [fieldname for fieldname in ("relieving_date", "modified") if fieldname in fields]

	rows = frappe.get_list(
		"Employee",
		filters=filters,
		fields=fields,
		order_by=", ".join(f"{fieldname} desc" for fieldname in order_fields) or "name desc",
		limit_page_length=0,
	)
	return [row for row in rows if _is_departed_employee(row)]


def _is_departed_employee(employee):
	return employee.get("status") == "Left"


def _get_latest_separations(employee_names):
	if not employee_names:
		return {}
	meta = frappe.get_meta("Employee Separation")
	if not _meta_has_field(meta, "employee"):
		return {}
	fields = [
		fieldname
		for fieldname in (
			"name",
			"employee",
			"docstatus",
			"boarding_begins_on",
			"boarding_status",
			"applied_on",
			"applied_by",
			"approved_on",
			"approved_by",
			"departed_on",
			"departed_by",
			"separation_reason_type",
			"separation_reason",
			"custom_separation_reason",
			"separation_reason_detail",
			"approver_reason_type",
			"approver_reason",
			"approver_custom_reason",
			"approver_reason_detail",
			"exit_interview",
			"modified",
		)
		if _meta_has_field(meta, fieldname)
	]
	order_fields = [
		fieldname
		for fieldname in ("employee", "docstatus", "boarding_begins_on", "modified")
		if fieldname in fields
	]

	rows = frappe.get_list(
		"Employee Separation",
		filters={
			"employee": ["in", employee_names],
			"docstatus": 1,
			"boarding_status": "Completed",
		},
		fields=fields,
		order_by=", ".join(
			f"{fieldname} {'asc' if fieldname == 'employee' else 'desc'}"
			for fieldname in order_fields
		),
		limit_page_length=0,
	)
	latest = {}
	for row in rows:
		latest.setdefault(row.employee, row)
	return latest


def _build_record(employee, separation=None, department_names=None):
	department = employee.get("department") or ""
	return frappe._dict(
		{
			"employee": employee.get("name"),
			"employee_code": employee.get("custom_employee_code") or "",
			"employee_name": employee.get("employee_name") or "",
			"company": employee.get("company") or "",
			"department": department,
			"department_display": _department_display_name(department, department_names),
			"designation": employee.get("designation") or "",
			"departure_date": employee.get("relieving_date"),
			"planned_departure_date": separation.get("boarding_begins_on") if separation else None,
			"application_time": separation.get("applied_on") if separation else None,
			"application_operator": separation.get("applied_by") if separation else None,
			"approval_time": separation.get("approved_on") if separation else None,
			"approval_operator": separation.get("approved_by") if separation else None,
			"actual_departure_time": separation.get("departed_on") if separation else None,
			"actual_departure_operator": separation.get("departed_by") if separation else None,
			"separation_name": separation.get("name") if separation else None,
			"separation_status": separation.get("boarding_status") if separation else None,
			"separation_reason_type": (
				separation.get("separation_reason_type") if separation else ""
			),
			"separation_reason": separation.get("separation_reason") if separation else "",
			"custom_separation_reason": (
				separation.get("custom_separation_reason") if separation else ""
			),
			"separation_reason_display": _separation_reason_display(separation),
			"separation_reason_detail": (
				separation.get("separation_reason_detail") if separation else ""
			),
			"approver_reason_type": separation.get("approver_reason_type") if separation else "",
			"approver_reason": separation.get("approver_reason") if separation else "",
			"approver_custom_reason": (
				separation.get("approver_custom_reason") if separation else ""
			),
			"approver_reason_detail": (
				separation.get("approver_reason_detail") if separation else ""
			),
			"approver_reason_display": _approver_reason_display(separation),
			"exit_interview": separation.get("exit_interview") if separation else None,
			"modified": separation.get("modified") if separation else employee.get("modified"),
		}
	)


def _separation_reason_display(separation):
	return _reason_display(
		separation,
		"separation_reason_type",
		"separation_reason",
		"custom_separation_reason",
	)


def _approver_reason_display(separation):
	return _reason_display(
		separation,
		"approver_reason_type",
		"approver_reason",
		"approver_custom_reason",
	)


def _reason_display(separation, type_field, reason_field, custom_field):
	if not separation:
		return ""

	reason_type = str(separation.get(type_field) or "").strip()
	if reason_type == "自定义":
		return str(separation.get(custom_field) or "").strip()
	return str(separation.get(reason_field) or "").strip()


def _strip_department_company_suffix(value):
	text = str(value or "").strip()
	return re.sub(r"\s+-\s+[^-]+$", "", text).strip()


def _get_department_display_names(department_values):
	values = sorted({value for value in department_values if value})
	if not values:
		return {}

	try:
		rows = frappe.get_all(
			"Department",
			filters={"name": ["in", values]},
			fields=["name", "department_name"],
			limit_page_length=0,
		)
	except Exception:
		# Department enrichment must not make the separation archive unavailable.
		frappe.log_error(
			title="Employee Separation Records Department Lookup Failed",
			message=frappe.get_traceback(),
		)
		return {}

	return {row.name: row.department_name for row in rows if row.get("department_name")}


def _department_display_name(value, department_names=None):
	if not value:
		return ""

	return (department_names or {}).get(value) or _strip_department_company_suffix(value)


def _meta_has_field(meta, fieldname):
	return fieldname in {"name", "docstatus", "modified"} or bool(meta.has_field(fieldname))


def _matches_search(row, needle):
	values = (
		row.employee_code,
		row.employee_name,
		row.department_display,
		row.designation,
		row.separation_reason_type,
		row.separation_reason_display,
		row.separation_reason_detail,
		row.approver_reason_type,
		row.approver_reason_display,
		row.approver_reason_detail,
	)
	return any(needle in str(value or "").casefold() for value in values)
