"""Pure parsing helpers for employee relationship workbooks.

This module deliberately has no Frappe dependency so the source workbook can
be validated without writing employee relationship documents.
"""

from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook


RELATIONSHIP_CATEGORIES = (
	"直系亲属",
	"旁系亲属",
	"姻亲",
	"男女朋友",
	"同学",
	"前同事",
	"朋友",
	"同村",
	"其他",
)


def text(value):
	if value is None:
		return ""
	if isinstance(value, datetime):
		return value.strftime("%Y-%m-%d")
	if isinstance(value, date):
		return value.isoformat()
	return " ".join(str(value).split())


def compact(value):
	return re.sub(r"\s+", "", text(value).replace("（", "(").replace("）", ")"))


def workbook_digest(content):
	return hashlib.sha256(content).hexdigest()


def preview_token(company, digest, rows):
	payload = {
		"company": text(company),
		"digest": digest,
		"rows": [
			{
				"sheet": row["source_sheet"],
				"row": row["source_row"],
				"a": row["employee_a_identity"],
				"b": row["employee_b_identity"],
				"relationship": row["relationship"],
			}
			for row in rows
		],
	}
	encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
	return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def source_pair_key(row):
	return " <> ".join(sorted((row["employee_a_identity"], row["employee_b_identity"])))


def summarize_source_conflicts(rows):
	grouped = defaultdict(list)
	for row in rows:
		grouped[source_pair_key(row)].append(row)
	conflicts = []
	duplicate_count = 0
	for pair_key, items in grouped.items():
		if len(items) < 2:
			continue
		categories = sorted({row["relationship"] for row in items})
		if len(categories) == 1:
			duplicate_count += len(items) - 1
			continue
		conflicts.append(
			{
				"pair_key": pair_key,
				"employee_a_name": items[0]["employee_a_name"],
				"employee_b_name": items[0]["employee_b_name"],
				"categories": categories,
				"occurrences": [
					{
						"source_sheet": row["source_sheet"],
						"source_row": row["source_row"],
						"relationship": row["relationship"],
					}
					for row in items
				],
			}
		)
	return conflicts, duplicate_count


def _workbook(content):
	if hasattr(content, "read"):
		content = content.read()
	if isinstance(content, str):
		with open(content, "rb") as source:
			content = source.read()
	return load_workbook(BytesIO(content), data_only=True, read_only=False)


def _header(sheet):
	for row_number in range(1, min(sheet.max_row, 15) + 1):
		values = [compact(sheet.cell(row_number, column).value) for column in range(1, sheet.max_column + 1)]
		name_columns = [index + 1 for index, value in enumerate(values) if "姓名" in value and "更新" not in value]
		department_columns = [index + 1 for index, value in enumerate(values) if value == "部门"]
		relationship_columns = [index + 1 for index, value in enumerate(values) if value in ("关系", "员工关系大类", "关系大类")]
		if len(name_columns) >= 2 and len(department_columns) >= 2 and relationship_columns:
			return row_number, values, name_columns[:2], department_columns[:2], relationship_columns[0]
	raise ValueError(f"工作表“{sheet.title}”未识别到双方姓名、部门和关系表头")


def _code_columns(headers, name_columns):
	explicit = {}
	for index, value in enumerate(headers, start=1):
		if value in ("员工一工号", "员工1工号", "人员一工号", "人员1工号"):
			explicit[0] = index
		elif value in ("员工二工号", "员工2工号", "人员二工号", "人员2工号"):
			explicit[1] = index
	if len(explicit) == 2:
		return [explicit[0], explicit[1]]
	code_columns = [index + 1 for index, value in enumerate(headers) if "工号" in value]
	if len(code_columns) >= 2:
		return code_columns[:2]
	# A source file without company codes remains readable, but preview must map
	# every identity to a company-scoped code before import.
	return [None for _ in name_columns]


def _cell_text(sheet, row_number, column):
	return text(sheet.cell(row_number, column).value) if column else ""


def _identity(name, department, code):
	return f"工号:{code}" if code else f"姓名部门:{name}|{department}"


def parse_employee_relationship_workbook(content):
	workbook = _workbook(content)
	rows = []
	parsed_sheets = []
	for sheet in workbook.worksheets:
		if "人员关系" not in compact(sheet.title) and "员工关系" not in compact(sheet.title):
			continue
		header_row, headers, name_columns, department_columns, relationship_column = _header(sheet)
		code_columns = _code_columns(headers, name_columns)
		parsed_sheets.append(sheet.title)
		for row_number in range(header_row + 1, sheet.max_row + 1):
			a_name = _cell_text(sheet, row_number, name_columns[0])
			b_name = _cell_text(sheet, row_number, name_columns[1])
			relationship = _cell_text(sheet, row_number, relationship_column)
			if not any((a_name, b_name, relationship)):
				continue
			a_department = _cell_text(sheet, row_number, department_columns[0])
			b_department = _cell_text(sheet, row_number, department_columns[1])
			a_code = _cell_text(sheet, row_number, code_columns[0])
			b_code = _cell_text(sheet, row_number, code_columns[1])
			errors = []
			if not a_name:
				errors.append("员工一姓名为空")
			if not b_name:
				errors.append("员工二姓名为空")
			if relationship not in RELATIONSHIP_CATEGORIES:
				errors.append(
					f"关系大类“{relationship or '空'}”不在允许范围：{'、'.join(RELATIONSHIP_CATEGORIES)}"
				)
			if a_name and b_name and a_name == b_name and a_department == b_department:
				errors.append("员工一与员工二疑似为同一人")
			row = {
				"source_sheet": sheet.title,
				"source_row": row_number,
				"source_status": "离职" if "离职" in sheet.title else "在职",
				"employee_a_name": a_name,
				"employee_a_department": a_department,
				"employee_a_code": a_code,
				"employee_b_name": b_name,
				"employee_b_department": b_department,
				"employee_b_code": b_code,
				"relationship": relationship,
				"employee_a_identity": _identity(a_name, a_department, a_code),
				"employee_b_identity": _identity(b_name, b_department, b_code),
				"errors": errors,
			}
			rows.append(row)
	if not parsed_sheets:
		raise ValueError("未找到名称包含“人员关系”或“员工关系”的工作表")
	if not rows:
		raise ValueError("人员关系工作表中未读取到数据行")
	return {"sheet_names": parsed_sheets, "rows": rows}
