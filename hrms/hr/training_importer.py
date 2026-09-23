"""Pure workbook parsing helpers for the Yongxin training import.

This module deliberately has no Frappe dependency.  Preview and tests can read
the two source workbooks without writing business documents.
"""

from __future__ import annotations

import hashlib
import json
import re
from collections import OrderedDict
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


PLAN_SHEET = "2026年计划总表"
RECORD_SHEET = "2026年安全培训教育记录表"


def text(value):
	if value is None:
		return ""
	if isinstance(value, datetime):
		return value.strftime("%Y-%m-%d")
	if isinstance(value, date):
		return value.isoformat()
	return str(value).strip()


def compact(value):
	return re.sub(r"\s+", "", text(value).replace("（", "(").replace("）", ")"))


def number(value):
	if value in (None, ""):
		return None
	try:
		return float(value)
	except (TypeError, ValueError):
		return None


def workbook_digest(content):
	return hashlib.sha256(content).hexdigest()


def stable_key(prefix, payload):
	encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
	return f"{prefix}-{hashlib.sha256(encoded.encode('utf-8')).hexdigest()[:24]}"


def _workbook(content):
	if hasattr(content, "read"):
		content = content.read()
	if isinstance(content, str):
		with open(content, "rb") as source:
			content = source.read()
	return load_workbook(BytesIO(content), data_only=True, read_only=False)


def parse_dates(value, default_year=2026):
	"""Return every real date represented by one source cell.

	The actual workbook contains both normal Excel dates and compact multi-date
	values such as ``260825\n/260831``.  All original text is retained separately.
	"""
	if value in (None, ""):
		return []
	if isinstance(value, datetime):
		return [value.date()]
	if isinstance(value, date):
		return [value]
	if isinstance(value, (int, float)):
		try:
			return [from_excel(value).date()]
		except (TypeError, ValueError, OverflowError):
			return []

	raw = text(value)
	results = []
	patterns = (
		r"(?<!\d)(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})(?:日)?",
		r"(?<!\d)(\d{2})(\d{2})(\d{2})(?!\d)",
	)
	for pattern_index, pattern in enumerate(patterns):
		for match in re.finditer(pattern, raw):
			parts = [int(item) for item in match.groups()]
			year = parts[0] if pattern_index == 0 else 2000 + parts[0]
			month, day = parts[1], parts[2]
			try:
				candidate = date(year, month, day)
			except ValueError:
				continue
			if candidate not in results:
				results.append(candidate)
	if not results:
		match = re.fullmatch(r"(\d{1,2})[-/.月](\d{1,2})(?:日)?", raw)
		if match:
			try:
				results.append(date(default_year, int(match.group(1)), int(match.group(2))))
			except ValueError:
				pass
	return results


def _merged_top_labels(sheet, start_column, end_column):
	labels = []
	current = ""
	for column in range(start_column, end_column + 1):
		value = text(sheet.cell(3, column).value)
		if value:
			current = value
		sub = text(sheet.cell(4, column).value)
		labels.append(" / ".join(item for item in (current, sub) if item))
	return labels


def parse_plan_workbook(content):
	workbook = _workbook(content)
	if PLAN_SHEET not in workbook.sheetnames:
		raise ValueError(f"缺少工作表：{PLAN_SHEET}")
	sheet = workbook[PLAN_SHEET]
	required = {
		2: "部门",
		3: "分类",
		4: "培训类型",
		5: "培训内容",
		6: "内/外",
		12: "预计上课时间(月份)",
	}
	missing = [label for column, label in required.items() if compact(sheet.cell(3, column).value) != compact(label)]
	if missing:
		raise ValueError("计划总表表头不匹配：" + "、".join(missing))

	title = text(sheet.cell(2, 2).value)
	year_match = re.search(r"(20\d{2})", title)
	plan_year = int(year_match.group(1)) if year_match else 2026
	audience_labels = _merged_top_labels(sheet, 16, 33)
	rows = []
	for row_number in range(5, sheet.max_row + 1):
		content_name = text(sheet.cell(row_number, 5).value)
		if not content_name:
			continue
		audience_matrix = []
		for offset, column in enumerate(range(16, 34)):
			value = text(sheet.cell(row_number, column).value)
			if value:
				audience_matrix.append({"unit": audience_labels[offset], "requirement": value})
		actual_values = [sheet.cell(row_number, column).value for column in range(13, 16)]
		actual_dates = []
		for value in actual_values:
			actual_dates.extend(day.isoformat() for day in parse_dates(value, plan_year))
		row = {
			"source_row": row_number,
			"plan_year": plan_year,
			"department": text(sheet.cell(row_number, 2).value),
			"classification": text(sheet.cell(row_number, 3).value),
			"training_type": text(sheet.cell(row_number, 4).value),
			"content": content_name,
			"internal_external": text(sheet.cell(row_number, 6).value),
			"course_hours": number(sheet.cell(row_number, 7).value),
			"convener": text(sheet.cell(row_number, 8).value),
			"convener_department": text(sheet.cell(row_number, 9).value),
			"location": text(sheet.cell(row_number, 10).value),
			"target": text(sheet.cell(row_number, 11).value),
			"planned_month": text(sheet.cell(row_number, 12).value),
			"actual_dates": list(OrderedDict.fromkeys(actual_dates)),
			"actual_dates_text": "；".join(text(value) for value in actual_values if text(value)),
			"audience_matrix": audience_matrix,
			"remarks": text(sheet.cell(row_number, 34).value),
		}
		row["source_key"] = stable_key(
			"TRAIN-PLAN",
			{key: value for key, value in row.items() if key != "source_key"},
		)
		row["errors"] = [
			label
			for label, value in (("部门为空", row["department"]), ("培训类型为空", row["training_type"]), ("培训内容为空", row["content"]))
			if not value
		]
		rows.append(row)
	return {"sheet_name": sheet.title, "plan_year": plan_year, "rows": rows}


def _find_record_header(sheet):
	required = {"序号", "实际上课时间", "部门", "姓名", "培训内容"}
	for row_number in range(1, min(sheet.max_row, 12) + 1):
		values = [compact(sheet.cell(row_number, column).value) for column in range(1, min(sheet.max_column, 30) + 1)]
		if required.issubset(set(values)):
			return row_number, {value: index + 1 for index, value in enumerate(values) if value}
	raise ValueError("实际记录表未识别到标准表头")


def _record_sheet(workbook):
	"""Select the actual-record sheet by structure, not by one historic title."""
	ordered_names = ([RECORD_SHEET] if RECORD_SHEET in workbook.sheetnames else []) + [
		name for name in workbook.sheetnames if name != RECORD_SHEET
	]
	for sheet_name in ordered_names:
		sheet = workbook[sheet_name]
		try:
			_find_record_header(sheet)
		except ValueError:
			continue
		return sheet
	raise ValueError("未找到包含序号、实际上课时间、部门、姓名和培训内容的实际记录工作表")


def parse_record_workbook(content):
	workbook = _workbook(content)
	sheet = _record_sheet(workbook)
	header_row, columns = _find_record_header(sheet)
	aliases = {
		"serial": "序号",
		"month": "月份",
		"actual_date": "实际上课时间",
		"department": "部门",
		"employee_name": "姓名",
		"course_type": "课程类型",
		"content": "培训内容",
		"owner_department": "课程归属部门",
		"internal_external": "内/外",
		"courseware": "课件方式",
		"hours": "课时",
		"study_hours": "学时",
		"instructor": "授课人",
		"location": "地点",
		"target": "培训对象",
		"score": "成绩",
		"remarks": "备注(评价标准)",
	}
	missing = [label for label in aliases.values() if compact(label) not in columns]
	if missing:
		raise ValueError("实际记录表缺少列：" + "、".join(missing))
	resolved = {key: columns[compact(label)] for key, label in aliases.items()}
	rows = []
	for row_number in range(header_row + 1, sheet.max_row + 1):
		serial = sheet.cell(row_number, resolved["serial"]).value
		content_name = text(sheet.cell(row_number, resolved["content"]).value)
		if serial in (None, "") and not content_name:
			continue
		date_value = sheet.cell(row_number, resolved["actual_date"]).value
		dates = parse_dates(date_value, 2026)
		score_value = sheet.cell(row_number, resolved["score"]).value
		row = {
			"source_row": row_number,
			"source_serial": text(serial),
			"month": text(sheet.cell(row_number, resolved["month"]).value),
			"actual_date_text": text(date_value),
			"actual_dates": [item.isoformat() for item in dates],
			"department": text(sheet.cell(row_number, resolved["department"]).value),
			"employee_name": text(sheet.cell(row_number, resolved["employee_name"]).value),
			"course_type": text(sheet.cell(row_number, resolved["course_type"]).value),
			"content": content_name,
			"owner_department": text(sheet.cell(row_number, resolved["owner_department"]).value),
			"internal_external": text(sheet.cell(row_number, resolved["internal_external"]).value),
			"courseware": text(sheet.cell(row_number, resolved["courseware"]).value),
			"hours": number(sheet.cell(row_number, resolved["hours"]).value),
			"study_hours": number(sheet.cell(row_number, resolved["study_hours"]).value),
			"instructor": text(sheet.cell(row_number, resolved["instructor"]).value),
			"location": text(sheet.cell(row_number, resolved["location"]).value),
			"target": text(sheet.cell(row_number, resolved["target"]).value),
			"score": number(score_value),
			"score_text": text(score_value),
			"remarks": text(sheet.cell(row_number, resolved["remarks"]).value),
		}
		row["errors"] = []
		for label, value in (
			("实际上课时间为空或无法识别", row["actual_dates"]),
			("部门为空", row["department"]),
			("姓名为空", row["employee_name"]),
			("培训内容为空", row["content"]),
			("课时为空或不是数字", row["hours"]),
		):
			if value in (None, "", []):
				row["errors"].append(label)
		row["identity_key"] = f"{row['employee_name']}|{row['department']}"
		rows.append(row)

	grouped = OrderedDict()
	for row in rows:
		payload = {
			"actual_date_text": row["actual_date_text"],
			"actual_dates": row["actual_dates"],
			"content": row["content"],
			"course_type": row["course_type"],
			"owner_department": row["owner_department"],
			"internal_external": row["internal_external"],
			"courseware": row["courseware"],
			"hours": row["hours"],
			"instructor": row["instructor"],
			"location": row["location"],
			"target": row["target"],
		}
		key = stable_key("TRAIN-EVENT", payload)
		if key not in grouped:
			grouped[key] = {**payload, "source_key": key, "source_rows": [], "participants": [], "errors": []}
		grouped[key]["source_rows"].append(row["source_row"])
		grouped[key]["participants"].append(row)
		grouped[key]["errors"].extend(row["errors"])
	for event in grouped.values():
		event["errors"] = list(OrderedDict.fromkeys(event["errors"]))
		event["source_sheet"] = sheet.title
	return {"sheet_name": sheet.title, "header_row": header_row, "rows": rows, "events": list(grouped.values())}


def preview_token(company, plan_digest, record_digest, plan_rows, record_rows):
	payload = {
		"company": text(company),
		"plan_digest": plan_digest,
		"record_digest": record_digest,
		"plan_keys": [row["source_key"] for row in plan_rows],
		"record_rows": [[row["source_row"], row["source_serial"], row["identity_key"]] for row in record_rows],
	}
	return hashlib.sha256(json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")).hexdigest()
