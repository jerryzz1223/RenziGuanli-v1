"""Restore the verified Q3 diagram's groups into the shared live organization.

Preview by default. Source coordinates are identities, not names (three production
groups and three IPQC groups share labels). No Employee or Department is written.
"""
import json
import re
from pathlib import Path

import frappe
from frappe.utils import now_datetime

# Explicit branches checked against the source connector bands at rows 16-17.
# Spatial nearest-neighbour is wrong for 保养组, 备品组 and the 工程课助理.
TEAMS = {
	"连续课": "C H M S X", "设备课": "AC AI", "品管课": "AN AS AX BC",
	"工程课": "BH BM BR", "总办室": "CB CG CL CP CS", "品保课": "CW DL DQ",
	"药水课": "DV EB", "生管课": "EI EN", "量试组": "ES",
	"行政课": "FA FF FK FP", "环安课": "FU FZ", "资财课": "GE GJ GO GT",
}
SOURCE_DOC = "组织架构图260626.xlsx"
DEPARTMENT_CELLS = {"连续课": "J14", "设备课": "AE14", "品管课": "AS14", "工程课": "BM14", "总办室": "CI14", "品保课": "DB14", "药水课": "DW14", "生管课": "EI14", "量试组": "ES14", "行政课": "FG14", "环安课": "FU14", "资财课": "GH14"}


def parse_sheet(sheet):
	if sheet.title != "26Q3组织架构图" or not str(sheet["J13"].value).startswith("连续课"):
		raise ValueError("原表版式不匹配，未导入")
	merged = {sheet.cell(r.min_row, r.min_col).coordinate: r for r in sheet.merged_cells.ranges}
	groups = {}
	for dept, columns in TEAMS.items():
		for column in columns.split():
			cell = column + "18"
			groups[cell] = dept
	nodes = []
	def add(cell, dept, parent, kind, name=None, body=None):
		raw = str(sheet[cell].value or "").strip()
		lines = [x.strip() for x in raw.splitlines() if x.strip()]
		if not lines:
			return
		name = name or lines[0]
		staff_text = raw if body is None else "\n".join(body)
		body = lines[1:] if body is None else body
		role = "组长" if kind == "组" else "线长" if kind == "线" and name == "直线级" else ""
		bindings = []
		for line in body:
			line = re.sub(r"TBA\s*[*×xX]?\s*\d*", "", line, flags=re.I).strip()
			if not line:
				continue
			if line in {"组长", "班长", "副组长"}:
				role = line
				continue
			if line in {"（代）", "(代)", "（兼）", "(兼)"}:
				if bindings: bindings[-1]["role"] += "（代）" if "代" in line else "（兼）"
				continue
			for token in re.split(r"\s+|[、，,]", line):
				m = re.fullmatch(r"([\u4e00-\u9fff]{2,4})([（(](?:代|兼)[）)])?", token)
				if not m:
					raise ValueError(f"{cell} 无法识别人员内容: {token}")
				bindings.append({"source_name": m[1], "role": role + (m[2] or "")})
		nodes.append({"cell": cell, "department_label": dept, "parent_cell": parent,
			"kind": kind, "name": name, "bindings": bindings,
			"source_vacancies": sum(int(n) for n in re.findall(r"TBA\s*[*×xX]?\s*(\d+)", staff_text, re.I))})
	for cell, dept in groups.items():
		add(cell, dept, None, "组")
	# All lines and employee columns use their actual containing group interval.
	for row in (21, 23):
		for cell, region in sorted(merged.items(), key=lambda x: x[1].min_col):
			if region.min_row != row or not sheet[cell].value:
				continue
			if cell == "BW23":
				dept, parent = "工程课", None
			elif cell in {"DG21", "DG23"}:
				# FQC's box extends one column beyond 进出货组; its connector
				# still returns to that group, not neighbouring 试验组.
				dept, parent = "品保课", "CW18"
			else:
				matches = [g for g in groups if merged[g].min_col <= region.min_col and region.max_col <= merged[g].max_col]
				if len(matches) != 1:
					raise ValueError(f"{cell} 未找到唯一上级组")
				parent = matches[0]
				dept = groups[parent]
			if row == 23:
				line = next((n for n in nodes if n["cell"].endswith("21") and merged[n["cell"]].min_col == region.min_col and merged[n["cell"]].max_col == region.max_col), None)
				if line: parent = line["cell"]
				# The three long lists are merged all the way through row 31;
				# reading the merged origin retains every name, including 马勇.
				people_cell = sheet.cell(24, region.min_col).coordinate
				body = [x.strip() for x in str(sheet[people_cell].value or "").splitlines() if x.strip()]
				add(cell, dept, parent, "岗位", body=body)
			else:
				add(cell, dept, parent, "线")
	for column in ("AN", "AS", "AX"):
		add(column + "26", "品管课", column + "23", "室", body=str(sheet[column + "27"].value or "").splitlines())
	add("FA25", "行政课", "FA18", "岗位")
	add("FA26", "行政课", "FA18", "岗位", name="厂区维护", body=str(sheet["FA26"].value or "").strip().splitlines()[2:])
	return nodes


def run(apply=False, commit=True):
	from openpyxl import load_workbook
	from hrms.hr.page.organizational_chart import organizational_chart as chart
	from hrms.api.organization_template import reconcile_bindings
	from hrms.api.organization_roster_sync import reconcile
	frappe.only_for("System Manager")
	company = "永新"
	source = chart._resolve_yongxin_org_workbook()
	if not source or source.name != SOURCE_DOC:
		raise ValueError("缺少已核对的组织架构图260626.xlsx，未导入")
	wb = load_workbook(source, data_only=True)
	try:
		sheet = wb["26Q3组织架构图"]
		plan = parse_sheet(sheet)
		staffing = {dept: int(re.search(r"(?:编制/实际|编制)[:：]?(\d+)/", str(sheet[cell].value))[1]) for dept, cell in DEPARTMENT_CELLS.items()}
		leadership = {}
		for dept, cell in DEPARTMENT_CELLS.items():
			text = str(sheet["ES18" if dept == "量试组" else cell.replace("14", "13")].value or "")
			entries = []
			role, slot = "负责人", "primary"
			for line in [s.strip() for s in text.splitlines() if s.strip()][1:]:
				if line in {"组长", "班长"}: role = line; continue
				if "：" in line or ":" in line:
					role, line = re.split("[:：]", line, maxsplit=1)
					slot = "proxy" if role == "代理人" else "primary"
				line = re.sub(r"TBA\s*[*×xX]?\s*\d*", "", line, flags=re.I).strip()
				for name in line.split():
					if not re.fullmatch(r"[\u4e00-\u9fff]{2,4}", name): raise ValueError(f"{dept}负责人无法识别")
					entries.append({"source_name": name, "role": role, "slot": slot, "display_only": slot == "proxy"})
			leadership[dept] = entries
	finally: wb.close()
	manual = chart._get_manual_organization_records(company)
	staff = chart._get_active_employees(company)
	departments = {d.department_name: d.name for d in chart._get_departments(company)}
	units = {n.manual_config.get("department"): n for n in manual["nodes"] if chart.whole_department(n.manual_config)}
	existing = {n.manual_config["template_source_cell"]: n for n in manual["nodes"] if n.manual_config.get("template_source_cell")}
	matched, issues = set(), []
	for item in plan:
		item["department"] = departments[item["department_label"]]
		config = {"department": item["department"], "template_bindings": item["bindings"]}
		item["resolved"] = reconcile_bindings(config, staff, allow_name_match=True)
		matched.update(item["resolved"]["assigned_employees"])
		issues.extend({"cell": item["cell"], **b} for b in item["resolved"]["template_bindings"] if b.get("issue"))
	result = {"source": source.name, "groups_and_roles": len(plan), "matched_active_people": len(matched),
		"unresolved_references": len(issues), "to_create": sum(i["cell"] not in existing for i in plan)}
	if not apply:
		return {**result, "issues": issues}
	frappe.db.sql("select name from `tabOrganization Structure Version` where name=%s for update", manual["version"] if isinstance(manual.get("version"), str) else chart._get_manual_organization_version(company))
	backup = Path(frappe.get_site_path("private", "backups", "organization-template-" + now_datetime().strftime("%Y%m%d-%H%M%S-%f") + ".json"))
	backup.parent.mkdir(parents=True, exist_ok=True)
	backup.write_text(frappe.as_json({"before": [frappe.get_doc("Organization Node", n.name).as_dict() for n in manual["nodes"]], "plan": plan}), encoding="utf-8")
	backup.chmod(0o600)
	for label, bindings in leadership.items():
		unit = units[departments[label]]
		config = unit.manual_config
		if config.get("template_leadership") or (not config.get("roster_auto_sync") and any(config.get(k) for k in ("primary_employee", "proxy_employee", "role_title"))):
			continue
		resolved = reconcile_bindings({"department": departments[label], "template_bindings": bindings}, staff, allow_name_match=True)
		doc = frappe.get_doc("Organization Node", unit.name)
		config = chart._manual_node_config(doc.source_text)
		config.update(template_leadership=True, template_bindings=resolved["template_bindings"], template_source_document=source.name, roster_auto_sync=True)
		doc.source_text = json.dumps(config, ensure_ascii=False)
		doc.save()
	# Fill only previously unconfigured plans. Live counts continue reading Employee.
	for label, planned in staffing.items():
		unit = units[departments[label]]
		if not unit.manual_config.get("planned_headcount_set", unit.planned_headcount > 0):
			doc = frappe.get_doc("Organization Node", unit.name)
			config = chart._manual_node_config(doc.source_text)
			config["planned_headcount_set"] = True
			doc.planned_headcount = planned
			doc.source_text = json.dumps(config, ensure_ascii=False)
			doc.save()
	# The complete source confirms the previously unplaced branches. Preserve any
	# already manually placed unit. Management stays in the reviewed combined row.
	managers = {n.manual_config.get("manager_name"): n for n in manual["nodes"] if chart._manual_node_kind(n) == "分管"}
	management = next(n for n in manual["nodes"] if chart._manual_node_kind(n) == "管理层")
	for manager_name, labels in {"陈文萍": ["品保课", "药水课"], "逯瑜": ["生管课", "量试组"], "杨玉婷": ["行政课", "环安课", "资财课"]}.items():
		if manager_name not in managers:
			matches = [e for e in staff if e.employee_name == manager_name]
			if len(matches) != 1: raise ValueError(f"分管 {manager_name} 未能唯一匹配，未导入")
			saved = chart.save_manual_organization_node("分管", company=company, display_name="分管四", manager_employee=matches[0].name,
				parent_node=management.name, role_title="分管", assignment_mode="正式")
			managers[manager_name] = frappe.get_doc("Organization Node", saved["name"])
		for label in labels:
			unit = units[departments[label]]
			if unit.parent_node or not unit.manual_config.get("reporting_scope_pending"): continue
			doc = frappe.get_doc("Organization Node", unit.name)
			doc.parent_node = managers[manager_name].name
			config = chart._manual_node_config(doc.source_text)
			config["reporting_scope_pending"] = False
			doc.source_text = json.dumps(config, ensure_ascii=False)
			doc.save()
	pending = [i for i in plan if i["cell"] not in existing]
	while pending:
		ready = [i for i in pending if not i["parent_cell"] or i["parent_cell"] in existing]
		if not ready: raise ValueError("原表层级有循环或缺少上级")
		for item in ready:
			parent = existing[item["parent_cell"]].name if item["parent_cell"] else units[item["department"]].name
			saved = chart.save_manual_organization_node(item["kind"], display_name=item["name"], company=company,
				department=item["department"], parent_node=parent, roster_subset=1,
				assigned_employees=item["resolved"]["assigned_employees"], assignment_mode="正式")
			doc = frappe.get_doc("Organization Node", saved["name"])
			config = chart._manual_node_config(doc.source_text)
			config.update(item["resolved"], template_source_cell=item["cell"], template_source_document=source.name,
				template_source_vacancies=item["source_vacancies"], roster_auto_sync=True)
			doc.source_text = json.dumps(config, ensure_ascii=False)
			doc.save()
			existing[item["cell"]] = doc
			pending.remove(item)
	result["sync"] = reconcile(company)
	result["backup"] = str(backup)
	if commit: frappe.db.commit()
	return result


if __name__ == "__main__":
	import sys
	frappe.init(site="hrms.localhost", sites_path="/home/frappe/frappe-bench/sites")
	frappe.connect()
	frappe.set_user("Administrator")
	try: print(frappe.as_json(run(apply="--apply" in sys.argv)))
	finally: frappe.destroy()
